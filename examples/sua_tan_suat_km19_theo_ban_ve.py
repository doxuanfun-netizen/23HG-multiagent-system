# -*- coding: utf-8 -*-
"""
Sửa file theo dõi tần suất thí nghiệm Cầu Km19 theo bản vẽ thẩm định (logic: tools/tan_suat_theo_ban_ve.py).

Dùng cho file làm việc của đơn vị (KHÔNG commit vào repo). Bản trong repo thì sinh thẳng bằng bộ sinh
tools/generate_tan_suat_thi_nghiem_excel.py — bộ sinh gọi cùng bước sửa này (apply_ban_ve).

Đầu vào : file Excel theo dõi tần suất (--vao)
          examples/data/km19_tan_suat_theo_ban_ve.json   quy tắc dòng bê tông → trang + nhãn bản vẽ
          BBS đã sửa theo bản vẽ + du_lieu.json (HSTK đã OCR)
Đầu ra  : bản đã sửa (--ra): bê tông theo bản vẽ, thép dựng lại từ BBS, công thức tần suất sửa, DM sinh lại,
          sheet NGHI_VAN (cột Đã soát / Ghim).

Chạy:  python examples/sua_tan_suat_km19_theo_ban_ve.py --vao <file.xlsx> --ra <file_moi.xlsx>
"""

import argparse
import json
import os
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from tools.tan_suat_theo_ban_ve import DEFAULT_CONFIG, apply_ban_ve  # noqa: E402


def run(vao: str, ra: str, config: str = DEFAULT_CONFIG, ngay: str = None) -> dict:
    from tools.office_scrub import scrub
    wb = openpyxl.load_workbook(vao)
    cached_wb = openpyxl.load_workbook(vao, data_only=True)
    # ngày đổ bê tông (giá trị Excel lưu sẵn trong file — nhiều ô lấy từ sheet tiến độ) để xếp thép theo thời gian
    with open(config, encoding="utf-8") as f:
        sheet_bt = json.load(f)["sheet_be_tong"]
    cached = cached_wb[sheet_bt]
    bt_dates = {r: cached.cell(r, 3).value for r in range(1, cached.max_row + 1)}
    stats = apply_ban_ve(wb, config, ngay, bt_dates)
    for sh in wb.sheetnames:
        if sh not in cached_wb.sheetnames:
            continue
        ws_f, ws_v = wb[sh], cached_wb[sh]
        for r in range(1, ws_f.max_row + 1):
            for c in range(1, ws_f.max_column + 1):
                v = ws_f.cell(r, c).value
                if isinstance(v, str) and v.startswith("=") and "[" in v and "]" in v:
                    ws_f.cell(r, c).value = ws_v.cell(r, c).value
    wb.save(ra)
    scrub(ra, lien_ket_ngoai=True)
    return stats


def main(argv=None):
    ap = argparse.ArgumentParser(description="Sửa file tần suất thí nghiệm Km19 theo bản vẽ thẩm định")
    ap.add_argument("--vao", required=True, help="file cần sửa")
    ap.add_argument("--ra", required=True, help="file ghi ra (nên khác --vao để giữ bản gốc)")
    ap.add_argument("--config", default=DEFAULT_CONFIG)
    ap.add_argument("--ngay", default=None)
    a = ap.parse_args(argv)
    st = run(a.vao, a.ra, a.config, a.ngay)
    print(f"-> {a.ra}: {st['dong_thep']} dòng thép, {st['lan_bt']} lần lấy mẫu vật liệu bê tông, {st['lan_thep']} lần "
          f"thép, DM {st['dm']} dòng, {st['nghi_van']} nghi vấn")
    return 0


if __name__ == "__main__":
    sys.exit(main())
