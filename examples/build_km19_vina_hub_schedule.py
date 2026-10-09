"""Tiến độ cầu Km19+529.080 chuẩn Vina Alpha (100% công thức sống) từ Master template.

I/J = công thức theo liên kết FS/SS+lag ở cột K; Gantt tuần theo ô ngày khởi công;
đường găng tính bằng công thức (LF/LS/TF). Công tắc Gantt: 1=thanh, 2=nhân công, 3=ca.
Dùng:  python examples/build_km19_vina_hub_schedule.py --start 2026-03-15 --deadline 2026-12-30 --out <file.xlsx>
"""
import argparse
import re
import shutil
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "templates" / "Ho_So_KCS_QS_TienDo_Cau_Km19+529.080.xlsx"
SHEET = "TIEN_DO_THI_CONG_WBS"
WEEKS = 44
G0 = 13  # cột M
DATE_FMT = "DD/MM/YYYY"


def parse_preds(txt):
    out = []
    for part in re.split(r"[;,]", str(txt or "")):
        m = re.match(r"\s*([\d.]+)\s*(FS|SS)?\s*(?:([+-]\d+)d?)?\s*$", part)
        if m:
            out.append((m.group(1), m.group(2) or "FS", int(m.group(3) or 0)))
    return out


def build(out, start, deadline, template=None):
    shutil.copy(template or TEMPLATE, out)
    wb = openpyxl.load_workbook(out)
    ws = wb[SHEET]
    rows, wbs_row = [], {}
    for r in range(7, ws.max_row + 1):
        w = ws.cell(r, 1).value
        if w is not None and ws.cell(r, 8).value is not None:
            rows.append(r)
            wbs_row[str(w).strip()] = r
    preds = {r: parse_preds(ws.cell(r, 11).value) for r in rows}
    succ = {r: [] for r in rows}
    for r, ps in preds.items():
        for pid, t, lag in ps:
            succ[wbs_row[pid]].append((r, t, lag))

    # Gỡ Gantt tĩnh cũ
    for rng in [m for m in ws.merged_cells.ranges if m.min_col >= G0 and m.min_row >= 5]:
        ws.unmerge_cells(str(rng))
    for r in range(5, ws.max_row + 1):
        for c in range(G0, G0 + 40):
            cell = ws.cell(r, c)
            cell.value = None
            cell.fill = PatternFill()

    # Khối điều khiển
    S, D, SW = "$R$2", "$R$3", "$R$4"
    lab = Font(bold=True)
    for cell, text in (("N2", "Ngày khởi công"), ("N3", "Hạn hoàn thành (mục tiêu)"),
                       ("N4", "Gantt: 1=thanh 2=nhân công 3=ca"),
                       ("V2", "Hoàn thành (tính toán)"), ("V3", "Dự phòng so với hạn (ngày)"),
                       ("AD2", "Tổng thời gian (ngày)"), ("AD3", "Trạng thái")):
        ws[cell] = text
        ws[cell].font = lab
    first, last = rows[0], rows[-1]
    ws["R2"], ws["R3"], ws["R4"] = start, deadline, 1
    ws["R2"].number_format = ws["R3"].number_format = DATE_FMT
    ws["Z2"] = f"=MAX($J${first}:$J${last})"
    ws["Z2"].number_format = DATE_FMT
    ws["Z3"] = f"={D}-Z2"
    ws["AH2"] = f"=Z2-{S}+1"
    ws["AH3"] = '=IF(Z2<=' + D + ',"ĐẠT","TRỄ")'
    for c in ("R2", "R3", "R4"):
        ws[c].fill = PatternFill("solid", fgColor="FFF2CC")

    BF, BG, BH = G0 + WEEKS + 1, G0 + WEEKS + 2, G0 + WEEKS + 3  # LF, LS, TF
    cLF, cLS, cTF = L(BF), L(BG), L(BH)
    for c, t in ((BF, "LF"), (BG, "LS"), (BH, "TF")):
        ws.cell(6, c, t).font = lab
    ws.cell(5, BF, "Đường găng (công thức)").font = lab

    for r in rows:
        ps = preds[r]
        if not ps:
            ws.cell(r, 9, f"={S}")
        else:
            terms = []
            for pid, t, lag in ps:
                pr = wbs_row[pid]
                terms.append(f"I{pr}{lag:+d}" if t == "SS" else f"J{pr}+1{lag:+d}")
            ws.cell(r, 9, "=MAX(" + ",".join(terms) + ")" if len(terms) > 1 else "=" + terms[0])
        ws.cell(r, 10, f"=I{r}+H{r}-1")
        ws.cell(r, 9).number_format = ws.cell(r, 10).number_format = DATE_FMT
        # Hậu tiến
        if not succ[r]:
            lf = f"=MAX($J${first}:$J${last})"
        else:
            ts = [f"{cLS}{s}-1-{lag}" if t == "FS" else f"{cLS}{s}-({lag})+H{r}-1"
                  for s, t, lag in succ[r]]
            lf = "=MIN(" + ",".join(ts) + ")" if len(ts) > 1 else "=" + ts[0]
        ws.cell(r, BF, lf)
        ws.cell(r, BG, f"={cLF}{r}-H{r}+1")
        ws.cell(r, BH, f"={cLS}{r}-I{r}")
        for c in (BF, BG):
            ws.cell(r, c).number_format = DATE_FMT
        ws.cell(r, 12, f'=IF({cTF}{r}=0,"YES","NO")')

    # Gantt tuần
    hdr_fill = PatternFill("solid", fgColor="244061")
    for k in range(WEEKS):
        c = G0 + k
        ws.column_dimensions[L(c)].width = 5.5
        h = ws.cell(6, c, f"={S}+7*{k}")
        h.number_format = "DD/MM"
        h.fill, h.font = hdr_fill, Font(color="FFFFFF", size=8)
        h.alignment = Alignment(horizontal="center", text_rotation=90)
        ws.cell(5, c, f"=\"T\"&{k + 1}").font = Font(size=8, bold=True)
        col = L(c)
        for r in rows:
            ws.cell(r, c, f'=IF(AND($I{r}<={col}$6+6,$J{r}>={col}$6),'
                          f'IF({SW}=1,"█",IF({SW}=2,$G{r},$G{r}*2)),"")')
            ws.cell(r, c).alignment = Alignment(horizontal="center")
            ws.cell(r, c).font = Font(size=8, color="FFFFFF")
    ws.row_dimensions[6].height = 38
    rng = f"{L(G0)}{first}:{L(G0 + WEEKS - 1)}{last}"
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'AND({L(G0)}{first}<>"",$L{first}="YES")'], fill=PatternFill("solid", bgColor="C00000")))
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'AND({L(G0)}{first}<>"",$L{first}="NO")'], fill=PatternFill("solid", bgColor="4F81BD")))

    # Ngày hiển thị đúng ở sheet KCS
    kcs = wb["HOSO_KCS_NGHIEM_THU"] if "HOSO_KCS_NGHIEM_THU" in wb.sheetnames else None
    if kcs is not None:
        for row in kcs.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith(f"={SHEET}!") and \
                        re.search(r"![IJ]\d+$", cell.value):
                    cell.number_format = DATE_FMT
    wb.save(out)
    return out


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--start", default="2026-03-15")
    ap.add_argument("--deadline", default="2026-12-30")
    ap.add_argument("--out", required=True)
    ap.add_argument("--template", default=None, help="Master .xlsx gốc (mặc định: template trong repo)")
    a = ap.parse_args()
    print(build(a.out, date.fromisoformat(a.start), date.fromisoformat(a.deadline), a.template))
