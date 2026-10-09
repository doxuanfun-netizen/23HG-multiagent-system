# -*- coding: utf-8 -*-
"""
QS EXPORT — Xuất bảng tổng hợp dự toán G_XD và bảng chi tiết công tác (.xlsx)
Giá trị là số đã tính (không phải công thức) kèm nguồn của từng tỷ lệ để đối chiếu.
"""

from __future__ import annotations

from tools.money import round_vnd
from tools.qs_loader import RATE_LABELS, QSEstimate


def write_gxd_workbook(path: str, est: QSEstimate) -> None:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    def header(ws, row, values):
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=c, value=v)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F3A5E")
            cell.alignment = Alignment(wrap_text=True, vertical="center")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TONG_HOP_GXD"
    ws["A1"] = "BẢNG TỔNG HỢP KINH PHÍ DỰ TOÁN XÂY DỰNG (G_XD) — TT 36/2026/TT-BXD"
    ws["A2"] = f"Nguồn bảng QS: {est.source} — {len(est.items)} công tác"
    header(ws, 4, ["TT", "Khoản mục chi phí", "Cách tính", "Ký hiệu", "Giá trị (VNĐ)", "Nguồn tỷ lệ"])
    r = est.rates
    src = est.rate_sources
    rows = [
        ("1", "Chi phí trực tiếp", "Σ khối lượng × đơn giá", "T", est.T, est.source),
        ("2", "Chi phí gián tiếp", "C + LT + TT", "GT", est.GT, ""),
        ("", "  - Chi phí chung", f"{r['chung'] * 100:g}% × T", "C", est.GT_components["chung"], src.get("chung", "")),
        ("", "  - Chi phí nhà tạm để ở và điều hành thi công", f"{r['nha_tam'] * 100:g}% × T", "LT",
         est.GT_components["nha_tam"], src.get("nha_tam", "")),
        ("", "  - Chi phí một số công việc không xác định được KL từ thiết kế", f"{r['kxd'] * 100:g}% × T", "TT",
         est.GT_components["kxd"], src.get("kxd", "")),
        ("3", "Thu nhập chịu thuế tính trước", f"{r['tl'] * 100:g}% × (T + GT)", "TL", est.TL, src.get("tl", "")),
        ("4", "Chi phí xây dựng trước thuế", "T + GT + TL", "G", est.G, ""),
        ("5", "Thuế giá trị gia tăng", f"{r['vat'] * 100:g}% × G", "VAT", est.VAT, src.get("vat", "")),
        ("6", "TỔNG CỘNG DỰ TOÁN CHI PHÍ XÂY DỰNG", "G + VAT", "G_XD", est.G_XD, ""),
    ]
    for row in rows:
        ws.append(list(row))
    for row in ws.iter_rows(min_row=5, max_row=4 + len(rows), min_col=5, max_col=5):
        row[0].number_format = "#,##0"
    if est.warnings:
        ws.cell(row=6 + len(rows), column=1, value="CẢNH BÁO ĐỐI CHIẾU")
        for k, w in enumerate(est.warnings, start=7 + len(rows)):
            ws.cell(row=k, column=2, value=w)
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 26
    ws.column_dimensions["E"].width = 20
    ws.column_dimensions["F"].width = 50

    ws = wb.create_sheet("CHI_TIET_CONG_TAC")
    header(ws, 1, ["Dòng file", "TT", "Mã hiệu", "Nội dung công tác", "ĐVT", "Khối lượng", "Đơn giá",
                   "Thành tiền (KL×ĐG)", "Thành tiền trong file", "Tổng diễn giải", "Phần"])
    for i in est.items:
        ws.append([i.row, i.stt, i.code, i.description, i.unit, i.quantity, i.unit_price, round_vnd(est.line_amount(i)),
                   i.file_amount, i.detail_quantity, i.section])
    ws.column_dimensions["D"].width = 60
    for col in ("G", "H", "I"):
        for cell in ws[col][1:]:
            cell.number_format = "#,##0"
    wb.save(path)
