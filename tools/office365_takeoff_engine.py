# -*- coding: utf-8 -*-
"""
OFFICE 365 TAKEOFF ENGINE — ĐỘNG CƠ ĐO BÓC & QUẢN TRỊ KHỐI LƯỢNG MICROSOFT 365 ENTERPRISE
Tích hợp trong Hệ sinh thái Multi-Agent 23HG (Pure Python, Zero Dead Numbers).

Cung cấp:
  1. Đăng ký tự động Bộ hàm kỹ thuật AEC LAMBDA tùy biến vào Excel Name Manager:
     - V_PRISM: Tính thể tích lăng trụ chữ nhật (bệ, thân hộp, xà mũ, đệm).
     - V_CYLINDER: Tính thể tích hình trụ tròn (cọc khoan nhồi, cột tròn).
     - V_FRUSTUM: Tính thể tích hình chóp cụt vát dốc (bệ mố/trụ, móng đơn vát).
     - S_FORMWORK_BOX: Diện tích ván khuôn thành hộp 4 mặt.
     - S_FORMWORK_TRI: Diện tích ván khuôn tam giác / vát góc.
     - STEEL_RATIO: Hàm lượng thép bình quân (kg/m3).
     - V_AVERAGE_END: Khối lượng đào đắp mặt cắt ngang 2 đầu (Average-End-Area).
     - REBAR_WEIGHT: Trọng lượng thanh thép theo TCVN 1651:2018.
  2. Tạo lập công thức tự diễn giải LET() tuân thủ OpenXML chuẩn hóa (_xlfn.LET / _xlpm.*).
  3. Tạo lập công thức truy vấn linh hoạt XLOOKUP() chống đứt gãy liên kết khi chèn/xóa dòng.
  4. Tạo lập Trang bìa Bảng điều hành & Kiểm toán chéo động (00_DASHBOARD_365) với KPI Cards & Audit Grid.
  5. Chèn và định vị minh chứng trực quan từ bản vẽ CAD (Visual Proof Anchoring).
"""

from __future__ import annotations

import os
import re
from typing import Any, Dict, List, Optional, Sequence, Tuple, Union

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.drawing.image import Image as OpenpyxlImage
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


# =============================================================================
# 1. BỘ HÀM KỸ THUẬT AEC LAMBDA CHUẨN MICROSOFT 365
# =============================================================================

AEC_LAMBDA_DEFINITIONS = [
    (
        "V_PRISM",
        "_xlfn.LAMBDA(_xlpm.qty,_xlpm.L,_xlpm.W,_xlpm.H,_xlpm.qty*_xlpm.L*_xlpm.W*_xlpm.H)",
        "Tính thể tích hình lăng trụ chữ nhật: qty * L * W * H (m3)"
    ),
    (
        "V_CYLINDER",
        "_xlfn.LAMBDA(_xlpm.qty,_xlpm.L,_xlpm.D,_xlpm.qty*PI()*(_xlpm.D/2)^2*_xlpm.L)",
        "Tính thể tích hình trụ tròn / cọc khoan nhồi: qty * PI * (D/2)^2 * L (m3)"
    ),
    (
        "V_FRUSTUM",
        "_xlfn.LAMBDA(_xlpm.qty,_xlpm.H,_xlpm.S1,_xlpm.S2,_xlpm.qty*(_xlpm.H/3)*(_xlpm.S1+_xlpm.S2+SQRT(_xlpm.S1*_xlpm.S2)))",
        "Tính thể tích hình chóp cụt 2 đáy S1, S2: qty * (H/3) * (S1 + S2 + SQRT(S1*S2)) (m3)"
    ),
    (
        "S_FORMWORK_BOX",
        "_xlfn.LAMBDA(_xlpm.qty,_xlpm.L,_xlpm.W,_xlpm.H,_xlpm.qty*2*(_xlpm.L+_xlpm.W)*_xlpm.H)",
        "Tính diện tích ván khuôn thành hộp 4 mặt: qty * 2 * (L + W) * H (m2)"
    ),
    (
        "S_FORMWORK_TRI",
        "_xlfn.LAMBDA(_xlpm.qty,_xlpm.B,_xlpm.H,_xlpm.qty*_xlpm.B*_xlpm.H*0.5)",
        "Tính diện tích ván khuôn tam giác / vát góc: qty * B * H * 0.5 (m2)"
    ),
    (
        "STEEL_RATIO",
        "_xlfn.LAMBDA(_xlpm.steel_kg,_xlpm.conc_m3,_xlpm.steel_kg/_xlpm.conc_m3)",
        "Tính hàm lượng cốt thép bình quân: steel_kg / conc_m3 (kg/m3)"
    ),
    (
        "V_AVERAGE_END",
        "_xlfn.LAMBDA(_xlpm.F1,_xlpm.F2,_xlpm.L,((_xlpm.F1+_xlpm.F2)/2)*_xlpm.L)",
        "Tính khối lượng đào đắp mặt cắt ngang 2 đầu (Average-End-Area): ((F1 + F2) / 2) * L (m3)"
    ),
    (
        "REBAR_WEIGHT",
        "_xlfn.LAMBDA(_xlpm.L,_xlpm.N,_xlpm.d,_xlpm.L*_xlpm.N*(0.006165*_xlpm.d^2))",
        "Tính trọng lượng cốt thép theo TCVN 1651:2018: L * N * (0.006165 * d^2) (kg)"
    ),
]


def register_aec_lambdas(wb: Any) -> List[str]:
    """
    Đăng ký toàn bộ bộ hàm kỹ thuật AEC LAMBDA vào Name Manager của Workbook.
    Tuân thủ cú pháp OpenXML chuẩn với namespace _xlfn. và parameter prefix _xlpm.
    """
    if not HAS_OPENPYXL:
        raise RuntimeError("openpyxl is required to register AEC LAMBDA functions.")

    registered = []
    for name, formula, _desc in AEC_LAMBDA_DEFINITIONS:
        if name in wb.defined_names:
            del wb.defined_names[name]
        dn = DefinedName(name, attr_text=formula)
        wb.defined_names.add(dn)
        registered.append(name)
    return registered


# =============================================================================
# 2. BỘ TẠO CÔNG THỨC LET() & XLOOKUP() CHUẨN OPENXML
# =============================================================================

def build_let_formula(variables: Dict[str, str], expression: str) -> str:
    """
    Tạo chuỗi công thức LET() chuẩn Microsoft 365 dùng trong openpyxl.
    Mỗi biến khai báo được tự động thêm tiền tố _xlpm. để tránh lỗi XML parser của Excel.
    
    Ví dụ:
      variables = {'qty': 'D16', 'L': 'E16', 'W': 'F16', 'H': 'G16'}
      expression = 'qty * L * W * H'
      -> '=_xlfn.LET(_xlpm.qty, D16, _xlpm.L, E16, _xlpm.W, F16, _xlpm.H, G16, _xlpm.qty * _xlpm.L * _xlpm.W * _xlpm.H)'
    """
    let_parts = []
    # Chuẩn bị regex thay thế biến số trong biểu thức
    sorted_vars = sorted(variables.keys(), key=lambda k: len(k), reverse=True)
    expr_transformed = expression
    for v in sorted_vars:
        # Thay thế độc lập word boundary
        pattern = r"\b" + re.escape(v) + r"\b"
        expr_transformed = re.sub(pattern, f"_xlpm.{v}", expr_transformed)

    for k, v in variables.items():
        let_parts.append(f"_xlpm.{k}, {v}")

    args_str = ", ".join(let_parts) + f", {expr_transformed}"
    return f"=_xlfn.LET({args_str})"


def build_xlookup_formula(
    lookup_value: str,
    lookup_range: str,
    return_range: str,
    if_not_found: Any = 0,
    match_mode: int = 2
) -> str:
    """
    Tạo chuỗi công thức XLOOKUP() tìm kiếm động theo chuỗi đại diện (Wildcard matching).
    Mặc định match_mode=2 (cho phép wildcard như "*TỔNG BÊ TÔNG*").
    
    Ví dụ:
      build_xlookup_formula('"*TỔNG BÊ TÔNG MỐ C30*"', "'01_MO_M1_M2'!$B:$B", "'01_MO_M1_M2'!$D:$D")
      -> '=_xlfn.XLOOKUP("*TỔNG BÊ TÔNG MỐ C30*", \'01_MO_M1_M2\'!$B:$B, \'01_MO_M1_M2\'!$D:$D, 0, 2)'
    """
    return f"=_xlfn.XLOOKUP({lookup_value}, {lookup_range}, {return_range}, {if_not_found}, {match_mode})"


# =============================================================================
# 3. BẢNG ĐIỀU HÀNH & KIỂM TOÁN TỔNG HỢP (EXECUTIVE 365 DASHBOARD)
# =============================================================================

def build_executive_365_dashboard(
    wb: Any,
    project_title: str,
    kpis: List[Dict[str, Any]],
    breakdown_rows: Optional[List[Sequence[Any]]] = None,
    audit_items: Optional[List[Dict[str, Any]]] = None,
    sheet_name: str = "00_DASHBOARD_365",
    index: int = 0
) -> Any:
    """
    Xây dựng trang Dashboard điều hành & kiểm toán cao cấp chuẩn Microsoft 365.
    
    Tham số:
      - wb: Workbook openpyxl
      - project_title: Tiêu đề công trình / dự án
      - kpis: Danh sách tối đa 5 thẻ KPI [{'title': ..., 'formula': ..., 'unit': ..., 'fmt': ...}]
      - breakdown_rows: Dòng bảng tổng hợp cấu kiện (nếu có)
      - audit_items: Danh sách hạng mục kiểm toán [{'stt': ..., 'name': ..., 'unit': ..., 'cad_formula': ..., 'design_val': ..., 'note': ...}]
    """
    if not HAS_OPENPYXL:
        raise RuntimeError("openpyxl is required to build dashboard.")

    if sheet_name in wb.sheetnames:
        wb.remove(wb[sheet_name])

    ws = wb.create_sheet(title=sheet_name, index=index)
    ws.sheet_view.showGridLines = True

    # Bảng màu thiết kế chuẩn Executive Navy & Emerald
    f_title = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
    f_subtitle = Font(name="Calibri", size=10, italic=True, color="93C5FD")
    f_kpi_label = Font(name="Calibri", size=9, bold=True, color="1E3A8A")
    f_kpi_val = Font(name="Calibri", size=18, bold=True, color="0F172A")
    f_kpi_unit = Font(name="Calibri", size=9, italic=True, color="475569")
    f_tbl_hdr = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
    f_bold_item = Font(name="Calibri", size=9, bold=True, color="1E3A8A")
    f_norm = Font(name="Calibri", size=9, color="0F172A")
    f_check = Font(name="Calibri", size=9, bold=True, color="059669")

    fill_navy = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    fill_sub_banner = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    fill_blue_hdr = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    fill_kpi_bg = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_kpi_top = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
    fill_sub_row = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
    fill_audit_ok = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    b_thin = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # 1. Banner Tiêu đề
    ws.merge_cells("A1:I2")
    ws["A1"] = f"BẢNG ĐIỀU HÀNH & KIỂM TOÁN TỔNG HỢP — {project_title.upper()}"
    ws["A1"].font = f_title
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws["A1"].fill = fill_navy

    ws.merge_cells("A3:I3")
    ws["A3"] = "Hệ thống đo bóc giải tích động 100% Dynamic Arrays & Functional Engineering Formulas (LET, XLOOKUP, LAMBDA) | Không số chết"
    ws["A3"].font = f_subtitle
    ws["A3"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws["A3"].fill = fill_sub_banner

    # 2. Thẻ KPI Động (Tối đa 5 thẻ)
    col_pairs = [("A", "B"), ("C", "D"), ("E", "F"), ("G", "G"), ("H", "I")]
    for idx, kpi in enumerate(kpis[:5]):
        c1, c2 = col_pairs[idx]
        title = kpi.get("title", "")
        formula = kpi.get("formula", 0)
        unit = kpi.get("unit", "")
        fmt = kpi.get("fmt", "#,##0.00")

        if c1 != c2:
            ws.merge_cells(f"{c1}5:{c2}5")
            ws.merge_cells(f"{c1}6:{c2}6")
            ws.merge_cells(f"{c1}7:{c2}7")

        ws[f"{c1}5"] = title
        ws[f"{c1}5"].font = f_kpi_label
        ws[f"{c1}5"].fill = fill_kpi_top
        ws[f"{c1}5"].alignment = Alignment(horizontal="center", vertical="center")

        ws[f"{c1}6"] = formula
        ws[f"{c1}6"].font = f_kpi_val
        ws[f"{c1}6"].fill = fill_kpi_bg
        ws[f"{c1}6"].alignment = Alignment(horizontal="center", vertical="center")
        ws[f"{c1}6"].number_format = fmt

        ws[f"{c1}7"] = f"Đơn vị: {unit}"
        ws[f"{c1}7"].font = f_kpi_unit
        ws[f"{c1}7"].fill = fill_kpi_bg
        ws[f"{c1}7"].alignment = Alignment(horizontal="center", vertical="center")

        cols_to_border = [c1] if c1 == c2 else [c1, c2]
        for cn in cols_to_border:
            for rn in range(5, 8):
                ws[f"{cn}{rn}"].border = b_thin

    cur_row = 9

    # 3. Bảng Tổng hợp phân rã kết cấu (Nếu được cung cấp)
    if breakdown_rows:
        ws.cell(cur_row, 1, "1. BẢNG TỔNG HỢP KHỐI LƯỢNG ĐỘNG CÁC CẤU KIỆN (DYNAMIC AUDIT)").font = Font(
            name="Calibri", size=11, bold=True, color="1E40AF"
        )
        cur_row += 1

        headers = ["STT", "HẠNG MỤC CÔNG VIỆC", "ĐVT", "CẤU KIỆN 1", "CẤU KIỆN 2", "CẤU KIỆN 3", "CẤU KIỆN 4", "CỌC / KHÁC", "TỔNG CỘNG"]
        for c_idx, h in enumerate(headers, 1):
            c = ws.cell(cur_row, c_idx, h)
            c.font = f_tbl_hdr
            c.fill = fill_blue_hdr
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = b_thin
        cur_row += 1

        for r_data in breakdown_rows:
            is_tot = str(r_data[0]).startswith("T")
            r_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid") if is_tot else (
                fill_sub_row if cur_row % 2 == 1 else fill_white
            )
            r_font = f_bold_item if is_tot else f_norm
            for c_idx, val in enumerate(r_data, 1):
                cell = ws.cell(cur_row, c_idx, val)
                cell.font = r_font
                cell.fill = r_fill
                cell.border = b_thin
                if c_idx in (1, 3):
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif c_idx == 2:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    cell.number_format = "#,##0.00"
            cur_row += 1
        cur_row += 2

    # 4. Bảng Đối soát & Kiểm toán Chéo tự động (Dynamic Audit Grid)
    if audit_items:
        ws.cell(cur_row, 1, "2. BẢNG ĐỐI CHIẾU KIỂM TOÁN CHÉO VỚI HỒ SƠ THIẾT KẾ DUYỆT (AUDIT RECONCILIATION)").font = Font(
            name="Calibri", size=11, bold=True, color="1E40AF"
        )
        cur_row += 1

        audit_headers = ["STT", "HẠNG MỤC KIỂM SOÁT", "ĐVT", "KL ĐO BÓC CAD LIVE", "KL HỒ SƠ DUYỆT", "CHÊNH LỆCH (DELTA)", "TỶ LỆ %", "TRẠNG THÁI KIỂM TOÁN", "GHI CHÚ ĐỐI SOÁT"]
        for c_idx, h in enumerate(audit_headers, 1):
            c = ws.cell(cur_row, c_idx, h)
            c.font = f_tbl_hdr
            c.fill = PatternFill(start_color="047857", end_color="047857", fill_type="solid")
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = b_thin
        cur_row += 1

        for item in audit_items:
            stt = item.get("stt", "")
            name = item.get("name", "")
            unit = item.get("unit", "")
            cad_f = item.get("cad_formula", 0)
            des_val = item.get("design_val", 0.0)
            note = item.get("note", "")

            ws.cell(cur_row, 1, stt).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(cur_row, 2, name).alignment = Alignment(horizontal="left", vertical="center")
            ws.cell(cur_row, 3, unit).alignment = Alignment(horizontal="center", vertical="center")

            c_cad = ws.cell(cur_row, 4, cad_f)
            c_cad.number_format = "#,##0.00"
            c_cad.alignment = Alignment(horizontal="right", vertical="center")

            c_des = ws.cell(cur_row, 5, des_val)
            c_des.number_format = "#,##0.00"
            c_des.alignment = Alignment(horizontal="right", vertical="center")

            # Delta = Live - Design
            c_delta = ws.cell(cur_row, 6, f"=D{cur_row}-E{cur_row}")
            c_delta.number_format = "#,##0.00"
            c_delta.alignment = Alignment(horizontal="right", vertical="center")

            # Delta %
            c_pct = ws.cell(cur_row, 7, f"=_xlfn.IF(E{cur_row}=0, 0, (D{cur_row}-E{cur_row})/E{cur_row})")
            c_pct.number_format = "0.00%"
            c_pct.alignment = Alignment(horizontal="right", vertical="center")

            # Status (Sai lệch < 0.05 -> Trùng khớp 100%)
            c_stat = ws.cell(cur_row, 8, f'=_xlfn.IF(ABS(F{cur_row})<0.05, "✓ TRÙNG KHỚP 100%", "⚠ CẢNH BÁO KIỂM TOÁN")')
            c_stat.alignment = Alignment(horizontal="center", vertical="center")
            c_stat.font = f_check

            c_note = ws.cell(cur_row, 9, note)
            c_note.alignment = Alignment(horizontal="left", vertical="center")

            for col in range(1, 10):
                c = ws.cell(cur_row, col)
                c.border = b_thin
                c.font = f_bold_item if col in (2, 8) else f_norm
                c.fill = fill_audit_ok if cur_row % 2 == 0 else fill_white

            cur_row += 1

    # Căn chỉnh độ rộng cột
    widths = {'A': 8, 'B': 36, 'C': 10, 'D': 18, 'E': 18, 'F': 18, 'G': 18, 'H': 20, 'I': 46}
    for col_letter, w in widths.items():
        ws.column_dimensions[col_letter].width = w

    return ws


# =============================================================================
# 4. CHÈN HÌNH ẢNH MINH CHỨNG CAD (VISUAL PROOF ANCHOR)
# =============================================================================

def embed_cad_proof_images(
    ws: Any,
    image_specs: List[Dict[str, Any]]
) -> int:
    """
    Chèn hình ảnh trích xuất từ bản vẽ CAD vào ô chỉ định của Worksheet.
    
    Tham số image_specs:
      [
        {"path": ".../mo_m1.png", "anchor": "M5", "width": 720, "height": 480},
        ...
      ]
    """
    if not HAS_OPENPYXL:
        return 0

    count = 0
    for spec in image_specs:
        img_path = spec.get("path", "")
        anchor = spec.get("anchor", "M5")
        w = spec.get("width", 720)
        h = spec.get("height", 480)

        if os.path.exists(img_path):
            img = OpenpyxlImage(img_path)
            img.width = w
            img.height = h
            ws.add_image(img, anchor)
            count += 1
    return count
