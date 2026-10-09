# -*- coding: utf-8 -*-
"""
Tự động xuất bảng tính Excel Bóc tách khối lượng 1 đốt (1m dài) cho các loại cống hộp:
- Cống hộp 2.0x2.0m
- Cống hộp 3.0x3.0m
- Cống hộp đôi 2x(3.0x3.0m)
100% CÔNG THỨC SỐNG LIÊN KẾT ĐỘNG, Chuẩn Vincons / AEC Multi-Agent System.
"""

import os
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

output_file = os.environ.get("AEC_OUTPUT_FILE", os.path.join(os.path.expanduser("~"), "Downloads", "Documents", "2026.09.12.OLP_SD_HT_ChiTietCongHop+TamGiamTai Model (1)_Marker", "Boc_Tach_Khoi_Luong_1m_Dot_Cong_A5.xlsx"))

wb = openpyxl.Workbook()
wb.remove(wb.active)  # Remove default sheet

# Fonts
FONT_TITLE = Font(name="Calibri", size=14, bold=True, color="1B365D")
FONT_SUBTITLE = Font(name="Calibri", size=11, italic=True, color="444444")
FONT_HEADER = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
FONT_SECTION = Font(name="Calibri", size=11, bold=True, color="1B365D")
FONT_BOLD = Font(name="Calibri", size=10, bold=True)
FONT_REGULAR = Font(name="Calibri", size=10)
FONT_NOTE = Font(name="Calibri", size=9, italic=True, color="555555")

# Fills
FILL_NAVY = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
FILL_HEADER_BLUE = PatternFill(start_color="2A52BE", end_color="2A52BE", fill_type="solid")
FILL_SECTION = PatternFill(start_color="E6EEF8", end_color="E6EEF8", fill_type="solid")
FILL_HIGHLIGHT = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
FILL_TOTAL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

# Alignments
ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")

# Borders
BORDER_THIN = Border(
    left=Side(style='thin', color='BFBFBF'),
    right=Side(style='thin', color='BFBFBF'),
    top=Side(style='thin', color='BFBFBF'),
    bottom=Side(style='thin', color='BFBFBF')
)

def style_cell(cell, font=None, fill=None, alignment=None, border=BORDER_THIN, num_format=None):
    if font: cell.font = font
    if fill: cell.fill = fill
    if alignment: cell.alignment = alignment
    if border: cell.border = border
    if num_format: cell.number_format = num_format

# ==========================================
# FUNCTION: BUILD DETAIL SHEET
# ==========================================
def build_detail_sheet(ws, title, b_tr, h_tr, t_nap, t_day, t_thanh, vat_x, vat_y, t_lot, b_morong_lot, is_double=False, t_giua=0.3):
    ws.views.sheetView[0].showGridLines = True
    
    # Title Header
    ws.cell(row=2, column=2, value=f"CHI TIẾT THÔNG SỐ VÀ BÓC TÁCH KHỐI LƯỢNG: {title.upper()}").font = FONT_TITLE
    ws.cell(row=3, column=2, value="Đơn vị tính cho L = 1.0 mét dài cống | 100% Công thức sống liên kết động").font = FONT_SUBTITLE
    
    # Section I: Parameters
    ws.cell(row=5, column=2, value="I. CÁC THÔNG SỐ HÌNH HỌC ĐẦU VÀO (INPUT PARAMETERS)").font = FONT_SECTION
    ws.merge_cells(start_row=5, start_column=2, end_row=5, end_column=7)
    ws.cell(row=5, column=2).fill = FILL_SECTION
    
    headers_param = ["STT", "Thông số hình học", "Ký hiệu", "Giá trị", "Đơn vị", "Ghi chú kỹ thuật"]
    for col_idx, h in enumerate(headers_param, start=2):
        c = ws.cell(row=6, column=col_idx, value=h)
        style_cell(c, font=FONT_HEADER, fill=FILL_HEADER_BLUE, alignment=ALIGN_CENTER)
    ws.row_dimensions[6].height = 26
    
    params = [
        (1, "Chiều rộng thông thủy", "B_tt", b_tr, "m", "Kích thước lòng cống"),
        (2, "Chiều cao thông thủy", "H_tt", h_tr, "m", "Kích thước lòng cống"),
        (3, "Chiều dày bản nắp", "t_nap", t_nap, "m", "Bản nắp bê tông chịu lực"),
        (4, "Chiều dày bản đáy", "t_day", t_day, "m", "Bản đáy móng cống"),
        (5, "Chiều dày thành biên (2 bên)", "t_thanh", t_thanh, "m", "Thành bên cống"),
    ]
    if is_double:
        params.append((6, "Chiều dày vách giữa", "t_giua", t_giua, "m", "Vách ngăn đôi 2 ngăn cống"))
    
    params.extend([
        (len(params) + 1, "Kích thước vát góc ngang", "v_x", vat_x, "m", "Góc vát gia cường góc trong"),
        (len(params) + 1, "Kích thước vát góc đứng", "v_y", vat_y, "m", "Góc vát gia cường góc trong"),
        (len(params) + 1, "Chiều dày bê tông lót M100", "t_lot", t_lot, "m", "Lớp lót đáy móng cống"),
        (len(params) + 1, "Độ mở rộng móng lót mỗi bên", "delta_lot", b_morong_lot, "m", "Mở rộng 100mm mỗi bên"),
        (len(params) + 1, "Chiều dài đốt tính toán chuẩn", "L", 1.00, "m", "Chuẩn hóa đốt dài 1.0m"),
        (len(params) + 1, "Định mức thép dự kiến", "DM_thep", 100.0, "kg/m³", "Kg thép trên 1m³ bê tông cống"),
    ])
    
    r = 6
    for p in params:
        r += 1
        ws.row_dimensions[r].height = 20
        c_stt = ws.cell(row=r, column=2, value=p[0])
        c_name = ws.cell(row=r, column=3, value=p[1])
        c_sym = ws.cell(row=r, column=4, value=p[2])
        c_val = ws.cell(row=r, column=5, value=p[3])
        c_unit = ws.cell(row=r, column=6, value=p[4])
        c_note = ws.cell(row=r, column=7, value=p[5])
        
        style_cell(c_stt, font=FONT_REGULAR, alignment=ALIGN_CENTER)
        style_cell(c_name, font=FONT_REGULAR, alignment=ALIGN_LEFT)
        style_cell(c_sym, font=FONT_REGULAR, alignment=ALIGN_CENTER)
        style_cell(c_val, font=FONT_BOLD, alignment=ALIGN_RIGHT, num_format="#,##0.00" if isinstance(p[3], float) and p[3] < 10 else "#,##0.0")
        style_cell(c_unit, font=FONT_REGULAR, alignment=ALIGN_CENTER)
        style_cell(c_note, font=FONT_NOTE, alignment=ALIGN_LEFT)
    
    # Section II: Calculations
    r += 2
    ws.cell(row=r, column=2, value="II. BẢNG TÍNH TOÁN BÓC TÁCH KHỐI LƯỢNG (LIVE FORMULAS)").font = FONT_SECTION
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
    ws.cell(row=r, column=2).fill = FILL_SECTION
    
    r += 1
    headers_calc = ["STT", "Nội dung công việc / Cấu kiện", "Diễn giải công thức toán", "Khối lượng tính toán", "Đơn vị", "Ghi chú kiểm tra"]
    for col_idx, h in enumerate(headers_calc, start=2):
        c = ws.cell(row=r, column=col_idx, value=h)
        style_cell(c, font=FONT_HEADER, fill=FILL_NAVY, alignment=ALIGN_CENTER)
    ws.row_dimensions[r].height = 26
    
    # Parameter cell references
    if not is_double:
        c_Btt = "E7"; c_Htt = "E8"; c_tnap = "E9"; c_tday = "E10"; c_tthanh = "E11"
        c_vx = "E12"; c_vy = "E13"; c_tlot = "E14"; c_dlot = "E15"; c_L = "E16"; c_dmthep = "E17"
    else:
        c_Btt = "E7"; c_Htt = "E8"; c_tnap = "E9"; c_tday = "E10"; c_tthanh = "E11"; c_tgiua = "E12"
        c_vx = "E13"; c_vy = "E14"; c_tlot = "E15"; c_dlot = "E16"; c_L = "E17"; c_dmthep = "E18"

    calc_start_row = r + 1
    
    # Define calculation lines
    # Rows will be:
    # 1. B_ngoai: r1 = calc_start_row
    # 2. H_ngoai: r2 = calc_start_row + 1
    # 3. B_lot:   r3 = calc_start_row + 2
    # 4. BT lot:  r4 = calc_start_row + 3  <-- TARGET 1
    # 5. BT cong: r5 = calc_start_row + 4  <-- TARGET 2
    # 6. VK ngoài: r6 = calc_start_row + 5 <-- TARGET 3
    # 7. VK trong: r7 = calc_start_row + 6 <-- TARGET 4
    # 8. Tổng VK:  r8 = calc_start_row + 7 <-- TARGET 5
    # 9. Cốt thép: r9 = calc_start_row + 8 <-- TARGET 6
    
    r1 = calc_start_row
    r2 = r1 + 1
    r3 = r2 + 1
    r4 = r3 + 1
    r5 = r4 + 1
    r6 = r5 + 1
    r7 = r6 + 1
    r8 = r7 + 1
    r9 = r8 + 1
    
    if not is_double:
        items = [
            (1, "Chiều rộng phủ bì cống (B_ngoai)", "B_tt + 2 * t_thanh", f"={c_Btt}+2*{c_tthanh}", "m", "Kích thước ngang ngoài cùng"),
            (2, "Chiều cao phủ bì cống (H_ngoai)", "H_tt + t_nap + t_day", f"={c_Htt}+{c_tnap}+{c_tday}", "m", "Kích thước đứng ngoài cùng"),
            (3, "Chiều rộng bản đáy móng lót (B_lot)", f"B_ngoai + 2 * delta_lot", f"=E{r1}+2*{c_dlot}", "m", "Mở rộng móng lót 2 bên"),
            (4, "Bê tông lót móng cống M100", f"B_lot * t_lot * L", f"=E{r3}*{c_tlot}*{c_L}", "m³", "Bê tông lót đáy dày 100mm"),
            (5, "Bê tông thân và nắp cống M250", f"(B_ngoai*H_ngoai - B_tt*H_tt + 4*0.5*v_x*v_y)*L", f"=(E{r1}*E{r2}-{c_Btt}*{c_Htt}+4*(0.5*{c_vx}*{c_vy}))*{c_L}", "m³", "Bê tông thành, nắp, đáy và vát góc"),
            (6, "Ván khuôn ngoài thành cống", f"2 * H_ngoai * L", f"=2*E{r2}*{c_L}", "m²", "2 mặt ngoài thành cống"),
            (7, "Ván khuôn trong lòng cống", f"(2*(H_tt-2*v_y) + (B_tt-2*v_x) + 4*SQRT(v_x^2+v_y^2))*L", f"=(2*({c_Htt}-2*{c_vy})+({c_Btt}-2*{c_vx})+4*SQRT({c_vx}^2+{c_vy}^2))*{c_L}", "m²", "2 thành trong + nắp trong + 4 vát góc"),
            (8, "Tổng diện tích ván khuôn cống", f"Ván khuôn ngoài + Ván khuôn trong", f"=E{r6}+E{r7}", "m²", "Tổng diện tích ván khuôn thành và nắp"),
            (9, "Khối lượng cốt thép cống", f"Thể tích BT cống * DM_thep", f"=E{r5}*{c_dmthep}", "kg", "Hàm lượng thép dự kiến 100 kg/m³"),
        ]
    else:
        items = [
            (1, "Chiều rộng phủ bì cống đôi (B_ngoai)", "2 * B_tt + 2 * t_thanh + t_giua", f"=2*{c_Btt}+2*{c_tthanh}+{c_tgiua}", "m", "2 khoang + 2 thành biên + 1 vách giữa"),
            (2, "Chiều cao phủ bì cống (H_ngoai)", "H_tt + t_nap + t_day", f"={c_Htt}+{c_tnap}+{c_tday}", "m", "Kích thước đứng ngoài cùng"),
            (3, "Chiều rộng bản đáy móng lót (B_lot)", f"B_ngoai + 2 * delta_lot", f"=E{r1}+2*{c_dlot}", "m", "Mở rộng móng lót 2 bên"),
            (4, "Bê tông lót móng cống M100", f"B_lot * t_lot * L", f"=E{r3}*{c_tlot}*{c_L}", "m³", "Bê tông lót đáy dày 100mm"),
            (5, "Bê tông thân và nắp cống M250", f"(B_ngoai*H_ngoai - 2*B_tt*H_tt + 8*0.5*v_x*v_y)*L", f"=(E{r1}*E{r2}-2*({c_Btt}*{c_Htt})+8*(0.5*{c_vx}*{c_vy}))*{c_L}", "m³", "Bê tông cống đôi 2 ngăn"),
            (6, "Ván khuôn ngoài thành cống", f"2 * H_ngoai * L", f"=2*E{r2}*{c_L}", "m²", "2 mặt ngoài thành biên"),
            (7, "Ván khuôn trong lòng cống (2 ngăn)", f"2 * [2*(H_tt-2*v_y) + (B_tt-2*v_x) + 4*SQRT(v_x^2+v_y^2)]*L", f"=2*(2*({c_Htt}-2*{c_vy})+({c_Btt}-2*{c_vx})+4*SQRT({c_vx}^2+{c_vy}^2))*{c_L}", "m²", "Ván khuôn mặt trong của 2 ngăn cống"),
            (8, "Tổng diện tích ván khuôn cống", f"Ván khuôn ngoài + Ván khuôn trong", f"=E{r6}+E{r7}", "m²", "Tổng diện tích ván khuôn cống đôi"),
            (9, "Khối lượng cốt thép cống", f"Thể tích BT cống * DM_thep", f"=E{r5}*{c_dmthep}", "kg", "Hàm lượng thép dự kiến 100 kg/m³"),
        ]

    curr_r = calc_start_row - 1
    for it in items:
        curr_r += 1
        ws.row_dimensions[curr_r].height = 22
        stt, name, expr_str, formula, unit, note = it
        
        c_stt = ws.cell(row=curr_r, column=2, value=stt)
        c_name = ws.cell(row=curr_r, column=3, value=name)
        c_expr = ws.cell(row=curr_r, column=4, value=expr_str)
        c_calc = ws.cell(row=curr_r, column=5, value=formula)
        c_unit = ws.cell(row=curr_r, column=6, value=unit)
        c_note = ws.cell(row=curr_r, column=7, value=note)
        
        style_cell(c_stt, font=FONT_REGULAR, alignment=ALIGN_CENTER)
        style_cell(c_name, font=FONT_REGULAR, alignment=ALIGN_LEFT)
        style_cell(c_expr, font=FONT_NOTE, alignment=ALIGN_LEFT)
        style_cell(c_unit, font=FONT_REGULAR, alignment=ALIGN_CENTER)
        style_cell(c_note, font=FONT_NOTE, alignment=ALIGN_LEFT)
        
        # Highlight important output rows
        if stt in [4, 5]: # Bê tông
            style_cell(c_name, font=FONT_BOLD)
            style_cell(c_calc, font=FONT_BOLD, fill=FILL_TOTAL, alignment=ALIGN_RIGHT, num_format="#,##0.000")
        elif stt in [8]: # Tổng VK
            style_cell(c_name, font=FONT_BOLD)
            style_cell(c_calc, font=FONT_BOLD, fill=FILL_TOTAL, alignment=ALIGN_RIGHT, num_format="#,##0.000")
        elif stt in [9]: # Thép
            style_cell(c_name, font=FONT_BOLD)
            style_cell(c_calc, font=FONT_BOLD, fill=FILL_HIGHLIGHT, alignment=ALIGN_RIGHT, num_format="#,##0.0")
        else:
            style_cell(c_calc, font=FONT_REGULAR, alignment=ALIGN_RIGHT, num_format="#,##0.000" if unit in ["m³", "m²"] else "#,##0.00")

    return {
        "bt_lot_row": r4,
        "bt_cong_row": r5,
        "vk_ngoai_row": r6,
        "vk_trong_row": r7,
        "vk_tong_row": r8,
        "thep_row": r9,
    }

# ==========================================
# BUILD DETAIL SHEETS FIRST
# ==========================================
ws2 = wb.create_sheet(title="ChiTiet_CH_2x2")
rows_2x2 = build_detail_sheet(ws2, "Cống hộp đơn 2.0x2.0m", 2.0, 2.0, 0.25, 0.25, 0.25, 0.15, 0.15, 0.10, 0.10)

ws3 = wb.create_sheet(title="ChiTiet_CH_3x3")
rows_3x3 = build_detail_sheet(ws3, "Cống hộp đơn 3.0x3.0m", 3.0, 3.0, 0.30, 0.30, 0.30, 0.20, 0.20, 0.10, 0.10)

ws4 = wb.create_sheet(title="ChiTiet_CH_Doi_2x3x3")
rows_2x3x3 = build_detail_sheet(ws4, "Cống hộp đôi 2x(3.0x3.0m)", 3.0, 3.0, 0.30, 0.30, 0.30, 0.20, 0.20, 0.10, 0.10, is_double=True, t_giua=0.30)

# ==========================================
# BUILD SHEET 1: TỔNG HỢP KHỐI LƯỢNG 1M DÀI
# ==========================================
ws1 = wb.create_sheet(title="01_TongHop_1m_Dai", index=0)
ws1.views.sheetView[0].showGridLines = True

ws1.cell(row=2, column=2, value="DỰ ÁN: TUYẾN CỐNG HỘP THOÁT NƯỚC A5 (T5-a*)").font = FONT_SUBTITLE
ws1.cell(row=3, column=2, value="BẢNG TỔNG HỢP BÓC TÁCH KHỐI LƯỢNG 1 ĐỐT (L = 1.0 MÉT DÀI)").font = FONT_TITLE
ws1.cell(row=4, column=2, value="Tiêu chuẩn kỹ thuật: TCVN 4453:1995, TCVN 1651:2018 | 100% Công thức liên kết động sang các Sheet chi tiết").font = FONT_NOTE

headers_ws1 = [
    "STT", "Hạng mục / Loại cống", "Kích thước thông thủy BxH (m)", "Đơn vị tính",
    "Bê tông lót M100 (m³)", "Bê tông thân nắp M250 (m³)", "Ván khuôn ngoài (m²)",
    "Ván khuôn trong (m²)", "Tổng ván khuôn (m²)", "Cốt thép cống (kg)", "Ghi chú kỹ thuật"
]

row = 6
for col_idx, h in enumerate(headers_ws1, start=2):
    cell = ws1.cell(row=row, column=col_idx, value=h)
    style_cell(cell, font=FONT_HEADER, fill=FILL_NAVY, alignment=ALIGN_CENTER)
ws1.row_dimensions[row].height = 30

data_summary = [
    (1, "Cống hộp đơn 2.0x2.0m", "2.0 x 2.0", "1m dài",
     f"='ChiTiet_CH_2x2'!E{rows_2x2['bt_lot_row']}",
     f"='ChiTiet_CH_2x2'!E{rows_2x2['bt_cong_row']}",
     f"='ChiTiet_CH_2x2'!E{rows_2x2['vk_ngoai_row']}",
     f"='ChiTiet_CH_2x2'!E{rows_2x2['vk_trong_row']}",
     f"='ChiTiet_CH_2x2'!E{rows_2x2['vk_tong_row']}",
     f"='ChiTiet_CH_2x2'!E{rows_2x2['thep_row']}",
     "Đốt cống hộp đơn điển hình"),
    (2, "Cống hộp đơn 3.0x3.0m", "3.0 x 3.0", "1m dài",
     f"='ChiTiet_CH_3x3'!E{rows_3x3['bt_lot_row']}",
     f"='ChiTiet_CH_3x3'!E{rows_3x3['bt_cong_row']}",
     f"='ChiTiet_CH_3x3'!E{rows_3x3['vk_ngoai_row']}",
     f"='ChiTiet_CH_3x3'!E{rows_3x3['vk_trong_row']}",
     f"='ChiTiet_CH_3x3'!E{rows_3x3['vk_tong_row']}",
     f"='ChiTiet_CH_3x3'!E{rows_3x3['thep_row']}",
     "Đốt cống hộp đơn điển hình"),
    (3, "Cống hộp đôi 2x(3.0x3.0m)", "2x(3.0 x 3.0)", "1m dài",
     f"='ChiTiet_CH_Doi_2x3x3'!E{rows_2x3x3['bt_lot_row']}",
     f"='ChiTiet_CH_Doi_2x3x3'!E{rows_2x3x3['bt_cong_row']}",
     f"='ChiTiet_CH_Doi_2x3x3'!E{rows_2x3x3['vk_ngoai_row']}",
     f"='ChiTiet_CH_Doi_2x3x3'!E{rows_2x3x3['vk_trong_row']}",
     f"='ChiTiet_CH_Doi_2x3x3'!E{rows_2x3x3['vk_tong_row']}",
     f"='ChiTiet_CH_Doi_2x3x3'!E{rows_2x3x3['thep_row']}",
     "Cống đôi 2 ngăn thoát nước chính"),
]

for item in data_summary:
    row += 1
    ws1.row_dimensions[row].height = 24
    for c_idx, val in enumerate(item, start=2):
        cell = ws1.cell(row=row, column=c_idx, value=val)
        if c_idx in [2, 5]:
            style_cell(cell, font=FONT_REGULAR, alignment=ALIGN_CENTER)
        elif c_idx in [3, 4, 12]:
            style_cell(cell, font=FONT_BOLD if c_idx == 3 else FONT_REGULAR, alignment=ALIGN_LEFT)
        else:
            style_cell(cell, font=FONT_BOLD, alignment=ALIGN_RIGHT,
                       num_format="#,##0.0" if c_idx == 11 else "#,##0.000",
                       fill=FILL_TOTAL if c_idx in [6, 7, 10] else (FILL_HIGHLIGHT if c_idx == 11 else None))

row += 2
ws1.cell(row=row, column=2, value="(*) Ghi chú & Hướng dẫn sử dụng:").font = FONT_BOLD
row += 1
ws1.cell(row=row, column=2, value="1. Bảng số liệu trên được tính cho đốt cống chuẩn hóa L = 1.0 mét dài. Khi thi công phân đốt thực tế (2m, 3m hoặc 4m), chỉ cần nhân theo chiều dài đốt.").font = FONT_NOTE
row += 1
ws1.cell(row=row, column=2, value="2. Toàn bộ các ô khối lượng ở Sheet Tổng hợp đều liên kết bằng công thức sống ('ChiTiet_CH_2x2'!E..., v.v.). Khi bạn thay đổi thông số kích thước ở các Sheet chi tiết, toàn bộ kết quả tự động cập nhật.").font = FONT_NOTE
row += 1
ws1.cell(row=row, column=2, value="3. Cốt thép tạm tính theo suất tiêu hao kinh nghiệm 100 kg/m³ bê tông cống. Có thể điều chỉnh hệ số này tại ô DM_thep ở từng Sheet chi tiết.").font = FONT_NOTE

# ==========================================
# AUTO-FIT COLUMN WIDTHS ACROSS ALL SHEETS
# ==========================================
for ws in wb.worksheets:
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if cell.number_format and ('#,##' in cell.number_format or '0.00' in cell.number_format):
                val_str += '      '
            if len(val_str) > max_len and '\n' not in val_str:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
    # Special adjustment
    ws.column_dimensions['A'].width = 3
    ws.column_dimensions['B'].width = 8
    ws.column_dimensions['C'].width = 34
    if ws.title == "01_TongHop_1m_Dai":
        ws.column_dimensions['C'].width = 30
        ws.column_dimensions['L'].width = 32

wb.save(output_file)
print(f"[OK] Đã xuất bản thành công bảng tính bóc tách hoàn hảo: {output_file}")
