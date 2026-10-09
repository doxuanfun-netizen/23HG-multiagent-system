# -*- coding: utf-8 -*-
"""
KCS REGISTER — sổ nghiệm thu KCS: đọc danh mục bản ghi, kiểm tra logic chéo ngày và sinh biên bản Word.

Thay cho bản cũ (examples/generate_kcs_word_package.py cũ), nơi cột "Đánh giá Logic" là chữ gõ tay
"Hợp lệ" ở cả 22 dòng, không hề được tính từ ngày:
  - Trạng thái từng bản ghi được TÍNH từ ngày bắt đầu / ngày nghiệm thu và cột `tham_chieu`
    (các bản ghi phải nghiệm thu trước khi bản ghi sau bắt đầu).
  - Bản ghi không có tham chiếu → "CHƯA KIỂM TỰ ĐỘNG". Không bao giờ ghi "Hợp lệ" ngầm định.
  - Biên bản chỉ kết luận "đủ điều kiện" khi trạng thái là HỢP LỆ.

Định dạng CSV (dấu chấm phẩy, UTF-8):
    ma;cong_viec;bat_dau;nghiem_thu;dieu_kien;tham_chieu
    BBNT-05;Định vị tim cọc ...;25/10/2026;28/10/2026;BBNT-02, BBNT-04;BBNT-02,BBNT-04
Ngày theo dd/mm/yyyy. `tham_chieu` là các mã cách nhau bằng dấu phẩy; để trống nếu không có bản ghi tiên quyết.
"""

from __future__ import annotations

import csv
import datetime as dt
import json
import os
import re
from dataclasses import dataclass
from typing import Dict, List, Optional, Sequence, Tuple

HOP_LE = "HỢP LỆ"
MAU_THUAN = "MÂU THUẪN"
CHUA_KIEM = "CHƯA KIỂM TỰ ĐỘNG"
NGAY_SAI = "NGÀY KHÔNG HỢP LỆ"
THAM_CHIEU_SAI = "THAM CHIẾU KHÔNG TỒN TẠI"
PASS_STATUSES = frozenset({HOP_LE})

_DATE = re.compile(r"^\s*(\d{1,2})/(\d{1,2})/(\d{4})\s*$")


def parse_date(text: str) -> dt.date:
    m = _DATE.match(text or "")
    if not m:
        raise ValueError(f"ngày '{text}' không đúng dạng dd/mm/yyyy")
    return dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))


def fmt_date(d: dt.date) -> str:
    return d.strftime("%d/%m/%Y")


@dataclass
class KcsRecord:
    code: str
    task: str
    start: dt.date
    finish: dt.date
    condition: str = ""
    refs: Tuple[str, ...] = ()


@dataclass
class Check:
    code: str
    status: str
    detail: str = ""

    @property
    def passed(self) -> bool:
        return self.status in PASS_STATUSES


class RegisterError(ValueError):
    """Sổ KCS không đọc được (thiếu cột, ngày sai, mã trùng)."""


def load_register(path: str) -> List[KcsRecord]:
    with open(path, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f, delimiter=";"))
    if not rows:
        raise RegisterError(f"{path}: không có bản ghi nào")
    need = {"ma", "cong_viec", "bat_dau", "nghiem_thu"}
    missing = need - set(rows[0].keys())
    if missing:
        raise RegisterError(f"{path}: thiếu cột {sorted(missing)}")
    out, seen = [], set()
    for i, row in enumerate(rows, start=2):
        code = (row.get("ma") or "").strip()
        if not code:
            raise RegisterError(f"{path} dòng {i}: thiếu mã bản ghi")
        if code in seen:
            raise RegisterError(f"{path} dòng {i}: mã {code} bị trùng")
        seen.add(code)
        try:
            start = parse_date(row.get("bat_dau", ""))
            finish = parse_date(row.get("nghiem_thu", ""))
        except ValueError as e:
            raise RegisterError(f"{path} dòng {i} ({code}): {e}") from e
        refs = tuple(x.strip() for x in (row.get("tham_chieu") or "").split(",") if x.strip())
        out.append(KcsRecord(code, (row.get("cong_viec") or "").strip(), start, finish,
                             (row.get("dieu_kien") or "").strip(), refs))
    return out


def check_register(records: Sequence[KcsRecord]) -> Dict[str, Check]:
    """Tính trạng thái logic cho từng bản ghi. Trả {mã: Check}."""
    by = {r.code: r for r in records}
    result: Dict[str, Check] = {}
    for r in records:
        if r.finish < r.start:
            result[r.code] = Check(r.code, NGAY_SAI,
                                   f"nghiệm thu {fmt_date(r.finish)} trước khi bắt đầu {fmt_date(r.start)}")
            continue
        if not r.refs:
            result[r.code] = Check(r.code, CHUA_KIEM, "không có bản ghi tiên quyết được khai báo")
            continue
        unknown = [x for x in r.refs if x not in by]
        if unknown:
            result[r.code] = Check(r.code, THAM_CHIEU_SAI, f"không có bản ghi {', '.join(unknown)}")
            continue
        late = [x for x in r.refs if r.start < by[x].finish]
        if late:
            parts = [f"{x} nghiệm thu {fmt_date(by[x].finish)}" for x in late]
            result[r.code] = Check(r.code, MAU_THUAN,
                                   f"bắt đầu {fmt_date(r.start)} trước khi " + "; ".join(parts))
        else:
            result[r.code] = Check(r.code, HOP_LE, "đúng thứ tự với " + ", ".join(r.refs))
    return result


def load_signers(path: Optional[str]) -> Dict[str, str]:
    """Người ký biên bản, đọc từ file cấu hình (không gõ trong code)."""
    keys = ("tvgs_ten", "tvgs_chuc_vu", "nt_ten", "nt_chuc_vu", "kcs_ten")
    if not path:
        return {k: "[.........]" for k in keys}
    with open(path, encoding="utf-8") as f:
        data = json.load(f)
    return {k: str(data.get(k, "[.........]")) for k in keys}


def _pass_text(check: Check) -> str:
    if check.status == HOP_LE:
        return "Đã kiểm tra logic chéo: đúng thứ tự với bước trước (" + check.detail.replace("đúng thứ tự với ", "") + ")."
    if check.status == CHUA_KIEM:
        return "Chưa kiểm tự động điều kiện tiên quyết; cần đối chiếu thủ công trước khi nghiệm thu."
    return f"Chưa đạt logic chéo ({check.status.lower()}: {check.detail}); cần xem lại ngày thực hiện trước khi nghiệm thu."


def build_docx(records: Sequence[KcsRecord], checks: Dict[str, Check], out_path: str, *,
               project_name: str, location: str, signers: Dict[str, str]) -> str:
    """Xuất sổ KCS: bảng kiểm tra logic chéo + một biên bản cho mỗi bản ghi. Trả đường dẫn file đã ghi."""
    import docx
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    def shade(cell, hex_fill):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_fill)
        tcPr.append(shd)

    def run(p, text, size=10, bold=False, italic=False, color=None):
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        return r

    doc = docx.Document()
    for s in doc.sections:
        s.top_margin, s.bottom_margin = Inches(0.75), Inches(0.75)
        s.left_margin, s.right_margin = Inches(0.85), Inches(0.75)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, "BẢNG KIỂM TRA LOGIC CHÉO NGÀY THÁNG NGHIỆM THU KCS", 13, bold=True, color="1F497D")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, f"CÔNG TRÌNH: {project_name}\n(Trạng thái được tính từ ngày và tham chiếu; không nhập tay)",
        10, italic=True, color="595959")

    n_pass = sum(1 for r in records if checks[r.code].passed)
    p = doc.add_paragraph()
    run(p, f"Kết quả: {n_pass}/{len(records)} bản ghi HỢP LỆ; "
           f"{len(records) - n_pass} bản ghi cần xem lại hoặc kiểm thủ công.", 10, bold=True)

    headers = ["Số hiệu", "Tên công việc nghiệm thu", "Ngày bắt đầu", "Ngày nghiệm thu",
               "Điều kiện tiên quyết", "Đánh giá logic"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        shade(cell, "1F497D")
        run(cell.paragraphs[0], h, 8.5, bold=True, color="FFFFFF")
    for r in records:
        c = checks[r.code]
        cells = table.add_row().cells
        values = [r.code, r.task, fmt_date(r.start), fmt_date(r.finish), r.condition or "—",
                  f"{c.status}" + (f" — {c.detail}" if c.detail else "")]
        for i, v in enumerate(values):
            cells[i].text = ""
            color = "008000" if c.passed and i == 5 else ("C00000" if (not c.passed and i == 5) else None)
            run(cells[i].paragraphs[0], v, 8, bold=(i == 5), color=color)

    doc.add_page_break()

    for idx, r in enumerate(records, start=1):
        c = checks[r.code]
        hdr = doc.add_table(rows=1, cols=2)
        hdr.alignment = WD_TABLE_ALIGNMENT.CENTER
        left, right = hdr.rows[0].cells
        left.text = ""
        run(left.paragraphs[0], "BAN QLDA ĐTXD CÔNG TRÌNH\n", 9, bold=True)
        run(left.paragraphs[0], f"CÔNG TRÌNH: {project_name}", 8.5, italic=True)
        right.text = ""
        run(right.paragraphs[0], "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\n", 9.5, bold=True)
        run(right.paragraphs[0], "Độc lập - Tự do - Hạnh phúc", 9, bold=True)

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(12)
        run(p, f"BIÊN BẢN NGHIỆM THU CÔNG VIỆC XÂY DỰNG\nSố: {r.code}/NT-GXD", 12, bold=True, color="1F497D")

        body = doc.add_paragraph()
        body.paragraph_format.line_spacing = 1.15
        conclusion = ("Đồng ý nghiệm thu công việc xây dựng nêu trên và cho phép nhà thầu tiếp tục thi công."
                      if c.passed else
                      "CHƯA đủ điều kiện nghiệm thu chuyển bước. Cần đối chiếu lại ngày thực hiện và điều kiện "
                      "tiên quyết trước khi ký.")
        spec = [
            ("1. Hạng mục công việc nghiệm thu: ", True), (f"{r.task}\n", False),
            ("2. Vị trí xây dựng: ", True), (f"{location}\n", False),
            ("3. Thời gian nghiệm thu: ", True),
            (f"Bắt đầu: {fmt_date(r.start)}; Kết thúc: {fmt_date(r.finish)}\n", False),
            ("4. Thành phần tham gia nghiệm thu:\n", True),
            ("   a) Đại diện Tư vấn giám sát: ", True),
            (f"{signers['tvgs_ten']} — {signers['tvgs_chuc_vu']}\n", False),
            ("   b) Đại diện Nhà thầu thi công: ", True),
            (f"{signers['nt_ten']} — {signers['nt_chuc_vu']}\n", False),
            ("5. Đánh giá công việc xây dựng đã thực hiện:\n", True),
            (f"   - {_pass_text(c)}\n", False),
            ("6. Kết luận nghiệm thu:\n", True), (f"   {conclusion}\n", False),
        ]
        for text, bold in spec:
            run(body, text, 9.5, bold=bold)

        sign = doc.add_table(rows=1, cols=2)
        sign.alignment = WD_TABLE_ALIGNMENT.CENTER
        a, b = sign.rows[0].cells
        for cell, title, name in ((a, "ĐẠI DIỆN NHÀ THẦU THI CÔNG", signers["nt_ten"]),
                                  (b, "ĐẠI DIỆN TƯ VẤN GIÁM SÁT", signers["tvgs_ten"])):
            cell.text = ""
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            run(cell.paragraphs[0], f"{title}\n\n\n\n\n{name}", 9.5, bold=True)

        if idx < len(records):
            doc.add_page_break()

    out_dir = os.path.dirname(os.path.abspath(out_path))
    os.makedirs(out_dir, exist_ok=True)
    doc.save(out_path)
    return out_path
