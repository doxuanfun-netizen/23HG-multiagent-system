# -*- coding: utf-8 -*-
"""
Sửa file theo dõi tần suất thí nghiệm (bê tông, cốt thép, tổng hợp, danh mục lấy mẫu DM) theo bản vẽ thẩm định.

- Bê tông: khối lượng từng đợt đổ lấy từ bản vẽ (du_lieu.json theo trang + nhãn, qua tools.estimate_reconcile);
  hạng mục bản vẽ chỉ ghi tổng thì chia cho các đợt theo tỷ lệ đợt cũ. Mỗi thay đổi ghi nguồn và NGHI VẤN.
- Cốt thép: dựng lại từ BBS đã sửa theo bản vẽ (tools.bbs_reconcile.read_bbs), cọc tách từng cọc.
- Tần suất: lần 1 lấy ngay khi bắt đầu dùng, sau đó mỗi khi lũy kế vượt bội số (50 t xi măng, 200 m³ cát,
  350 m³ đá, 20 t thép mỗi Ø) — số lần = ROUNDUP(lũy kế / tần suất). File cũ dùng INT nên Ø < 20 t có 0 lần,
  và cột cát/đá chia ngược 350/200.
- Danh mục lấy mẫu DM sinh lại theo đúng các lần lấy mẫu mới (giữ mẫu công thức của dòng DM cũ).
"""

from __future__ import annotations

import copy
import datetime as _dt
import json
import math
import os
import re
from dataclasses import dataclass
from typing import Dict, List, Optional, Tuple

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from tools.estimate_reconcile import NghiVan, resolve_f, write_flag_sheet

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_CONFIG = os.path.join(ROOT, "examples", "data", "km19_tan_suat_theo_ban_ve.json")

FILL_SUA = PatternFill("solid", fgColor="FFF2CC")
FILL_NGHI_VAN = PatternFill("solid", fgColor="F8CBAD")

STEEL_COLS = {10: "P", 12: "Q", 14: "R", 16: "S", 18: "T", 20: "U", 22: "V", 25: "W", 28: "X", 32: "Y", "cap": "Z"}
TRIG_COLS = {10: "AA", 12: "AB", 14: "AC", 16: "AD", 18: "AE", 20: "AF", 22: "AG", 25: "AH", 28: "AI", 32: "AJ",
             "cap": "AK"}
TONG_HOP_ROW = {10: 12, 12: 13, 14: 14, 16: 15, 18: 16, 20: 17, 22: 18, 25: 19, 28: 20, 32: 21, "cap": 22}
CAP = "ASTM A416"


def _num(v) -> Optional[float]:
    """Giá trị số của ô: số, hoặc công thức số học đơn giản ('=184.24/2', '=52.96')."""
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str) and re.fullmatch(r"=\s*[\d.+\-*/() ]+", v):
        return float(eval(v[1:], {"__builtins__": {}}))  # chỉ số và + - * / ( )
    return None


def _value(spec, pages) -> float:
    if isinstance(spec, (int, float)):
        return float(spec)
    spec = dict(spec)
    chia = spec.pop("chia", 1)
    return resolve_f(spec, pages)[0] / chia


def _fmt(x: Optional[float]) -> str:
    return "" if x is None else f"{x:,.2f}".replace(",", " ").replace(".", ",").replace(" ", ".")


def _split(total: float, weights: List[float]) -> List[float]:
    """Chia tổng theo tỷ lệ, làm tròn 0,01; dòng cuối nhận phần dư để tổng đúng bằng tổng bản vẽ."""
    w = sum(weights)
    parts = [round(total * x / w, 2) if w else round(total / len(weights), 2) for x in weights]
    parts[-1] = round(total - sum(parts[:-1]), 2)
    return parts


def _lan(x: float, step: float) -> int:
    """Số lần thí nghiệm khi lũy kế = x: lần 1 ngay khi bắt đầu dùng, thêm 1 lần mỗi khi vượt bội số của step."""
    return 0 if x <= 1e-9 else math.ceil(round(x / step, 9))


# ----------------------------------------------------------------------------------------------- bê tông
def data_rows(ws) -> List[int]:
    return [r for r in range(1, ws.max_row + 1) if isinstance(ws.cell(r, 1).value, (int, float))
            and isinstance(ws.cell(r, 2).value, str)]


def find_row(ws, rows: List[int], pattern: str) -> Optional[int]:
    return next((r for r in rows if re.search(pattern, ws.cell(r, 2).value, re.I)), None)


def concrete_offset(ws) -> int:
    """
    Bố cục sheet bê tông: file của đơn vị có cột E = khối lượng; file sinh từ repo không có cột D phụ nên lệch trái
    1 cột (D = khối lượng). Chỉ số cột trong code là theo bố cục E; cột >= 4 cộng thêm độ lệch này.
    """
    col = next((c for c in range(1, 10) if str(ws.cell(1, c).value or "").strip().upper().startswith("KL")), 5)
    return col - 5


def apply_concrete(ws, cfg: dict, pages, flags: List[NghiVan]) -> List[int]:
    off = concrete_offset(ws)
    C = lambda n: n + off if n >= 4 else n                      # noqa: E731
    rows = data_rows(ws)
    sheet = ws.title
    note_col = ws.max_column + 1
    ws.cell(1, note_col, "NGUỒN BẢN VẼ / NGHI VẤN (soát & ghim)").font = Font(bold=True)
    ws.column_dimensions[get_column_letter(note_col)].width = 70
    ts = cfg["tan_suat"]

    def note(r, text, flag=True, old=None, new=None):
        c = ws.cell(r, note_col)
        c.value = (c.value + " | " if c.value else "") + text
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if flag:
            c.fill = FILL_NGHI_VAN
            flags.append(NghiVan(sheet, r, ws.cell(r, 2).value, text, _fmt(old), _fmt(new)))

    matched = set()
    has = lambda pattern: any(re.search(pattern, ws.cell(r, 2).value, re.I) for r in rows)   # noqa: E731
    for rule in cfg["be_tong"]:
        if rule.get("chi_khi_co") and not has(rule["chi_khi_co"]):
            continue
        if rule.get("chi_khi_khong_co") and has(rule["chi_khi_khong_co"]):
            continue
        hits = [r for r in rows if re.search(rule["mau"], ws.cell(r, 2).value, re.I) and r not in matched]
        if not hits:
            if not rule.get("khong_bat_buoc"):
                flags.append(NghiVan(sheet, None, rule["mau"], "Không tìm thấy dòng đợt đổ ứng với hạng mục bản vẽ này"))
            continue
        matched.update(hits)
        olds = [_num(ws.cell(r, C(5)).value) for r in hits]
        if rule.get("chia_theo_ty_le_cu"):
            news = _split(_value(rule["E"], {} if isinstance(rule["E"], (int, float)) else pages),
                          [o or 0.0 for o in olds])
        else:
            v = round(_value(rule["E"], pages), 2)
            news = [v] * len(hits)
        for r, old, new in zip(hits, olds, news):
            ws.cell(r, C(5)).value = new
            ws.cell(r, C(5)).fill = FILL_SUA
            if rule.get("B"):
                ws.cell(r, 2).value = rule["B"]
            if rule.get("doi_mac"):
                ws.cell(r, 2).value = ws.cell(r, 2).value.replace(*rule["doi_mac"])
            changed = old is None or abs(new - old) > 0.005
            text = (f"Bản vẽ: {_fmt(old)} → {_fmt(new)} m³" if changed else "Khớp bản vẽ") + f" ({rule['nguon']})"
            if rule.get("nghi_van"):
                text += ". " + rule["nghi_van"]
            note(r, text, flag=changed or bool(rule.get("nghi_van")), old=old, new=new)
    for r in rows:
        if r not in matched:
            note(r, "Chưa đối chiếu bản vẽ (không có quy tắc cho dòng này)", old=_num(ws.cell(r, C(5)).value))

    # cấp phối: chỉ có cấp phối cho các mác trong 'cap_phoi_co_san'; mác khác đang dùng cấp phối C30
    for r in rows:
        grades = set(re.findall(r"\b(C\d{2}|M\d{3})\b", ws.cell(r, 2).value))
        missing = grades - set(cfg.get("cap_phoi_co_san", []))
        if missing and not ws.cell(r, C(8)).value:
            note(r, f"Mác {', '.join(sorted(missing))} chưa có cấp phối riêng — đang dùng cấp phối C30 "
                    f"(0,425 t XM; 0,555 m³ cát; 0,700 m³ đá / m³). Nhập cấp phối được duyệt.")

    # tần suất vật liệu: XM / cát / đá — ROUNDUP, đúng số chia theo tiêu đề cột
    steps = {17: ts["xi_mang_tan"], 18: ts["cat_m3"], 19: ts["da_m3"]}
    src = {17: 14, 18: 15, 19: 16}
    prev = None
    for r in rows:
        for col, step in steps.items():
            c = get_column_letter(C(src[col]))
            cur, before = f"{c}{r}", (f"{c}{prev}" if prev else None)
            n = f"ROUNDUP({cur}/{step},0)"
            cond = f"{n}>ROUNDUP({before}/{step},0)" if before else f"{n}>0"
            ws.cell(r, C(col)).value = f'=IF({cur}="","",IF({cond},"Lần " & {n},""))'
        prev = r
    flags.append(NghiVan(sheet, None, "Cột Q/R/S", f"Sửa công thức số lần: cát chia {ts['cat_m3']} m³, đá chia "
                                                   f"{ts['da_m3']} m³ (dòng 4–73 cũ chia ngược 350/200); lần 1 lấy khi "
                                                   f"bắt đầu dùng (ROUNDUP thay INT)."))
    for item in cfg.get("thieu_be_tong", []):
        flags.append(NghiVan(sheet, None, item["hang_muc"],
                             f"Có trên bản vẽ nhưng file chưa có dòng đợt đổ: {item['khoi_luong']}", "",
                             item["khoi_luong"], item["trang"]))
    return rows


def concrete_events(ws, rows: List[int], cfg: dict) -> List[Tuple[int, int, int]]:
    """(dòng, cột Q/R/S, lần) — tính lại đúng như công thức trong file."""
    ts = cfg["tan_suat"]
    off = concrete_offset(ws)
    cum = {14: 0.0, 15: 0.0, 16: 0.0}
    ratio = {14: (8, 11), 15: (9, 12), 16: (10, 13)}       # cọc nhồi (H/I/J) hoặc thường (K/L/M), bố cục E
    steps = {14: ts["xi_mang_tan"], 15: ts["cat_m3"], 16: ts["da_m3"]}
    out = []
    for r in rows:
        e = _num(ws.cell(r, 5 + off).value) or 0.0
        for c in (14, 15, 16):
            a, b = ratio[c]
            k = _num(ws.cell(r, a + off).value) or 0.0
            k = k if k > 0 else (_num(ws.cell(r, b + off).value) or 0.0)
            before = cum[c]
            cum[c] = round(before + e * k, 2)
            for n in range(_lan(before, steps[c]) + 1, _lan(cum[c], steps[c]) + 1):   # mỗi lần một sự kiện
                out.append((r, c + 3, n))
    return out


# ----------------------------------------------------------------------------------------------- cốt thép
@dataclass
class SteelRow:
    hang_muc: str
    bo_phan: str
    mark: str
    d: float
    mac: str
    hinh: str
    L: float
    n: float
    unit: float
    loai: str
    bt_row: Optional[int]

    @property
    def key(self):
        return "cap" if self.mac == CAP else (10 if self.d <= 10 else int(self.d))


def steel_rows_from_bbs(sections, cfg: dict, ws_bt, bt_rows: List[int], flags: List[NghiVan],
                        bt_dates: Optional[Dict[int, object]] = None) -> List[Tuple[str, List[SteelRow]]]:
    """
    [(tiêu đề nhóm, [dòng thép])]: cọc tách từng cọc, các hạng mục theo BBS. Nếu có ngày đổ bê tông của các dòng
    liên kết (bt_dates) thì xếp nhóm theo ngày — lũy kế từng Ø và các lần lấy mẫu đi theo thời gian thi công;
    nhóm chưa có ngày xếp sau, giữ thứ tự BBS.
    """
    groups: List[Tuple[str, List[SteelRow]]] = []
    link = cfg["thep_lien_ket"]

    def bt(pattern):
        r = find_row(ws_bt, bt_rows, pattern) if pattern else None
        if pattern and r is None:
            flags.append(NghiVan(ws_bt.title, None, pattern, "Không tìm thấy dòng bê tông để lấy ngày lắp dựng thép"))
        return r

    def mk(row, hang_muc, n, loai, bt_row):
        cable = row.is_cable
        return SteelRow(hang_muc, row.bo_phan, row.mark, 15.2 if cable else float(row.d), CAP if cable else row.mac,
                        row.hinh or "", float(row.L_m), float(n), float(row.unit), loai, bt_row)

    for sec in sections:
        for name, rows in sec.items.items():
            rows = [r for r in rows if not r.placeholder]
            if name.startswith("Cọc khoan nhồi"):
                for c in cfg["coc"]:
                    mine = [r for r in rows if c["bo_phan"] in r.bo_phan]
                    for k in range(1, c["so_coc"] + 1):
                        title = f"Cọc khoan nhồi C{k} {c['ten']} (D1.2m, L={c['L']}m)"
                        r_bt = bt(c["mau_dong"].format(k=k))
                        groups.append((title.upper(), [mk(r, title, r.n_per, "Coc", r_bt) for r in mine]))
                continue
            rule = next((x for x in link if x["hang_muc"] in name), None)
            if rule is None:
                flags.append(NghiVan("Theo_Doi_Tan_Suat_Cot_Thep", None, name, "Hạng mục BBS chưa gắn dòng bê tông"))
            out = []
            for r in rows:
                pattern = None
                if rule and "theo_ky_hieu" in rule:
                    pattern = next((p for k, p in rule["theo_ky_hieu"].items() if k and re.search(k, r.mark or "")),
                                   rule["theo_ky_hieu"].get(""))
                elif rule and "theo_bo_phan" in rule:
                    pattern = next((p for k, p in rule["theo_bo_phan"].items() if k and k in (r.bo_phan or "")),
                                   rule["theo_bo_phan"].get(""))
                elif rule:
                    pattern = rule["dong"]
                out.append(mk(r, name, r.n_per * r.n_ck, rule["loai"] if rule else "", bt(pattern)))
            groups.append((name.upper(), out))
    if bt_dates:
        far = _dt.datetime.max

        def when(g):
            ds = [bt_dates.get(x.bt_row) for x in g[1] if isinstance(bt_dates.get(x.bt_row), _dt.datetime)]
            return min(ds) if ds else far
        groups = [g for _, g in sorted(enumerate(groups), key=lambda t: (when(t[1]), t[0]))]
    return groups


def _copy_style(src, dst):
    if src.has_style:
        dst._style = copy.copy(src._style)


def apply_steel(ws, groups: List[Tuple[str, List[SteelRow]]], cfg: dict, flags: List[NghiVan]):
    """Ghi lại sheet thép (giữ dòng tiêu đề 1–2, bảng ngày lùi AQ:AR). Trả (dòng tổng, {dòng: SteelRow}, sự kiện)."""
    step = cfg["tan_suat"]["thep_tan"]
    bt_sheet = cfg["sheet_be_tong"]
    tmpl_data = {c: ws.cell(5, c) for c in range(1, 45)}
    tmpl_sec = {c: ws.cell(4, c) for c in range(1, 45)}
    styles_data = {c: copy.copy(x._style) for c, x in tmpl_data.items() if x.has_style}
    styles_sec = {c: copy.copy(x._style) for c, x in tmpl_sec.items() if x.has_style}
    fmt_data = {c: x.number_format for c, x in tmpl_data.items()}
    offsets = [(ws.cell(r, 43).value, ws.cell(r, 44).value) for r in range(2, 8)]
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row >= 3:
            ws.unmerge_cells(str(rng))
    ws.delete_rows(3, ws.max_row)
    if not any(a for a, _ in offsets):
        # file chưa có bảng ngày lùi (lắp dựng thép trước ngày đổ bê tông, ngày) — dùng bảng của file đơn vị
        offsets = [("Coc", 1), ("Be", 3), ("Than", 2), ("Dam", 4), ("HoanThien", 2), (None, None)]
        ws.cell(1, 40, "LOẠI HẠNG MỤC (OFFSET)")
        ws.cell(1, 41, "DÒNG BÊ TÔNG ĐỐI CHIẾU")
        ws.cell(1, 43, "LOẠI HẠNG MỤC")
        ws.cell(1, 44, "LÙI (ngày)")
        ws.cell(1, 13, "N/T/N\nLẮP DỰNG")
        flags.append(NghiVan(ws.title, None, "Ngày lắp dựng", "Thêm bảng ngày lắp dựng thép trước ngày đổ bê tông (cọc 1, "
                                                              "bệ 3, thân 2, dầm 4, hoàn thiện 2 ngày) — chỉnh theo tiến độ."))
    for i, (a, b) in enumerate(offsets):
        ws.cell(2 + i, 43, a)
        ws.cell(2 + i, 44, b)
    ws["P2"] = "Ø10 (kèm Ø6, Ø8)\n(CB240-T)"

    r, stt, prev = 3, 0, None
    first_row = None
    rowmap: Dict[int, SteelRow] = {}
    cum = {k: 0.0 for k in STEEL_COLS}
    events = []
    for title, rows in groups:
        ws.cell(r, 1, f"--- {title} ---")
        for c, s in styles_sec.items():
            ws.cell(r, c)._style = copy.copy(s)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=38)
        r += 1
        for s in rows:
            stt += 1
            for c, st in styles_data.items():
                ws.cell(r, c)._style = copy.copy(st)
            vals = [stt, s.hang_muc, s.bo_phan, s.mark, s.d, s.mac, s.hinh, s.L, s.n, s.unit]
            for c, v in enumerate(vals, start=1):
                ws.cell(r, c, v)
            ws.cell(r, 11, f"=ROUND(H{r}*I{r}*J{r}, 2)")
            ws.cell(r, 12, f"=ROUND(K{r}/1000, 3)")
            ws.cell(r, 13, f'=IF(AO{r}="","",IF(INDEX({bt_sheet}!$C:$C,AO{r})="","",'
                           f'INDEX({bt_sheet}!$C:$C,AO{r})-VLOOKUP(AN{r},$AQ$2:$AR$7,2,FALSE)))')
            ws.cell(r, 13).number_format = fmt_data.get(13, "dd/mm/yyyy")
            for key, col in STEEL_COLS.items():
                cond = (f'F{r}="{CAP}"' if key == "cap" else
                        f'AND(E{r}<=10,F{r}<>"{CAP}")' if key == 10 else f'AND(E{r}={key},F{r}<>"{CAP}")')
                base = f"{col}{prev}+" if prev else ""
                ws[f"{col}{r}"] = f"=ROUND({base}IF({cond},L{r},0),3)"
                tcol = TRIG_COLS[key]
                n = f"ROUNDUP({col}{r}/{step},0)"
                cmp = f"{n}>ROUNDUP({col}{prev}/{step},0)" if prev else f"{n}>0"
                ws[f"{tcol}{r}"] = f'=IF({cmp},"Lần " & {n},"")'
            key = s.key
            col = STEEL_COLS[key]
            label = "Cáp 15.2" if key == "cap" else f"Ø{key}"
            n = f"ROUNDUP({col}{r}/{step},0)"
            cmp = f"{n}>ROUNDUP({col}{prev}/{step},0)" if prev else f"{n}>0"
            ws.cell(r, 15, f'=IF({cmp},"TN {label} (Lần " & {n} & ")","")')
            ws.cell(r, 38, None)                      # kết luận: ghi theo kết quả thí nghiệm, không ghi sẵn "ĐẠT"
            ws.cell(r, 40, s.loai or None)
            ws.cell(r, 41, s.bt_row)
            before = cum[key]
            cum[key] = round(before + round(round(s.L * s.n * s.unit, 2) / 1000, 3), 3)
            for n in range(_lan(before, step) + 1, _lan(cum[key], step) + 1):   # dòng lớn có thể vượt nhiều lô
                events.append((r, key, n))
            rowmap[r] = s
            first_row = first_row or r
            prev = r
            r += 1
    total = r
    ws.cell(total, 1, "TỔNG CỘNG KHỐI LƯỢNG CỐT THÉP & CÁP DƯL TOÀN BỘ CÔNG TRÌNH (theo BBS đã sửa theo bản vẽ)")
    for c, s in styles_sec.items():
        ws.cell(total, c)._style = copy.copy(s)
    ws.cell(total, 11, f"=SUM(K{first_row}:K{prev})")
    ws.cell(total, 12, f"=SUM(L{first_row}:L{prev})")
    for key, col in STEEL_COLS.items():
        ws[f"{col}{total}"] = f"={col}{prev}"
        ws[f"{TRIG_COLS[key]}{total}"] = f'="Tổng: " & ROUNDUP({col}{prev}/{step},0) & " lần"'
    flags.append(NghiVan(ws.title, None, "Toàn sheet", "Dựng lại từ BBS đã sửa theo bản vẽ thẩm định (thép cũ 793 t theo "
                                                        "BBS chưa sửa). Số lần TN = ROUNDUP(lũy kế/20 t) — Ø dưới 20 t "
                                                        "vẫn có 1 lần (file cũ dùng INT nên Ø10, Ø22 = 0 lần). Cột kết "
                                                        "luận để trống, ghi theo kết quả thí nghiệm (file cũ ghi sẵn "
                                                        "'ĐẠT TCVN 1651:2018')."))
    return total, rowmap, events


# ----------------------------------------------------------------------------------------------- TH_Thep_BQD_GLC
def apply_th_bqd(ws, steel_sheet: str, rowmap: Dict[int, SteelRow], flags: List[NghiVan]):
    """Sheet tổng hợp thép bản quá độ & gờ lan can trên mố: 3 nhóm, tham chiếu dòng mới của sheet thép."""
    tmpl = {c: copy.copy(ws.cell(6, c)._style) for c in range(1, 23) if ws.cell(6, c).has_style}
    tsec = {c: copy.copy(ws.cell(5, c)._style) for c in range(1, 23) if ws.cell(5, c).has_style}
    tsum = {c: copy.copy(ws.cell(20, c)._style) for c in range(1, 23) if ws.cell(20, c).has_style}
    tot = {c: copy.copy(ws.cell(54, c)._style) for c in range(1, 23) if ws.cell(54, c).has_style}
    extra = [[ws.cell(rr, c).value for c in range(1, 23)] for rr in (55, 56, 57)]
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row >= 5:
            ws.unmerge_cells(str(rng))
    ws.delete_rows(5, ws.max_row)
    groups = [("GỜ LAN CAN - MỐ M1", "Cộng Gờ lan can Mố M1",
               [r for r, s in rowmap.items() if s.hang_muc == "Mố M1" and (s.mark or "").startswith("G")]),
              ("GỜ LAN CAN - MỐ M2", "Cộng Gờ lan can Mố M2",
               [r for r, s in rowmap.items() if s.hang_muc == "Mố M2" and (s.mark or "").startswith("G")]),
              ("BẢN QUÁ ĐỘ (Mố M1 + Mố M2)", "Cộng Bản quá độ",
               [r for r, s in rowmap.items() if s.hang_muc.startswith("Bản quá độ")])]
    cols = {"O": 10, "P": 12, "Q": 14, "R": 16, "S": 18, "T": 20, "U": 22}
    r, sums, first, last = 5, [], None, None
    for title, sum_title, rows in groups:
        ws.cell(r, 1, title)
        for c, s in tsec.items():
            ws.cell(r, c)._style = copy.copy(s)
        r += 1
        start = r
        for i, sr in enumerate(rows):
            for c, s in tmpl.items():
                ws.cell(r, c)._style = copy.copy(s)
            ws.cell(r, 1, "=1" if i == 0 else f"=A{r - 1}+1")
            for c in range(2, 13):
                L = get_column_letter(c)
                ws[f"{L}{r}"] = f"={steel_sheet}!{L}{sr}"
            ws[f"N{r}"] = f"=I{r}*H{r}"
            for L, d in cols.items():
                ws[f"{L}{r}"] = f"=IF(E{r}<=10,N{r},0)" if d == 10 else f"=IF(E{r}={d},N{r},0)"
            first = first or r
            last = r
            r += 1
        for c, s in tsum.items():
            ws.cell(r, c)._style = copy.copy(s)
        ws.cell(r, 3, sum_title)
        ws[f"L{r}"] = f"=SUM(L{start}:L{r - 1})"
        sums.append(f"L{r}")
        r += 1
        if any(rowmap[x].d > 22 for x in rows):
            flags.append(NghiVan(ws.title, None, title, "Có thanh Ø > 22 — bảng chỉ có cột Ø10..Ø22"))
    total = r
    for c, s in tot.items():
        ws.cell(total, c)._style = copy.copy(s)
    ws.cell(total, 1, "TỔNG CỘNG KHỐI LƯỢNG THÉP (TẤN)")
    ws[f"L{total}"] = "=" + "+".join(sums)
    ws[f"N{total}"] = f"=SUBTOTAL(109,N{first}:N{last})"
    for L in cols:
        ws[f"{L}{total}"] = f"=SUM({L}{first}:{L}{last})"
    for i, vals in enumerate(extra, start=1):           # dòng số cây ÷ 11,7 và các phương án PA4/PA1
        for c, v in enumerate(vals, start=1):
            if isinstance(v, str) and v.startswith("="):  # dòng 54..57 cũ → total..total+3 (một lượt thay)
                v = re.sub(r"\b([A-Z]+)5([4-7])\b", lambda m: f"{m.group(1)}{total + int(m.group(2)) - 4}", v)
            ws.cell(total + i, c, v)
    flags.append(NghiVan(ws.title, None, "Phạm vi", "Sheet chỉ gồm gờ lan can TRÊN MỐ và bản quá độ. Gờ lan can chính của "
                                                    "cầu (tr.119, 20,07 t) không nằm ở sheet này — nếu dùng để đặt thép "
                                                    "gờ lan can thì cần thêm."))
    flags.append(NghiVan(ws.title, total + 2, "PA4 / PA1", "Số cây PA4/PA1 nhập tay theo khối lượng cũ — cập nhật theo "
                                                           "bộ cắt thép (02_XUONG_TIEN_CHE_COT_THEP)."))
    return total, first, last


def update_bang_nhap(wb, th_total: int, first: int, last: int):
    for ws in wb:
        if not ws.title.startswith("Bang_Nhap_Thep"):
            continue
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and "TH_Thep_BQD_GLC" in c.value:
                    v = re.sub(r"TH_Thep_BQD_GLC!([A-Z]+)54\b", lambda m: f"TH_Thep_BQD_GLC!{m.group(1)}{th_total}", c.value)
                    c.value = re.sub(r"TH_Thep_BQD_GLC!\$E\$\d+:\$J\$\d+", f"TH_Thep_BQD_GLC!$E${first}:$J${last}", v)


# ----------------------------------------------------------------------------------------------- tổng hợp
def apply_tong_hop(ws, steel_sheet: str, steel_total: int, cfg: dict, flags: List[NghiVan], cap_ban_ve: float,
                   cap_bbs: float):
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and steel_sheet in c.value:
                c.value = re.sub(rf"{steel_sheet}!([A-Z]+)\d+", lambda m: f"{steel_sheet}!{m.group(1)}{steel_total}", c.value)
            if isinstance(c.value, str) and re.search(r"INT\(F\d+/\d+\)\+1", c.value):
                c.value = re.sub(r"INT\((F\d+)/(\d+)\)\+1", r"ROUNDUP(\1/\2,0)", c.value)
    if abs(cap_ban_ve - cap_bbs) > 0.005:
        flags.append(NghiVan(ws.title, TONG_HOP_ROW["cap"], "Cáp DƯL 15.2", f"BBS {cap_bbs:.3f} t; bảng tổng hợp bản vẽ "
                                                                           f"tr.175 ghi {cap_ban_ve:.3f} t",
                             f"{cap_bbs:.3f}", f"{cap_ban_ve:.3f}", "tr.175"))
    for rr in range(23, ws.max_row + 1):
        b = str(ws.cell(rr, 2).value or "")
        if any(k in b for k in ("nén mẫu bê tông", "siêu âm", "PDA", "độ sụt", "cóc nối", "coupler")):
            flags.append(NghiVan(ws.title, rr, b, f"Khối lượng '{ws.cell(rr, 6).value}' / '{ws.cell(rr, 7).value}' "
                                                  f"đang nhập tay — chưa đối chiếu bản vẽ và chỉ dẫn kỹ thuật"
                                                  + (" (tr.175: cóc nối cọc 4.608 bộ)" if "cóc nối" in b else "")))


# ----------------------------------------------------------------------------------------------- DM
def apply_dm(ws, bt_sheet: str, steel_sheet: str, conc_events, steel_events, flags: List[NghiVan]):
    """Sinh lại danh mục lấy mẫu: mỗi lần lấy mẫu một dòng, giữ mẫu công thức/định dạng của dòng DM cũ."""
    def template(r):
        return [(ws.cell(r, c).value, copy.copy(ws.cell(r, c)._style) if ws.cell(r, c).has_style else None,
                 ws.cell(r, c).number_format) for c in range(1, ws.max_column + 1)]
    conc_t = template(14)
    steel_r = next((r for r in range(14, ws.max_row + 1) if steel_sheet in str(ws.cell(r, 7).value or "")), 14)
    steel_t = template(steel_r)
    old_n = sum(1 for r in range(14, ws.max_row + 1) if ws.cell(r, 2).value)
    ws.delete_rows(14, ws.max_row)
    mat_row = {17: 6, 18: 7, 19: 8}
    r = 14
    events = [("bt", row, col, n) for row, col, n in conc_events] + [("thep", row, key, n) for row, key, n in steel_events]
    for kind, src, col, n in events:
        t = conc_t if kind == "bt" else steel_t
        for c, (v, st, nf) in enumerate(t, start=1):
            cell = ws.cell(r, c)
            if st is not None:
                cell._style = copy.copy(st)
            cell.number_format = nf
            if isinstance(v, str) and v.startswith("="):
                v = re.sub(r"([A-Z])(14|%d)\b" % steel_r, lambda m: f"{m.group(1)}{r}", v) if c in (6, 8, 9, 14) else v
            cell.value = v
        ws.cell(r, 1, 1 if r == 14 else f"=+A{r - 1}+1")
        # số lần ghi thẳng (một dòng đợt đổ/dòng thép có thể ứng với nhiều lần lấy mẫu); tên lấy theo liên kết
        if kind == "bt":
            ws.cell(r, 3, f'="Lấy mẫu tần suất: "&Tong_Hop_Tan_Suat_Vat_Lieu!B{mat_row[col]}')
            ws.cell(r, 7, f"={bt_sheet}!C{src}")
            ws.cell(r, 11, f'="Lần {n} - "&{bt_sheet}!B{src}')
        else:
            ws.cell(r, 3, f'="Lấy mẫu tần suất: "&Tong_Hop_Tan_Suat_Vat_Lieu!B{TONG_HOP_ROW[col]}')
            ws.cell(r, 7, f"={steel_sheet}!M{src}")
            ws.cell(r, 11, f'="Lần {n} - "&{steel_sheet}!B{src}&" - "&{steel_sheet}!C{src}')
        ws.cell(r, 6, f'=IF(G{r}="","",G{r}-1)')
        r += 1
    flags.append(NghiVan(ws.title, None, "Danh mục lấy mẫu", f"Sinh lại theo các lần lấy mẫu mới: {len(events)} dòng "
                                                             f"(cũ {old_n}). Ngày lấy mẫu trống khi đợt đổ chưa có ngày "
                                                             f"(trước đây báo #VALUE!). Giờ bắt đầu/kết thúc lấy theo dòng "
                                                             f"mẫu cũ — chỉnh theo thực tế."))
    return len(events)


def apply_ban_ve(wb, config: str = DEFAULT_CONFIG, ngay: Optional[str] = None,
                 bt_dates: Optional[Dict[int, object]] = None) -> dict:
    """Sửa cả workbook tần suất theo bản vẽ: bê tông, thép từ BBS, TH_Thep_BQD_GLC/Bang_Nhap_Thep (nếu có), tổng hợp,
    DM (nếu có) và sheet NGHI_VAN. Dùng chung cho bộ sinh (tools/generate_tan_suat_thi_nghiem_excel.py) và script sửa
    file của đơn vị (examples/sua_tan_suat_km19_theo_ban_ve.py). bt_dates: {dòng bê tông: ngày đổ} để xếp thép theo
    thời gian; không có thì giữ thứ tự hạng mục."""
    from tools.bbs_reconcile import load_pages, read_bbs

    with open(config, encoding="utf-8") as f:
        cfg = json.load(f)
    pages = load_pages(os.path.join(ROOT, cfg["du_lieu_ban_ve"]))
    _, sections = read_bbs(os.path.join(ROOT, cfg["bbs"]))
    flags: List[NghiVan] = []
    ws_bt = wb[cfg["sheet_be_tong"]]
    bt_rows = apply_concrete(ws_bt, cfg, pages, flags)
    conc_ev = concrete_events(ws_bt, bt_rows, cfg)
    groups = steel_rows_from_bbs(sections, cfg, ws_bt, bt_rows, flags, bt_dates)
    steel_total, rowmap, steel_ev = apply_steel(wb[cfg["sheet_thep"]], groups, cfg, flags)

    if cfg.get("sheet_thep_bqd_glc") in wb.sheetnames:
        th_total, first, last = apply_th_bqd(wb[cfg["sheet_thep_bqd_glc"]], cfg["sheet_thep"], rowmap, flags)
        update_bang_nhap(wb, th_total, first, last)
    cap_bbs = round(sum(s.L * s.n * s.unit for s in rowmap.values() if s.key == "cap") / 1000, 3)
    cap_ban_ve = resolve_f(cfg["cap_dul_ban_ve_tan"], pages)[0]
    apply_tong_hop(wb[cfg["sheet_tong_hop"]], cfg["sheet_thep"], steel_total, cfg, flags, cap_ban_ve, cap_bbs)
    n_dm = apply_dm(wb["DM"], cfg["sheet_be_tong"], cfg["sheet_thep"], conc_ev, steel_ev, flags) if "DM" in wb.sheetnames else 0
    if "BBLM" in wb.sheetnames:
        flags.append(NghiVan("BBLM", 4, "Nút chọn phiếu", "Nút xoay (spinner) chọn số phiếu không giữ được khi công cụ lưu "
                                                          "file — nhập số thứ tự phiếu vào ô N4, hoặc chép sheet BBLM từ "
                                                          "file gốc sang."))
    ngay = ngay or _dt.date.today().strftime("%d/%m/%Y")
    write_flag_sheet(wb, flags, [], f"NGHI VẤN TẦN SUẤT THÍ NGHIỆM — sửa theo bản vẽ thẩm định ngày {ngay} (soát và "
                                    f"ghim ở cột I)")
    return {"nghi_van": len(flags), "dong_thep": len(rowmap), "dm": n_dm, "lan_bt": len(conc_ev), "lan_thep": len(steel_ev),
            "cap_bbs": cap_bbs, "cap_ban_ve": cap_ban_ve}
