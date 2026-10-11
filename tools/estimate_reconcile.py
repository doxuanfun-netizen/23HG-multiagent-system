# -*- coding: utf-8 -*-
"""
Sửa sheet diễn giải khối lượng (QS) của dự toán / phụ lục 03a theo bản vẽ thẩm định, đánh dấu NGHI VẤN.

Số liệu bản vẽ lấy tự động từ dữ liệu OCR (du_lieu.json) theo trang + nhãn dòng, không gõ tay:
  {"trang": 64, "nhan": "THE TICH BE TONG COC"}            → số cuối cùng trên dòng đầu tiên chứa nhãn (chấm thập phân)
  {"trang": 175, "sau": "BAN MAT CAU", "nhan": "..."}      → chỉ tìm sau dòng chứa nhãn 'sau'
  {"trang": 31, "thep": true}                              → tổng khối lượng thép ghi trên bảng thép của trang (tấn)
  {"trang": 40, "thep_ghi_0": true}                        → các thanh bản vẽ ghi 0 kg, tính dài × SL × kg/m (tấn)
  {"tong": [{"he_so": 3, ...}, ...]}                       → tổng có hệ số, ghi thành công thức Excel để soát được
So nhãn sau khi bỏ dấu, viết hoa và bỏ khoảng trắng (OCR hay dính chữ: 'Cốt thepD<=10').

Dòng diễn giải QS có dạng J = E × F × G × H × I. Dòng sửa: E = số cấu kiện, F = khối lượng bản vẽ, G = H = 1,
I = hệ số ghi trong cấu hình (mặc định 1),
cột C ghi nguồn trang, cột M ghi nghi vấn, cột N giữ khối lượng cũ.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from typing import Dict, List, Optional

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from tools.drawing_rebar_tables import norm, page_rebar_table, zero_weight_rows_kg

FILL_SUA = PatternFill("solid", fgColor="FFF2CC")       # vàng: đã sửa theo bản vẽ
FILL_NGHI_VAN = PatternFill("solid", fgColor="F8CBAD")  # cam: nghi vấn cần soát
FILL_HDR = PatternFill("solid", fgColor="1F497D")
THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


@dataclass
class NghiVan:
    sheet: str
    dong: Optional[int]
    hang_muc: str
    noi_dung: str
    gia_tri_cu: str = ""
    gia_tri_moi: str = ""
    trang: str = ""


def _key(text) -> str:
    return norm(text).replace(" ", "")


def _drawing_number(cell) -> Optional[float]:
    """
    Số trên bản vẽ đọc từ chữ OCR gốc. Bản vẽ dùng dấu chấm thập phân ('148.992' m³ = 148,992, '15.000' phiến = 15)
    và dấu phẩy ngăn nghìn ('1,359.80') — khác cột 'values' của du_lieu.json (chuẩn hóa theo kiểu Việt Nam, đọc
    '148.992' thành 148992), nên không dùng cột đó ở đây.
    """
    t = str(cell or "").strip().replace(" ", "")
    if re.fullmatch(r"\d{1,3}(,\d{3})+(\.\d+)?", t):
        t = t.replace(",", "")
    return float(t) if re.fullmatch(r"\d+(\.\d+)?", t) else None


def ocr_value(page: dict, nhan: str, sau: Optional[str] = None) -> float:
    """Số cuối cùng trên dòng đầu tiên (theo thứ tự bảng/dòng) có chứa nhãn. Không thấy → lỗi (không đoán)."""
    want, anchor = _key(nhan), _key(sau) if sau else None
    seen_anchor = anchor is None
    for t in page.get("tables") or []:
        for row in t.get("rows") or []:
            text = _key(" ".join(str(c) for c in row))
            if not seen_anchor:
                seen_anchor = anchor in text
                continue
            if want in text:
                nums = [x for x in (_drawing_number(c) for c in row) if x is not None]
                if not nums:
                    raise ValueError(f"tr.{page['page']}: dòng '{nhan}' không có số")
                return nums[-1]
    raise ValueError(f"tr.{page['page']}: không thấy dòng '{nhan}'" + (f" sau '{sau}'" if sau else ""))


def drawing_value(spec: dict, pages: Dict[int, dict]) -> float:
    page = pages[spec["trang"]]
    if spec.get("thep") or spec.get("thep_ghi_0"):
        t = page_rebar_table(page)
        if t is None or t.stated_total_kg is None:
            raise ValueError(f"tr.{spec['trang']}: không đọc được tổng thép")
        if spec.get("thep_ghi_0"):                  # các thanh bản vẽ ghi 0 kg: dài × SL × kg/m
            return round(zero_weight_rows_kg(t) / 1000.0, 5)
        return round(t.stated_total_kg / 1000.0, 5)
    return ocr_value(page, spec["nhan"], spec.get("sau"))


def _num_text(v: float) -> str:
    """Số ghi vào công thức Excel, giữ đủ 6 chữ số thập phân (không cắt còn 6 chữ số có nghĩa như '{:g}')."""
    return f"{v:.6f}".rstrip("0").rstrip(".")


def resolve_f(spec, pages: Dict[int, dict]):
    """(giá trị số, giá trị ghi vào ô F: số hoặc công thức '=3*1.13+...')."""
    if isinstance(spec, (int, float)):
        return float(spec), spec
    if "tong" in spec:
        terms = [(t.get("he_so", 1), drawing_value(t, pages)) for t in spec["tong"]]
        value = round(sum(k * v for k, v in terms), 5)
        return value, "=" + "+".join(f"{k:g}*{_num_text(v)}" if k != 1 else _num_text(v) for k, v in terms)
    v = drawing_value(spec, pages)
    return v, v


def _old_j(ws, r: int) -> Optional[float]:
    """Khối lượng cũ của dòng diễn giải (J = E·F·G·H·I) hoặc số ghi cứng."""
    j = ws.cell(r, 10).value
    if isinstance(j, (int, float)):
        return float(j)
    if isinstance(j, str) and re.fullmatch(r"=E\d+\*F\d+\*G\d+\*H\d+\*I\d+", j.replace(" ", "")):
        prod = 1.0
        for c in range(5, 10):
            v = ws.cell(r, c).value
            if not isinstance(v, (int, float)):
                return None
            prod *= v
        return prod
    return None


def _item_title(ws, r: int) -> str:
    """Tên hạng mục (dòng tiêu đề có số TT ở cột A) chứa dòng r."""
    for rr in range(r, 5, -1):
        a = ws.cell(rr, 1).value
        if a is not None and str(a).strip().isdigit():
            return f"{a}. {ws.cell(rr, 3).value}"
    return ""


def fmt(x: Optional[float]) -> str:
    return "" if x is None else f"{x:,.3f}".replace(",", " ").replace(".", ",").replace(" ", ".")


def apply_qs(ws, cfg: dict, pages: Dict[int, dict], sheet_label: str, note: str) -> List[NghiVan]:
    flags: List[NghiVan] = []
    ws["M5"], ws["N5"] = "NGHI VẤN (soát & ghim)", "Khối lượng cũ (trước sửa)"
    for c in ("M5", "N5"):
        ws[c].font, ws[c].fill, ws[c].border = Font(bold=True, color="FFFFFF"), FILL_HDR, BORDER
        ws[c].alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.column_dimensions["M"].width, ws.column_dimensions["N"].width = 70, 16
    ws["A1"] = note
    ws["A1"].font = Font(bold=True, color="C00000")

    def flag(r: int, text: str, cu="", moi="", trang=""):
        cell = ws.cell(r, 13)
        cell.value = (cell.value + " | " if cell.value else "") + text
        cell.fill, cell.border = FILL_NGHI_VAN, BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        flags.append(NghiVan(sheet_label, r, _item_title(ws, r), text, cu, moi, trang))

    for r_txt, formula in cfg.get("tieu_de_thanh_cong_thuc", {}).items():
        r = int(r_txt)
        old = ws.cell(r, 10).value
        if isinstance(old, (int, float)):
            ws.cell(r, 14).value = old
            ws.cell(r, 10).value = formula
            ws.cell(r, 10).fill = FILL_SUA
            flag(r, f"Tiêu đề ghi cứng {fmt(old)} — đổi thành tổng các dòng diễn giải ({formula})", fmt(old), formula)

    for item in cfg["dong"]:
        r = item["dong"]
        old = _old_j(ws, r)
        e = item["E"]
        if "F" in item:
            value, f_cell = resolve_f(item["F"], pages)
        else:
            value, f_cell = 0.0, 0
        he_so = item.get("I", 1)                  # hệ số (vd. nở thành cọc 1,12 — chủ dự án quyết định)
        ws.cell(r, 3).value = item["C"]
        for c, v in zip(range(5, 10), (e, f_cell, 1, 1, he_so)):
            ws.cell(r, c).value = v
            ws.cell(r, c).fill = FILL_SUA
        ws.cell(r, 10).value = f"=E{r}*F{r}*G{r}*H{r}*I{r}"
        ws.cell(r, 10).fill = FILL_SUA
        ws.cell(r, 14).value = round(old, 3) if old is not None else None
        new = e * value * he_so
        pages_txt = ", ".join(sorted(set(re.findall(r"tr\.\d+", item["C"]))))
        changed = old is None or abs(new - old) > max(0.005, 0.001 * abs(old))
        text = f"Theo bản vẽ: {fmt(old)} → {fmt(new)}" if changed else "Khớp bản vẽ"
        if item.get("nghi_van"):
            text += ". " + item["nghi_van"]
        if changed or item.get("nghi_van"):
            flag(r, text, fmt(old), fmt(new), pages_txt)

    for item in cfg.get("nghi_van_dong", []):
        flag(item["dong"], item["noi_dung"])
    return flags


def write_flag_sheet(wb, flags: List[NghiVan], thieu: List[dict], title: str) -> None:
    if "NGHI_VAN" in wb.sheetnames:
        del wb["NGHI_VAN"]
    ws = wb.create_sheet("NGHI_VAN")
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=12)
    hdr = ["STT", "Sheet", "Dòng", "Hạng mục", "Nội dung nghi vấn", "Giá trị cũ", "Giá trị mới", "Trang bản vẽ",
           "Đã soát / Ghim"]
    for c, h in enumerate(hdr, start=1):
        cell = ws.cell(3, c, h)
        cell.font, cell.fill, cell.border = Font(bold=True, color="FFFFFF"), FILL_HDR, BORDER
        cell.alignment = Alignment(wrap_text=True, horizontal="center")
    r = 4
    for i, f in enumerate(flags, start=1):
        for c, v in enumerate([i, f.sheet, f.dong, f.hang_muc, f.noi_dung, f.gia_tri_cu, f.gia_tri_moi, f.trang, ""],
                              start=1):
            cell = ws.cell(r, c, v)
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    for i, t in enumerate(thieu, start=len(flags) + 1):
        for c, v in enumerate([i, "(thiếu trong dự toán)", None, t["hang_muc"],
                               f"Có trên bản vẽ nhưng chưa có trong dự toán: {t['khoi_luong']}", "", t["khoi_luong"],
                               t["trang"], ""], start=1):
            cell = ws.cell(r, c, v)
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if c == 5:
                cell.fill = FILL_NGHI_VAN
        r += 1
    for col, w in zip("ABCDEFGHI", (6, 24, 7, 40, 90, 14, 16, 18, 16)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"


def apply_03a(wb, cfg03: dict, qs_sheet: str, flags: List[NghiVan]) -> None:
    ws = wb[cfg03["sheet"]]
    ws["M6"] = "NGHI VẤN (soát & ghim)"
    ws["M6"].font, ws["M6"].fill, ws["M6"].border = Font(bold=True, color="FFFFFF"), FILL_HDR, BORDER
    ws["M6"].alignment = Alignment(wrap_text=True, horizontal="center")
    ws.column_dimensions["M"].width = 70
    q = qs_sheet

    def flag(r: int, text: str, cu="", moi=""):
        cell = ws.cell(r, 13)
        cell.value = (cell.value + " | " if cell.value else "") + text
        cell.fill, cell.border = FILL_NGHI_VAN, BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        flags.append(NghiVan(cfg03["sheet"], r, f"{ws.cell(r, 1).value}. {ws.cell(r, 2).value}", text, cu, moi))

    for r_txt, qs_row in cfg03["dong_qs"].items():
        r = int(r_txt)
        old_d, old_e = ws.cell(r, 4).value, ws.cell(r, 5).value
        if isinstance(qs_row, list):
            d = "=" + "+".join(f"{q}!J{x}" for x in qs_row)
            e = f"=({'+'.join(f'{q}!L{x}' for x in qs_row)})/D{r}"
        else:
            d, e = f"={q}!J{qs_row}", f"={q}!K{qs_row}"
        if (old_d, old_e) != (d, e):
            ws.cell(r, 4).value, ws.cell(r, 5).value = d, e
            ws.cell(r, 4).fill = ws.cell(r, 5).fill = FILL_SUA
            flag(r, f"Sửa tham chiếu QS: {old_d} → {d} (dòng cũ trỏ vào dòng diễn giải, không phải dòng hạng mục — "
                    f"đơn giá về 0)", str(old_d), d)
        if r_txt in cfg03.get("B", {}):
            ws.cell(r, 2).value = cfg03["B"][r_txt]
        if r_txt in cfg03.get("nghi_van_dong", {}):
            flag(r, cfg03["nghi_van_dong"][r_txt])
        h = ws.cell(r, 8).value
        lnote = ws.cell(r, 12).value
        text = cfg03["nghi_van_chung"] + f" (H = {h})"
        if isinstance(lnote, str) and re.search(r"\d", lnote):
            text += f". Ghi chú cột L '{lnote}' là số cũ, không cập nhật theo D — soát lại"
        flag(r, text)
    if cfg03.get("thieu"):
        flag(28, cfg03["thieu"])
