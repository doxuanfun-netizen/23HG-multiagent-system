# -*- coding: utf-8 -*-
import datetime as dt
import os
import sys
import tempfile
import unittest

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "1_Scripts_TuDongHoa"))
import cpm_compare  # noqa: E402


def make_workbook(path, es_b):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TIEN_DO"
    ws["F2"] = dt.datetime(2024, 1, 1)
    data = [  # stt, wbs, name, kind, pred, dur, K, ES, EF
        (1, "1.1", "A", "Task", None, 3, 1, dt.datetime(2024, 1, 1), dt.datetime(2024, 1, 3)),
        (2, "1.2", "B", "Task", "1.1FS", 2, 1, es_b, dt.datetime(2024, 1, 5)),
    ]
    for i, (stt, w, n, k, p, d, kk, es, ef) in enumerate(data, start=6):
        for c, v in enumerate((stt, w, n, k, p, d, kk, es, ef), start=1):
            ws.cell(i, c, v)
    wb.create_sheet("NGAY_NGHI_LE")
    wb.save(path)


class TestCompare(unittest.TestCase):
    def test_matching_workbook_has_no_mismatch(self):
        with tempfile.TemporaryDirectory() as td:
            p = os.path.join(td, "ok.xlsx")
            make_workbook(p, dt.datetime(2024, 1, 4))
            bad, n = cpm_compare.compare(p)
        self.assertEqual((bad, n), ([], 2))

    def test_wrong_stored_date_is_reported(self):
        with tempfile.TemporaryDirectory() as td:
            p = os.path.join(td, "bad.xlsx")
            make_workbook(p, dt.datetime(2024, 1, 5))  # đáng lẽ 04/01
            bad, _ = cpm_compare.compare(p)
        self.assertEqual([b[0] for b in bad], ["1.2"])
        self.assertEqual(bad[0][2], dt.date(2024, 1, 4))

    def test_main_exit_code(self):
        with tempfile.TemporaryDirectory() as td:
            p = os.path.join(td, "bad.xlsx")
            make_workbook(p, dt.datetime(2024, 1, 5))
            self.assertEqual(cpm_compare.main([p]), 1)
            make_workbook(p, dt.datetime(2024, 1, 4))
            self.assertEqual(cpm_compare.main([p]), 0)


if __name__ == "__main__":
    unittest.main()
