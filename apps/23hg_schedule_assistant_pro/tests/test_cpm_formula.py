# -*- coding: utf-8 -*-
"""Kiểm thử công thức Excel CPM: dựng bảng -> LibreOffice tính lại -> so với giá trị tính tay và bộ Python.
Tự bỏ qua nếu không có LibreOffice. Mất khoảng 20-40 giây."""
import datetime as dt
import os
import shutil
import sys
import tempfile
import unittest

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "1_Scripts_TuDongHoa"))
from _libreoffice import calc_available  # noqa: E402
import build_cpm_formula as bf  # noqa: E402
import cpm_engine as ce  # noqa: E402

D = dt.date
START = dt.datetime(2024, 1, 1)
HAVE_SOFFICE = calc_available()


def d(x):
    return x.date() if isinstance(x, dt.datetime) else x


def build_and_read(rows, start, holidays):
    td = tempfile.mkdtemp()
    path = os.path.join(td, "t.xlsx")
    bf.build_workbook(rows, start, holidays, path)
    bf.recalc_with_libreoffice(path)
    ws = openpyxl.load_workbook(path, data_only=True)["TIEN_DO_CPM"]
    out = {}
    for i, row in enumerate(rows):
        r = bf.FIRST + i
        g = lambda k: ws.cell(r, bf.C[k]).value  # noqa: E731
        out[row["wbs"]] = dict(es=d(g("d_es")), ef=d(g("d_ef")), ls=d(g("d_ls")), lf=d(g("d_lf")),
                               tf=g("d_tf"), crit=g("crit") == "GĂNG", chk=g("chk"))
    errors = [(r, c, ws.cell(r, c).value) for r in range(bf.FIRST, bf.FIRST + len(rows)) for c in range(1, 34)
              if isinstance(ws.cell(r, c).value, str) and ws.cell(r, c).value[:1] == "#" or
              isinstance(ws.cell(r, c).value, str) and ws.cell(r, c).value.startswith("Err:")]
    shutil.rmtree(td, ignore_errors=True)
    return out, errors


def row(wbs, dur, preds=(), kind="Task"):
    return dict(stt=None, wbs=wbs, name=wbs, kind=kind, dur=dur, preds=list(preds), stored_es=None, stored_ef=None)


@unittest.skipUnless(HAVE_SOFFICE, "Cần LibreOffice để tính lại công thức")
class TestFormulaHandCalculated(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.net, cls.err1 = build_and_read([
            row("A", 3), row("B", 2, [("A", "FS", 0)]), row("C", 5, [("A", "FS", 0)]),
            row("D", 1, [("B", "FS", 0), ("C", "FS", 0)])], START, [])
        cls.ss, cls.err2 = build_and_read([row("A", 10), row("B", 3, [("A", "SS", 2)])], START, [])
        cls.hol, cls.err3 = build_and_read([row("A", 3)], START, [D(2024, 1, 2)])

    def test_no_formula_errors(self):
        self.assertEqual(self.err1 + self.err2 + self.err3, [])

    def test_forward_pass(self):
        n = self.net
        self.assertEqual((n["A"]["es"], n["A"]["ef"]), (D(2024, 1, 1), D(2024, 1, 3)))
        self.assertEqual((n["C"]["es"], n["C"]["ef"]), (D(2024, 1, 4), D(2024, 1, 9)))
        self.assertEqual((n["D"]["es"], n["D"]["ef"]), (D(2024, 1, 10), D(2024, 1, 10)))

    def test_backward_pass_and_float(self):
        n = self.net
        self.assertEqual((n["B"]["ls"], n["B"]["lf"], n["B"]["tf"]), (D(2024, 1, 8), D(2024, 1, 9), 3))
        self.assertEqual([w for w in "ABCD" if n[w]["crit"]], ["A", "C", "D"])
        for w in "ACD":
            self.assertEqual(n[w]["tf"], 0)

    def test_ss_lag_backward(self):
        s = self.ss
        self.assertEqual((s["B"]["es"], s["B"]["ef"]), (D(2024, 1, 3), D(2024, 1, 5)))
        self.assertEqual((s["B"]["ls"], s["B"]["tf"]), (D(2024, 1, 9), 5))
        self.assertEqual(s["A"]["tf"], 0)

    def test_holiday_is_skipped(self):
        self.assertEqual(self.hol["A"]["ef"], D(2024, 1, 4))


@unittest.skipUnless(HAVE_SOFFICE and os.path.exists(bf.SRC_XLSM), "Cần LibreOffice và file .xlsm")
class TestFormulaMatchesEngineOnRealData(unittest.TestCase):
    def test_all_rows_match(self):
        rows, start, hol = bf.rows_from_xlsm()
        got, errors = build_and_read(rows, start, hol)
        self.assertEqual(errors, [])
        tasks = [ce.Task(r["wbs"], r["dur"], r["kind"], r["preds"]) for r in rows]
        res = ce.schedule(tasks, d(start), ce.Calendar(hol))
        for r in rows:
            e, g = res[r["wbs"]], got[r["wbs"]]
            with self.subTest(wbs=r["wbs"]):
                self.assertEqual((g["es"], g["ef"], g["ls"], g["lf"], g["tf"], g["crit"]),
                                 (e.es, e.ef, e.ls, e.lf, e.tf, e.critical))
                self.assertIn(g["chk"], (None, ""))


if __name__ == "__main__":
    unittest.main()
