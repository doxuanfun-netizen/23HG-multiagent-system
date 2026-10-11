# -*- coding: utf-8 -*-
"""Tạo bảng CPM bằng công thức Excel (tính xuôi + tính ngược + TF + đường găng).

Quy ước: ngày công (Thứ Hai–Thứ Bảy, nghỉ CN và lễ), quan hệ FS/SS với trễ (ngày công).
Tiền nhiệm phải nằm TRÊN công tác trong bảng (tránh tham chiếu vòng).
Dùng:  python build_cpm_formula.py [--out FILE] [--recalc]
"""
from __future__ import annotations

import argparse
import datetime as dt
import os
import re
import shutil
import subprocess
import tempfile
from typing import Dict, List, Optional

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC_XLSM = os.path.join(ROOT, "23HG_DU_AN_MAU_TIEN_DO_CHUAN_G1_PRO.xlsm")
DEFAULT_OUT = os.path.join(ROOT, "4_CPM_CONG_THUC_THU_NGHIEM", "TIEN_DO_CPM_CONG_THUC_NHAP.xlsx")
WE = '"0000001"'  # Chủ nhật nghỉ
FIRST = 6

# Cột (chỉ số 1-based)
C = dict(stt=1, wbs=2, name=3, kind=4, dur=5, p1=6, r1=7, l1=8, p2=9, r2=10, l2=11, snet=12,
         es=13, ef=14, ls=15, lf=16, tf=17, lfp=18, lfc1=19, lsc1=20, lfc2=21, lsc2=22,
         d_es=23, d_ef=24, d_ls=25, d_lf=26, d_tf=27, crit=28, s_es=29, s_ef=30, dif_es=31, dif_ef=32, chk=33)
CL = {k: L(v) for k, v in C.items()}


def holidays_from_workbook(ws) -> List[dt.date]:
    out = set()
    for r in range(6, 16):
        ds = re.findall(r"(\d{2}/\d{2}/\d{4})", str(ws.cell(r, 2).value))
        if not ds:
            continue
        a = dt.datetime.strptime(ds[0], "%d/%m/%Y").date()
        b = dt.datetime.strptime(ds[-1], "%d/%m/%Y").date()
        d = a
        while d <= b:
            out.add(d)
            d += dt.timedelta(days=1)
    return sorted(out)


def parse_preds(text) -> List[tuple]:
    out = []
    if not text:
        return out
    for part in str(text).replace(";", ",").split(","):
        m = re.fullmatch(r"\s*([\d\.]*\d)(FS|SS)?([+-]\d+)?\s*", part)
        if not m:
            raise ValueError(f"Không đọc được tiền nhiệm '{part}' (bản nháp chỉ hỗ trợ FS/SS)")
        out.append((m.group(1), m.group(2) or "FS", int(m.group(3) or 0)))
    return out


def rows_from_xlsm(path: str = SRC_XLSM):
    wb = openpyxl.load_workbook(path)
    ws = wb["TIEN_DO"]
    rows = []
    for r in range(FIRST, ws.max_row + 1):
        w = ws.cell(r, 2).value
        if w is None:
            continue
        rows.append(dict(
            stt=ws.cell(r, 1).value, wbs=str(w), name=ws.cell(r, 3).value, kind=ws.cell(r, 4).value,
            dur=ws.cell(r, 6).value or 0, rain=ws.cell(r, 7).value or 1.0, preds=parse_preds(ws.cell(r, 5).value)[:2],
            stored_es=ws.cell(r, 8).value, stored_ef=ws.cell(r, 9).value))
    start = ws["F2"].value or rows and ws.cell(FIRST, 8).value
    return rows, start, holidays_from_workbook(wb["NGAY_NGHI_LE"])


def build_workbook(rows: List[dict], start: dt.datetime, holidays: List[dt.date], out_path: str) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TIEN_DO_CPM"
    hs = wb.create_sheet("NGHI_LE")
    hs["A1"] = "Ngày nghỉ lễ (công trường nghỉ thi công)"
    hs["A1"].font = Font(bold=True)
    for i, d in enumerate(holidays, start=2):
        hs.cell(i, 1, dt.datetime(d.year, d.month, d.day)).number_format = "dd/mm/yyyy"
    hs.column_dimensions["A"].width = 16
    last_h = max(2, 1 + len(holidays))
    wb.defined_names["HOL"] = DefinedName("HOL", attr_text=f"NGHI_LE!$A$2:$A${last_h}")
    wb.defined_names["PROJ_START"] = DefinedName("PROJ_START", attr_text="TIEN_DO_CPM!$F$2")
    wb.defined_names["PROJ_END"] = DefinedName("PROJ_END", attr_text="TIEN_DO_CPM!$M$2")
    wb.defined_names["PROJ_LIMIT"] = DefinedName("PROJ_LIMIT", attr_text="TIEN_DO_CPM!$M$3")

    n = len(rows)
    last = FIRST + n - 1
    ws["A1"] = "TIẾN ĐỘ CPM – BẢN NHÁP CÔNG THỨC (ngày công: nghỉ CN và lễ; quan hệ FS/SS)"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"], ws["F2"] = "Ngày khởi công:", start
    ws["A3"], ws["F3"] = "Hạn chót hợp đồng (tuỳ chọn):", None
    ws["K2"], ws["M2"] = "Kết thúc tính được:", f"=MAX(${CL['ef']}${FIRST}:${CL['ef']}${last})"
    ws["K3"], ws["M3"] = "Mốc tính ngược:", '=IF($F$3="",$M$2,$F$3)'
    for c in ("F2", "F3", "M2", "M3"):
        ws[c].number_format = "dd/mm/yyyy"
    ws["A4"] = ("Ô nhập: E–L. Cột M–V là công thức/phụ trợ (nhóm cột R–V đã thu gọn). "
                "Tiền nhiệm phải nằm TRÊN công tác. Găng = TF ≤ 0.")
    ws["A4"].font = Font(italic=True, color="555555")

    heads = {
        "stt": "STT", "wbs": "Mã WBS", "name": "Nội dung", "kind": "Loại", "dur": "Thời lượng (ngày công)",
        "p1": "Tiền nhiệm 1", "r1": "Quan hệ 1", "l1": "Trễ 1", "p2": "Tiền nhiệm 2", "r2": "Quan hệ 2", "l2": "Trễ 2",
        "snet": "Ràng buộc sớm nhất", "es": "ES", "ef": "EF", "ls": "LS", "lf": "LF", "tf": "TF (ngày công)",
        "lfp": "(LF sơ bộ)", "lfc1": "(LF từ KN1)", "lsc1": "(LS từ KN1)", "lfc2": "(LF từ KN2)", "lsc2": "(LS từ KN2)",
        "d_es": "ES hiển thị", "d_ef": "EF hiển thị", "d_ls": "LS hiển thị", "d_lf": "LF hiển thị",
        "d_tf": "TF hiển thị", "crit": "Găng", "s_es": "ES đã lưu", "s_ef": "EF đã lưu",
        "dif_es": "Lệch ES (ngày lịch)", "dif_ef": "Lệch EF (ngày lịch)", "chk": "Kiểm tra tiền nhiệm"}
    for k, h in heads.items():
        cell = ws.cell(5, C[k], h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    ws.row_dimensions[5].height = 42

    def rng(col, r1, r2):
        return f"${CL[col]}{r1}:${CL[col]}${r2}"

    for i, row in enumerate(rows):
        r = FIRST + i
        kind = row["kind"]
        ws.cell(r, C["stt"], row.get("stt"))
        ws.cell(r, C["wbs"], str(row["wbs"]))
        ws.cell(r, C["name"], row.get("name"))
        ws.cell(r, C["kind"], kind)
        ws.cell(r, C["dur"], row.get("dur", 0))
        for slot, (pk, rk, lk) in enumerate((("p1", "r1", "l1"), ("p2", "r2", "l2"))):
            if slot < len(row.get("preds", [])):
                pw, rel, lag = row["preds"][slot]
                ws.cell(r, C[pk], pw), ws.cell(r, C[rk], rel), ws.cell(r, C[lk], lag)
        if row.get("snet"):
            ws.cell(r, C["snet"], row["snet"]).number_format = "dd/mm/yyyy"
        E, B = f"${CL['dur']}{r}", f"${CL['wbs']}{r}"
        is_leaf = kind != "Summary"

        if is_leaf:
            # ---- tính xuôi ----
            def cand(pk, rk, lk):
                p, rel, lag = f"${CL[pk]}{r}", f"${CL[rk]}{r}", f"${CL[lk]}{r}"
                es_p = f"INDEX({rng('es', 5, r - 1)},MATCH({p},{rng('wbs', 5, r - 1)},0))"
                ef_p = f"INDEX({rng('ef', 5, r - 1)},MATCH({p},{rng('wbs', 5, r - 1)},0))"
                return (f'IF({p}="",0,IF({rel}="SS",WORKDAY.INTL({es_p},{lag},{WE},HOL),'
                        f'WORKDAY.INTL({ef_p},1+{lag},{WE},HOL)))')
            ws.cell(r, C["es"], (f'=MAX(WORKDAY.INTL(PROJ_START-1,1,{WE},HOL),N(${CL["snet"]}{r}),'
                                 f'{cand("p1", "r1", "l1")},{cand("p2", "r2", "l2")})'))
            ws.cell(r, C["ef"], f'=IF({E}=0,${CL["es"]}{r},WORKDAY.INTL(${CL["es"]}{r},{E}-1,{WE},HOL))')

            # ---- tính ngược ----
            below = r + 1 <= last
            def con(slot_p, col_c, rr=r):
                crit = f"{rng(slot_p, rr + 1, last)},{B}"
                return (f'IF(COUNTIFS({crit},{rng(col_c, rr + 1, last)},">0")=0,{{d}},'
                        f'_xlfn.MINIFS({rng(col_c, rr + 1, last)},{crit}))')
            if below:
                lfp = f'=MIN(PROJ_LIMIT,{con("p1", "lfc1").format(d="PROJ_LIMIT")},{con("p2", "lfc2").format(d="PROJ_LIMIT")})'
            else:
                lfp = "=PROJ_LIMIT"
            ws.cell(r, C["lfp"], lfp)
            lfp_c = f"${CL['lfp']}{r}"
            base = f'IF({E}=0,{lfp_c},WORKDAY.INTL({lfp_c},-({E}-1),{WE},HOL))'
            if below:
                ls = f'=MIN({base},{con("p1", "lsc1").format(d=base)},{con("p2", "lsc2").format(d=base)})'
            else:
                ls = "=" + base
            ws.cell(r, C["ls"], ls)
            ws.cell(r, C["lf"], f'=IF({E}=0,${CL["ls"]}{r},WORKDAY.INTL(${CL["ls"]}{r},{E}-1,{WE},HOL))')
            es_c, ls_c = f"${CL['es']}{r}", f"${CL['ls']}{r}"
            ws.cell(r, C["tf"], (f'=IF({ls_c}>={es_c},NETWORKDAYS.INTL({es_c},{ls_c},{WE},HOL)-1,'
                                 f'-(NETWORKDAYS.INTL({ls_c},{es_c},{WE},HOL)-1))'))
            # đóng góp của dòng này lên tiền nhiệm của nó
            for pk, rk, lk, cf, cs in (("p1", "r1", "l1", "lfc1", "lsc1"), ("p2", "r2", "l2", "lfc2", "lsc2")):
                p, rel, lag = f"${CL[pk]}{r}", f"${CL[rk]}{r}", f"${CL[lk]}{r}"
                ws.cell(r, C[cf], f'=IF(OR({p}="",{rel}="SS"),"",WORKDAY.INTL({ls_c},-(1+{lag}),{WE},HOL))')
                ws.cell(r, C[cs], f'=IF(OR({p}="",{rel}<>"SS"),"",WORKDAY.INTL({ls_c},-{lag},{WE},HOL))')
            if row.get("preds"):
                ws.cell(r, C["chk"], f'=IF(ISNA(MATCH(${CL["p1"]}{r},{rng("wbs", 5, r - 1)},0)),"TIỀN NHIỆM PHẢI Ở TRÊN","")'
                        if r > FIRST else "TIỀN NHIỆM PHẢI Ở TRÊN")

        # ---- cột hiển thị ----
        wr, last_r = FIRST, last
        for dk, lk, fn in (("d_es", "es", "MINIFS"), ("d_ef", "ef", "MAXIFS"), ("d_ls", "ls", "MINIFS"),
                           ("d_lf", "lf", "MAXIFS"), ("d_tf", "tf", "MINIFS")):
            if is_leaf:
                ws.cell(r, C[dk], f"=${CL[lk]}{r}")
            else:
                ws.cell(r, C[dk], f'=_xlfn.{fn}({rng(lk, wr, last_r)},{rng("wbs", wr, last_r)},{B}&".*")')
        ws.cell(r, C["crit"], f'=IF(${CL["d_tf"]}{r}<=0,"GĂNG","")')
        ws.cell(r, C["s_es"], row.get("stored_es"))
        ws.cell(r, C["s_ef"], row.get("stored_ef"))
        ws.cell(r, C["dif_es"], f'=IF(${CL["s_es"]}{r}="","",${CL["d_es"]}{r}-${CL["s_es"]}{r})')
        ws.cell(r, C["dif_ef"], f'=IF(${CL["s_ef"]}{r}="","",${CL["d_ef"]}{r}-${CL["s_ef"]}{r})')
        for k in ("es", "ef", "ls", "lf", "lfp", "lfc1", "lsc1", "lfc2", "lsc2", "d_es", "d_ef", "d_ls", "d_lf", "s_es", "s_ef"):
            ws.cell(r, C[k]).number_format = "dd/mm/yyyy"
        if kind == "Summary":
            for k in ("wbs", "name", "kind"):
                ws.cell(r, C[k]).font = Font(bold=True)

    # định dạng, kiểm tra dữ liệu, tô găng
    widths = dict(stt=6, wbs=10, name=46, kind=10, dur=11, p1=11, r1=9, l1=8, p2=11, r2=9, l2=8, snet=13,
                  es=12, ef=12, ls=12, lf=12, tf=10, d_es=12, d_ef=12, d_ls=12, d_lf=12, d_tf=10, crit=8,
                  s_es=12, s_ef=12, dif_es=11, dif_ef=11, chk=24)
    for k, v in widths.items():
        ws.column_dimensions[CL[k]].width = v
    for k in ("lfp", "lfc1", "lsc1", "lfc2", "lsc2"):
        ws.column_dimensions[CL[k]].width = 12
        ws.column_dimensions[CL[k]].hidden = True
        ws.column_dimensions[CL[k]].outlineLevel = 1
    dv = DataValidation(type="list", formula1='"FS,SS"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{CL['r1']}{FIRST}:{CL['r1']}{last}")
    dv.add(f"{CL['r2']}{FIRST}:{CL['r2']}{last}")
    dv2 = DataValidation(type="list", formula1='"Task,Milestone,Summary"', allow_blank=False)
    ws.add_data_validation(dv2)
    dv2.add(f"{CL['kind']}{FIRST}:{CL['kind']}{last}")
    ws.conditional_formatting.add(
        f"A{FIRST}:{CL['crit']}{last}",
        FormulaRule(formula=[f'${CL["crit"]}{FIRST}="GĂNG"'], fill=PatternFill("solid", bgColor="FFC7CE")))
    ws.freeze_panes = f"{CL['p1']}{FIRST}"
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    wb.save(out_path)
    return out_path


def recalc_with_libreoffice(path: str, timeout: int = 170) -> str:
    """Tính lại bằng LibreOffice để file có giá trị cache; trả về đường dẫn file đã tính (ghi đè)."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError("Không tìm thấy LibreOffice (soffice)")
    with tempfile.TemporaryDirectory() as td:
        subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", td, path],
                       check=True, capture_output=True, timeout=timeout)
        shutil.copyfile(os.path.join(td, os.path.basename(path)), path)
    return path


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--out", default=DEFAULT_OUT)
    ap.add_argument("--recalc", action="store_true", help="tính lại bằng LibreOffice để có giá trị cache")
    a = ap.parse_args()
    rows, start, hol = rows_from_xlsm()
    build_workbook(rows, start, hol, a.out)
    if a.recalc:
        recalc_with_libreoffice(a.out)
    print("Đã tạo:", a.out)


if __name__ == "__main__":
    main()
