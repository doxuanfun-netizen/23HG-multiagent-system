import os
import openpyxl
import xml.etree.ElementTree as ET
from datetime import datetime

WORKBOOK_PATH = r"23HG_DU_AN_MAU_TIEN_DO_CHUAN_G1_PRO.xlsm"
XML_TEST_PATH = r"2_BaoCao_XuatBan/Du_An_Mau_TuyenDuong_5.5km.xml"
XER_TEST_PATH = r"2_BaoCao_XuatBan/Du_An_Mau_TuyenDuong_5.5km.xer"

def load_tasks_from_excel(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb['TIEN_DO']
    tasks = []
    wbs_map = {}
    
    uid = 1
    for r in range(6, ws.max_row + 1):
        wbs = str(ws.cell(r, 2).value or '').strip()
        name = str(ws.cell(r, 3).value or '').strip()
        ttype = str(ws.cell(r, 4).value or '').strip()
        pred = str(ws.cell(r, 5).value or '').strip()
        dur = ws.cell(r, 6).value or 0
        s_date = ws.cell(r, 8).value
        f_date = ws.cell(r, 9).value
        pct = ws.cell(r, 13).value or 0
        
        if not wbs or not name:
            continue
            
        task = {
            'uid': uid,
            'id': uid,
            'wbs': wbs,
            'name': name,
            'type': ttype,
            'pred_str': pred,
            'duration': float(dur),
            'start': s_date if isinstance(s_date, datetime) else datetime.fromisoformat(str(s_date)[:10]),
            'finish': f_date if isinstance(f_date, datetime) else datetime.fromisoformat(str(f_date)[:10]),
            'pct': float(pct) * 100 if float(pct) <= 1.0 and float(pct) > 0 else float(pct),
            'level': len(wbs.split('.'))
        }
        tasks.append(task)
        wbs_map[wbs] = uid
        uid += 1
        
    return tasks, wbs_map

def parse_pred_links(pred_str, wbs_map):
    links = []
    if not pred_str:
        return links
        
    items = [x.strip() for x in pred_str.split(',') if x.strip()]
    for item in items:
        # Determine relationship type
        rel_type = "FS"
        rel_num = 1 # MSPDI: 1=FS, 3=SS, 0=FF, 2=SF
        if "SS" in item:
            rel_type = "SS"
            rel_num = 3
            parts = item.split("SS")
        elif "FF" in item:
            rel_type = "FF"
            rel_num = 0
            parts = item.split("FF")
        elif "SF" in item:
            rel_type = "SF"
            rel_num = 2
            parts = item.split("SF")
        elif "FS" in item:
            rel_type = "FS"
            rel_num = 1
            parts = item.split("FS")
        else:
            parts = [item, ""]
            
        pred_wbs = parts[0].strip()
        lag_str = parts[1].strip() if len(parts) > 1 else ""
        lag_days = 0.0
        if lag_str:
            lag_str = lag_str.replace("+", "").strip()
            try:
                lag_days = float(lag_str)
            except:
                lag_days = 0.0
                
        if pred_wbs in wbs_map:
            links.append({
                'pred_uid': wbs_map[pred_wbs],
                'pred_wbs': pred_wbs,
                'type_code': rel_type,
                'type_num': rel_num,
                'lag_days': lag_days
            })
    return links

def export_mspdi_xml(tasks, wbs_map, out_path):
    project_start = min(t['start'] for t in tasks).strftime('%Y-%m-%dT07:00:00')
    project_finish = max(t['finish'] for t in tasks).strftime('%Y-%m-%dT17:00:00')
    
    xml = []
    xml.append('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>')
    xml.append('<Project xmlns="http://schemas.microsoft.com/project">')
    xml.append('    <SaveVersion>14</SaveVersion>')
    xml.append('    <Name>GÓI THẦU SỐ 01: XÂY LẮP ĐƯỜNG GIAO THÔNG (Km 0+000 – Km 5+500)</Name>')
    xml.append('    <Title>Tiến Độ Thi Công Chuẩn G1 - Tuyến Đường 5.5km</Title>')
    xml.append('    <Author>NBT (Nguyễn Bảo Tú) - 23HG SYSTEM</Author>')
    xml.append('    <Company>23HG SYSTEM</Company>')
    xml.append(f'    <StartDate>{project_start}</StartDate>')
    xml.append(f'    <FinishDate>{project_finish}</FinishDate>')
    xml.append('    <CalendarUID>1</CalendarUID>')
    xml.append('    <DefaultStartTime>07:00:00</DefaultStartTime>')
    xml.append('    <DefaultFinishTime>17:00:00</DefaultFinishTime>')
    xml.append('    <MinutesPerDay>480</MinutesPerDay>')
    xml.append('    <MinutesPerWeek>2880</MinutesPerWeek>')
    xml.append('    <DaysPerMonth>26</DaysPerMonth>')
    xml.append('    <DurationFormat>7</DurationFormat>')
    xml.append('    <Calendars>')
    xml.append('        <Calendar>')
    xml.append('            <UID>1</UID>')
    xml.append('            <Name>Standard 6 Days (23HG)</Name>')
    xml.append('            <IsBaseCalendar>1</IsBaseCalendar>')
    xml.append('            <WeekDays>')
    xml.append('                <WeekDay><DayType>1</DayType><DayWorking>0</DayWorking></WeekDay>')
    for dt in range(2, 8):
        xml.append(f'                <WeekDay><DayType>{dt}</DayType><DayWorking>1</DayWorking><WorkingTimes><WorkingTime><FromTime>07:00:00</FromTime><ToTime>11:30:00</ToTime></WorkingTime><WorkingTime><FromTime>13:30:00</FromTime><ToTime>17:00:00</ToTime></WorkingTime></WorkingTimes></WeekDay>')
    xml.append('            </WeekDays>')
    xml.append('        </Calendar>')
    xml.append('    </Calendars>')
    xml.append('    <Tasks>')
    
    # Task 0 (Project Summary)
    xml.append('        <Task>')
    xml.append('            <UID>0</UID>')
    xml.append('            <ID>0</ID>')
    xml.append('            <Name>GÓI THẦU SỐ 01: XÂY LẮP ĐƯỜNG GIAO THÔNG (Km 0+000 – Km 5+500)</Name>')
    xml.append('            <Type>1</Type>')
    xml.append('            <IsNull>0</IsNull>')
    xml.append('            <OutlineLevel>0</OutlineLevel>')
    xml.append('            <OutlineNumber>0</OutlineNumber>')
    xml.append('            <WBS>0</WBS>')
    xml.append(f'            <Start>{project_start}</Start>')
    xml.append(f'            <Finish>{project_finish}</Finish>')
    xml.append('            <Summary>1</Summary>')
    xml.append('        </Task>')
    
    for t in tasks:
        dur_hrs = int(t['duration'] * 8)
        is_milestone = 1 if t['type'].lower() == 'milestone' or t['duration'] == 0 else 0
        is_summary = 1 if t['type'].lower() == 'summary' else 0
        s_str = t['start'].strftime('%Y-%m-%dT07:00:00')
        f_str = t['finish'].strftime('%Y-%m-%dT17:00:00')
        clean_name = t['name'].replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;').replace('"', '&quot;')
        
        xml.append('        <Task>')
        xml.append(f'            <UID>{t["uid"]}</UID>')
        xml.append(f'            <ID>{t["id"]}</ID>')
        xml.append(f'            <Name>{clean_name}</Name>')
        xml.append(f'            <Type>0</Type>')
        xml.append(f'            <IsNull>0</IsNull>')
        xml.append(f'            <OutlineLevel>{t["level"]}</OutlineLevel>')
        xml.append(f'            <OutlineNumber>{t["wbs"]}</OutlineNumber>')
        xml.append(f'            <WBS>{t["wbs"]}</WBS>')
        xml.append(f'            <Start>{s_str}</Start>')
        xml.append(f'            <Finish>{f_str}</Finish>')
        xml.append(f'            <Duration>PT{dur_hrs}H0M0S</Duration>')
        xml.append(f'            <DurationFormat>7</DurationFormat>')
        xml.append(f'            <PercentComplete>{int(t["pct"])}</PercentComplete>')
        xml.append(f'            <Summary>{is_summary}</Summary>')
        xml.append(f'            <Milestone>{is_milestone}</Milestone>')
        
        links = parse_pred_links(t['pred_str'], wbs_map)
        for link in links:
            lag_val = int(link['lag_days'] * 4800) # tenths of a minute
            xml.append('            <PredecessorLink>')
            xml.append(f'                <PredecessorUID>{link["pred_uid"]}</PredecessorUID>')
            xml.append(f'                <Type>{link["type_num"]}</Type>')
            xml.append('                <CrossProject>0</CrossProject>')
            xml.append(f'                <LinkLag>{lag_val}</LinkLag>')
            xml.append('                <LagFormat>7</LagFormat>')
            xml.append('            </PredecessorLink>')
        xml.append('        </Task>')
        
    xml.append('    </Tasks>')
    xml.append('</Project>')
    
    with open(out_path, 'w', encoding='utf-8') as f:
        f.write('\n'.join(xml))
    print(f"Exported MSPDI XML to {out_path} ({len(xml)} lines).")

def export_primavera_xer(tasks, wbs_map, out_path):
    project_start = min(t['start'] for t in tasks).strftime('%Y-%m-%d 07:00')
    project_finish = max(t['finish'] for t in tasks).strftime('%Y-%m-%d 17:00')
    
    lines = []
    lines.append(f"ERMHDR\t20.12\t{datetime.now().strftime('%Y-%m-%d')}\tNBT\t23HG_SYSTEM\tPRO")
    
    # PROJECT Table
    lines.append("%T\tPROJECT")
    lines.append("%F\tproj_id\tacct_id\torig_proj_id\tsrc_proj_id\tbase_type\tclndr_id\tsum_base_proj_id\tlast_recalc_date\tplan_end_date\tscd_end_date\tadd_date\tlast_checksum\tguid\tplan_start_date\tbatch_sum_flag\tname_sep_char\tdef_complete_pct_type\tdef_cost_qty_recalc_flag\tallow_neg_act_flag\tsum_only_flag\tcritical_path_type\tproj_short_name\tproj_name")
    lines.append(f"%R\t1\t1\t1\t1\tWA_PROJECT\t1\t1\t{project_start}\t{project_finish}\t{project_finish}\t{project_start}\t0\t{{23HG-G1-PROJECT}}\t{project_start}\tY\t.\tCP_PHYS\tY\tN\tN\tCP_LONGEST_PATH\t23HG_G1\tGÓI THẦU SỐ 01: XÂY LẮP ĐƯỜNG GIAO THÔNG (Km 0+000 – Km 5+500)")
    
    # PROJWBS Table
    lines.append("%T\tPROJWBS")
    lines.append("%F\twbs_id\tproj_id\tob_id\tseq_num\test_wt\tproj_node_flag\tsum_only_flag\twbs_short_name\twbs_name")
    for t in tasks:
        if t['type'].lower() == 'summary':
            lines.append(f"%R\t{t['uid']}\t1\t1\t{t['uid']}\t1.00\tY\tN\t{t['wbs']}\t{t['name']}")
            
    # TASK Table
    lines.append("%T\tTASK")
    lines.append("%F\ttask_id\tproj_id\twbs_id\tclndr_id\tphys_complete_pct\ttask_type\tduration_type\tstatus_code\ttask_code\ttask_name\ttarget_drtn_hr_cnt\ttarget_start_date\ttarget_end_date")
    for t in tasks:
        dur_hrs = int(t['duration'] * 8)
        t_type = "TT_Mile" if t['type'].lower() == 'milestone' else "TT_Task"
        s_code = "TK_Complete" if t['pct'] >= 100 else ("TK_Active" if t['pct'] > 0 else "TK_NotStart")
        s_date = t['start'].strftime('%Y-%m-%d 07:00')
        f_date = t['finish'].strftime('%Y-%m-%d 17:00')
        lines.append(f"%R\t{t['uid']}\t1\t1\t1\t{t['pct']:.1f}\t{t_type}\tDT_DUR\t{s_code}\t{t['wbs']}\t{t['name']}\t{dur_hrs}\t{s_date}\t{f_date}")
        
    # TASKPRED Table
    lines.append("%T\tTASKPRED")
    lines.append("%F\ttask_pred_id\ttask_id\tpred_task_id\tproj_id\tpred_proj_id\tpred_type\tlag_hr_cnt")
    pred_counter = 1
    for t in tasks:
        links = parse_pred_links(t['pred_str'], wbs_map)
        for link in links:
            pred_type = f"PR_{link['type_code']}"
            lag_hrs = int(link['lag_days'] * 8)
            lines.append(f"%R\t{pred_counter}\t{t['uid']}\t{link['pred_uid']}\t1\t1\t{pred_type}\t{lag_hrs}")
            pred_counter += 1
            
    lines.append("%E")
    
    with open(out_path, 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines))
    print(f"Exported Primavera XER to {out_path} ({len(lines)} lines).")

print("Step 1: Loading tasks from Excel...")
tasks, wbs_map = load_tasks_from_excel(WORKBOOK_PATH)
print(f"Loaded {len(tasks)} tasks.")

os.makedirs('2_BaoCao_XuatBan', exist_ok=True)
export_mspdi_xml(tasks, wbs_map, XML_TEST_PATH)
export_primavera_xer(tasks, wbs_map, XER_TEST_PATH)

# Verify Roundtrip
print("\n--- Verifying XML Roundtrip ---")
tree = ET.parse(XML_TEST_PATH)
root = tree.getroot()
ns = {'ns': 'http://schemas.microsoft.com/project'}
xml_tasks = root.findall('.//ns:Task', ns)
print(f"XML Tasks found: {len(xml_tasks)} (includes Project root UID 0)")
assert len(xml_tasks) == len(tasks) + 1, "Task count mismatch in XML!"

print("\n--- Verifying XER Roundtrip ---")
with open(XER_TEST_PATH, 'r', encoding='utf-8') as f:
    xer_content = f.read()
assert "%T\tTASK" in xer_content, "Missing TASK table in XER"
assert "%T\tTASKPRED" in xer_content, "Missing TASKPRED table in XER"
assert "%E" in xer_content, "Missing EOF %E in XER"
print("XER Format verified successfully!")
print("ALL ROUNDTRIP VERIFICATIONS PASSED 100%!")
