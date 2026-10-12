# -*- coding: utf-8 -*-
"""
Bộ kiểm thử đơn vị toàn diện cho Động cơ Bóc tách Khối lượng QS Takeoff (Native COM Interop).
Kiểm tra: get_units, quantity_by_layer, count_blocks, export_boq (Zero Dead Numbers).
"""
import unittest
from unittest.mock import MagicMock, patch
import os
import math
import tempfile
import openpyxl

from cad_controller import CADController


class TestQSTakeoffEngine(unittest.TestCase):
    """Kiểm thử toàn diện 4 công cụ bóc tách khối lượng và xuất BoQ."""

    def setUp(self):
        self.controller = CADController()
        self.mock_doc = MagicMock()
        self.mock_doc.Name = "MatBangKetCau_Tang1.dwg"
        self.mock_doc.FullName = "C:/Projects/MatBangKetCau_Tang1.dwg"
        self.mock_doc.Path = "C:/Projects"
        self.mock_doc.Saved = True

        self.mock_app = MagicMock()
        self.mock_app.ActiveDocument = self.mock_doc
        self.controller.app = self.mock_app
        self.controller.doc = self.mock_doc

    def test_get_units_default_mm(self):
        """Kiểm tra get_units đọc INSUNITS=4 (mm) và INSUNITS=0 (Unspecified quy về mm)."""
        # Case 1: INSUNITS = 4 (Millimeters)
        self.mock_doc.GetVariable.side_effect = lambda var: {
            "INSUNITS": 4,
            "MEASUREMENT": 1,
            "DIMLFAC": 1.0,
            "LUNITS": 2
        }.get(var, 0)

        res = self.controller.get_units()
        self.assertTrue(res["success"])
        data = res["data"]
        self.assertEqual(data["insunits"], 4)
        self.assertEqual(data["unit_name"], "Millimeters")
        self.assertAlmostEqual(data["scale_to_m"], 0.001)
        self.assertAlmostEqual(data["scale_to_m2"], 1e-6)
        self.assertIsNone(data["warning"])

        # Case 2: INSUNITS = 0 (Unspecified -> mặc định mm tại VN)
        self.mock_doc.GetVariable.side_effect = lambda var: {
            "INSUNITS": 0,
            "MEASUREMENT": 1,
            "DIMLFAC": 1.0,
            "LUNITS": 2
        }.get(var, 0)

        res0 = self.controller.get_units()
        self.assertTrue(res0["success"])
        self.assertAlmostEqual(res0["data"]["scale_to_m"], 0.001)

    def test_get_units_meters_and_dimlfac_warning(self):
        """Kiểm tra get_units phát hiện đơn vị Meters và cảnh báo đỏ khi DIMLFAC != 1.0."""
        self.mock_doc.GetVariable.side_effect = lambda var: {
            "INSUNITS": 6,  # Meters
            "MEASUREMENT": 1,
            "DIMLFAC": 0.8,  # Lỗi co giãn kích thước
            "LUNITS": 2
        }.get(var, 0)

        res = self.controller.get_units()
        self.assertTrue(res["success"])
        data = res["data"]
        self.assertEqual(data["insunits"], 6)
        self.assertEqual(data["unit_name"], "Meters")
        self.assertAlmostEqual(data["scale_to_m"], 1.0)
        self.assertAlmostEqual(data["scale_to_m2"], 1.0)
        self.assertIsNotNone(data["warning"])
        self.assertIn("DIMLFAC=0.8", data["warning"])

    def test_quantity_by_layer_geometry_aggregation(self):
        """Kiểm tra quantity_by_layer duyệt ModelSpace, cộng dồn chiều dài và diện tích theo từng layer."""
        # Giả lập đơn vị mm (scale_to_m = 0.001, scale_to_m2 = 1e-6)
        self.mock_doc.GetVariable.side_effect = lambda var: {
            "INSUNITS": 4,
            "MEASUREMENT": 1,
            "DIMLFAC": 1.0,
            "LUNITS": 2
        }.get(var, 0)

        # Tạo danh sách thực thể giả lập
        ent_line1 = MagicMock()
        ent_line1.ObjectName = "AcDbLine"
        ent_line1.Layer = "KET_CAU"
        ent_line1.Length = 5000.0  # 5m

        ent_pline1 = MagicMock()
        ent_pline1.ObjectName = "AcDbPolyline"
        ent_pline1.Layer = "KET_CAU"
        ent_pline1.Length = 10000.0  # 10m
        ent_pline1.Closed = True
        ent_pline1.Area = 6000000.0  # 6 m2

        ent_circle = MagicMock()
        ent_circle.ObjectName = "AcDbCircle"
        ent_circle.Layer = "COT_TRON"
        ent_circle.Radius = 500.0  # r = 0.5m -> Area = pi * 0.25 m2
        ent_circle.Area = math.pi * 500.0 * 500.0

        ent_hatch = MagicMock()
        ent_hatch.ObjectName = "AcDbHatch"
        ent_hatch.Layer = "HOAN_THIEN"
        ent_hatch.Area = 12000000.0  # 12 m2

        mock_entities = [ent_line1, ent_pline1, ent_circle, ent_hatch]
        self.mock_doc.ModelSpace.Count = len(mock_entities)
        self.mock_doc.ModelSpace.Item.side_effect = lambda idx: mock_entities[idx]

        res = self.controller.quantity_by_layer()
        self.assertTrue(res["success"])
        data = res["data"]
        self.assertEqual(data["layers_count"], 3)

        # Tìm layer KET_CAU
        kc_layer = next((l for l in data["layers"] if l["layer"] == "KET_CAU"), None)
        self.assertIsNotNone(kc_layer)
        self.assertEqual(kc_layer["total_entities"], 2)
        # raw_length = 5000 + 10000 = 15000 mm -> length_m = 15.0 m
        self.assertAlmostEqual(kc_layer["raw_length"], 15000.0)
        self.assertAlmostEqual(kc_layer["length_m"], 15.0)
        # raw_area = 6000000 mm2 -> area_m2 = 6.0 m2
        self.assertAlmostEqual(kc_layer["area_m2"], 6.0)
        self.assertEqual(kc_layer["breakdown"]["lines"]["count"], 1)
        self.assertEqual(kc_layer["breakdown"]["polylines"]["count"], 1)
        self.assertEqual(kc_layer["breakdown"]["polylines"]["closed_count"], 1)

        # Kiểm tra lọc theo layer chỉ định
        res_filtered = self.controller.quantity_by_layer(layers=["HOAN_THIEN"])
        self.assertTrue(res_filtered["success"])
        self.assertEqual(res_filtered["data"]["layers_count"], 1)
        self.assertEqual(res_filtered["data"]["layers"][0]["layer"], "HOAN_THIEN")
        self.assertAlmostEqual(res_filtered["data"]["layers"][0]["area_m2"], 12.0)

    def test_count_blocks_dynamic_and_attributes(self):
        """Kiểm tra count_blocks giải mã EffectiveName của Dynamic Block và trích xuất Attributes."""
        # Block 1: Dynamic block Cửa D1 (EffectiveName="CUA_DI_D1", Name="*U10")
        att_d1_type = MagicMock(TagString="LOAI_CUA", TextString="D1")
        att_d1_w = MagicMock(TagString="KICH_THUOC", TextString="900x2200")
        b1 = MagicMock()
        b1.ObjectName = "AcDbBlockReference"
        b1.EffectiveName = "CUA_DI_D1"
        b1.Name = "*U10"
        b1.Layer = "CUA_GO"
        b1.Handle = "A101"
        b1.InsertionPoint = (100.0, 200.0, 0.0)
        b1.HasAttributes = True
        b1.GetAttributes.return_value = [att_d1_type, att_d1_w]

        # Block 2: Cùng loại CUA_DI_D1
        b2 = MagicMock()
        b2.ObjectName = "AcDbBlockReference"
        b2.EffectiveName = "CUA_DI_D1"
        b2.Name = "*U11"
        b2.Layer = "CUA_GO"
        b2.Handle = "A102"
        b2.InsertionPoint = (300.0, 200.0, 0.0)
        b2.HasAttributes = True
        b2.GetAttributes.return_value = [att_d1_type, att_d1_w]

        # Block 3: Static block Cột Bê tông (EffectiveName=None, Name="COT_C1")
        b3 = MagicMock()
        b3.ObjectName = "AcDbBlockReference"
        b3.EffectiveName = None
        b3.Name = "COT_C1"
        b3.Layer = "KET_CAU_COT"
        b3.Handle = "B201"
        b3.InsertionPoint = (0.0, 0.0, 0.0)
        b3.HasAttributes = False

        mock_blocks = [b1, b2, b3]
        self.mock_doc.ModelSpace.Count = len(mock_blocks)
        self.mock_doc.ModelSpace.Item.side_effect = lambda idx: mock_blocks[idx]

        res = self.controller.count_blocks()
        self.assertTrue(res["success"])
        data = res["data"]
        self.assertEqual(data["total_blocks_found"], 3)
        self.assertEqual(data["unique_blocks_count"], 2)

        # CUA_DI_D1 count = 2
        d1_item = next((b for b in data["blocks"] if b["block_name"] == "CUA_DI_D1"), None)
        self.assertIsNotNone(d1_item)
        self.assertEqual(d1_item["count"], 2)
        self.assertEqual(d1_item["layers"], ["CUA_GO"])
        self.assertIn("LOAI_CUA", d1_item["attributes_summary"])
        self.assertIn("D1", d1_item["attributes_summary"]["LOAI_CUA"])

        # COT_C1 count = 1
        c1_item = next((b for b in data["blocks"] if b["block_name"] == "COT_C1"), None)
        self.assertIsNotNone(c1_item)
        self.assertEqual(c1_item["count"], 1)

    def test_export_boq_excel_live_formulas(self):
        """Kiểm tra export_boq xuất file Excel với 100% công thức sống chuẩn Zero Dead Numbers."""
        quantities_data = {
            "units": {"unit_name": "Millimeters", "scale_to_m": 0.001, "scale_to_m2": 1e-6},
            "layers": [
                {
                    "layer": "KET_CAU_DAM",
                    "total_entities": 10,
                    "raw_length": 50000.0,
                    "length_m": 50.0,
                    "raw_area": 25000000.0,
                    "area_m2": 25.0,
                    "breakdown": {"lines": {"count": 6, "length": 30000.0}, "polylines": {"count": 4, "length": 20000.0, "area": 25000000.0, "closed_count": 4}}
                },
                {
                    "layer": "TUONG_220",
                    "total_entities": 8,
                    "raw_length": 40000.0,
                    "length_m": 40.0,
                    "raw_area": 0.0,
                    "area_m2": 0.0,
                    "breakdown": {"lines": {"count": 8, "length": 40000.0}}
                }
            ]
        }
        blocks_data = {
            "blocks": [
                {
                    "block_name": "CUA_DI_D1",
                    "count": 12,
                    "layers": ["CUA"],
                    "attributes_summary": {"WIDTH": {"900": 12}}
                },
                {
                    "block_name": "COT_C1",
                    "count": 8,
                    "layers": ["KET_CAU"],
                    "attributes_summary": {}
                }
            ]
        }

        with tempfile.TemporaryDirectory() as tmp_dir:
            out_file = os.path.join(tmp_dir, "BANG_BOQ_KIEM_THU.xlsx")
            res = self.controller.export_boq(
                file_path=out_file,
                format="xlsx",
                quantities_data=quantities_data,
                blocks_data=blocks_data,
                title="BẢNG TIÊN LƯỢNG KHỐI LƯỢNG THỬ NGHIỆM"
            )

            self.assertTrue(res["success"])
            self.assertTrue(os.path.exists(out_file))

            # Mở file Excel kiểm tra công thức sống (data_only=False để đọc chuỗi công thức)
            wb = openpyxl.load_workbook(out_file, data_only=False)
            self.assertIn("1_Tong_Hop_BoQ", wb.sheetnames)
            self.assertIn("2_Chi_Tiet_Layer", wb.sheetnames)
            self.assertIn("3_Thong_Ke_Block", wb.sheetnames)

            ws1 = wb["1_Tong_Hop_BoQ"]
            ws2 = wb["2_Chi_Tiet_Layer"]
            ws3 = wb["3_Thong_Ke_Block"]

            # 1. Kiểm tra Sheet 2 Chi Tiết Layer:
            # Hàng 4: D4 = 50000, E4 = '=D4*0.001', F4 = 25000000, G4 = '=F4*1e-06'
            self.assertEqual(ws2["E4"].value, "=D4*0.001")
            self.assertEqual(ws2["G4"].value, "=F4*1e-06")
            # Dòng tổng hàng 6 (vì 2 layer: 4, 5 -> tổng hàng 6):
            self.assertEqual(ws2["E6"].value, "=SUM(E4:E5)")
            self.assertEqual(ws2["G6"].value, "=SUM(G4:G5)")

            # 2. Kiểm tra Sheet 3 Thống Kê Block:
            # Hàng 6 (2 block: 4, 5 -> tổng hàng 6):
            self.assertEqual(ws3["D6"].value, "=SUM(D4:D5)")

            # 3. Kiểm tra Sheet 1 Bảng Tổng Hợp:
            # Phải có công thức liên kết trực tiếp tới Sheet 2 và Sheet 3
            self.assertEqual(ws1["D7"].value, "='2_Chi_Tiet_Layer'!E6")
            self.assertEqual(ws1["D8"].value, "='2_Chi_Tiet_Layer'!G6")
            self.assertEqual(ws1["D9"].value, "='3_Thong_Ke_Block'!D6")

    def test_export_boq_csv(self):
        """Kiểm tra export_boq hỗ trợ xuất định dạng CSV sạch sẽ có BOM UTF-8."""
        quantities_data = {
            "units": {"unit_name": "Millimeters", "scale_to_m": 0.001, "scale_to_m2": 1e-6},
            "layers": [{"layer": "DAM", "total_entities": 2, "raw_length": 1000.0, "length_m": 1.0, "raw_area": 0.0, "area_m2": 0.0}]
        }
        blocks_data = {"blocks": [{"block_name": "B1", "count": 5, "layers": ["0"], "attributes_summary": {}}]}

        with tempfile.TemporaryDirectory() as tmp_dir:
            out_file = os.path.join(tmp_dir, "test_boq.csv")
            res = self.controller.export_boq(
                file_path=out_file,
                format="csv",
                quantities_data=quantities_data,
                blocks_data=blocks_data
            )
            self.assertTrue(res["success"])
            self.assertTrue(os.path.exists(out_file))

            with open(out_file, "r", encoding="utf-8-sig") as f:
                content = f.read()
            self.assertIn("CHI TIẾT LAYER", content)
            self.assertIn("THỐNG KÊ BLOCK", content)
            self.assertIn("DAM", content)
            self.assertIn("B1", content)

    def test_smart_cad_command_routing_takeoff(self):
        """Kiểm tra smart_cad_command điều hướng chính xác các lệnh QS Takeoff."""
        self.mock_doc.GetVariable.side_effect = lambda var: 4 if var == "INSUNITS" else 1

        # 1. get_units
        res_u = self.controller.smart_cad_command("get_units")
        self.assertTrue(res_u["success"])
        self.assertEqual(res_u["executionMethod"], "native_com")

        # 2. quantity_by_layer
        self.mock_doc.ModelSpace.Count = 0
        res_q = self.controller.smart_cad_command("quantity_by_layer")
        self.assertTrue(res_q["success"])
        self.assertEqual(res_q["executionMethod"], "native_com")

        # 3. count_blocks
        res_b = self.controller.smart_cad_command("count_blocks")
        self.assertTrue(res_b["success"])
        self.assertEqual(res_b["executionMethod"], "native_com")

        # 4. export_boq
        with tempfile.TemporaryDirectory() as tmp_dir:
            csv_path = os.path.join(tmp_dir, "smart_boq.csv")
            res_e = self.controller.smart_cad_command(
                "export_boq",
                parameters={"filePath": csv_path, "format": "csv"}
            )
            self.assertTrue(res_e["success"])
            self.assertTrue(os.path.exists(csv_path))

    def test_disconnect_and_auto_cleanup(self):
        """Kiểm tra disconnect dọn dẹp biến COM và auto_cleanup trong export_boq."""
        self.assertIsNotNone(self.controller.doc)
        res = self.controller.disconnect()
        self.assertTrue(res["success"])
        self.assertIsNone(self.controller.doc)
        self.assertIsNone(self.controller.app)

        # Test auto_cleanup trong export_boq
        self.controller.doc = self.mock_doc
        self.controller.app = self.mock_app
        quantities_data = {"units": {"unit_name": "mm", "scale_to_m": 0.001, "scale_to_m2": 1e-6}, "layers": []}
        blocks_data = {"blocks": []}
        with tempfile.TemporaryDirectory() as tmp_dir:
            out_f = os.path.join(tmp_dir, "clean_boq.csv")
            exp_res = self.controller.export_boq(file_path=out_f, format="csv", quantities_data=quantities_data, blocks_data=blocks_data, auto_cleanup=True)
            self.assertTrue(exp_res["success"])
            self.assertTrue(exp_res.get("cleaned_up"))
            self.assertIsNone(self.controller.doc)
            self.assertIsNone(self.controller.app)


if __name__ == "__main__":
    unittest.main()
