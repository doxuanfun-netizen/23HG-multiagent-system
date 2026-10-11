# -*- coding: utf-8 -*-
"""
Test sơ đồ tiến độ As-Built (tools/generate_visual_bridge_progress.py + examples/data/asbuilt_km19.json).

Lỗi cũ: ngày kế hoạch đúc dầm và "0/15 phiến" gõ cứng ở hàng chục chỗ trong HTML và Excel; đổi tiến độ
phải sửa tay từng chỗ. Nay lịch được tính từ ke_hoach_duc_dam, các test dưới đây giữ điều đó.
"""

import copy
import datetime as dt
import importlib.util
import json
import os
import tempfile
import unittest

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SCRIPT = os.path.join(ROOT, "tools", "generate_visual_bridge_progress.py")


def load_module():
    spec = importlib.util.spec_from_file_location("asbuilt_gen", SCRIPT)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


M = load_module()
CFG = M.load_config()


def with_plan(**changes):
    cfg = copy.deepcopy(CFG)
    cfg["ke_hoach_duc_dam"].update(changes)
    return cfg


class ConfigTest(unittest.TestCase):
    def test_real_config_loads(self):
        items = CFG["hang_muc"]
        self.assertEqual(sum(1 for i in items if i.get("chen_ke_hoach_dam")), 1)
        self.assertEqual(len([i for i in items if not i.get("chen_ke_hoach_dam")]), 17)

    def test_cast_count_out_of_range_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = os.path.join(tmp, "c.json")
            with open(p, "w", encoding="utf-8") as f:
                json.dump(with_plan(da_duc=16), f, ensure_ascii=False)
            with self.assertRaisesRegex(M.AsBuiltConfigError, "da_duc"):
                M.load_config(p)

    def test_unknown_status_rejected(self):
        cfg = copy.deepcopy(CFG)
        cfg["hang_muc"][0]["tinh_trang"] = "XONG"
        with tempfile.TemporaryDirectory() as tmp:
            p = os.path.join(tmp, "c.json")
            with open(p, "w", encoding="utf-8") as f:
                json.dump(cfg, f, ensure_ascii=False)
            with self.assertRaisesRegex(M.AsBuiltConfigError, "tinh_trang"):
                M.load_config(p)

    def test_missing_beam_marker_rejected(self):
        cfg = copy.deepcopy(CFG)
        cfg["hang_muc"] = [i for i in cfg["hang_muc"] if not i.get("chen_ke_hoach_dam")]
        with tempfile.TemporaryDirectory() as tmp:
            p = os.path.join(tmp, "c.json")
            with open(p, "w", encoding="utf-8") as f:
                json.dump(cfg, f, ensure_ascii=False)
            with self.assertRaisesRegex(M.AsBuiltConfigError, "chen_ke_hoach_dam"):
                M.load_config(p)


class BeamPlanTest(unittest.TestCase):
    def test_default_plan_matches_old_dossier_dates(self):
        plan = M.ke_hoach_dam(CFG["ke_hoach_duc_dam"])
        spans = [(n["bat_dau"], n["ket_thuc"]) for n in plan["nhip"]]
        self.assertEqual(spans, [(dt.date(2026, 10, 25), dt.date(2026, 11, 29)),
                                 (dt.date(2026, 11, 29), dt.date(2027, 1, 3)),
                                 (dt.date(2027, 1, 3), dt.date(2027, 2, 7))])
        self.assertEqual(plan["tong"], 15)

    def test_partial_progress_splits_across_spans(self):
        plan = M.ke_hoach_dam(with_plan(da_duc=7)["ke_hoach_duc_dam"])
        self.assertEqual([n["da_duc"] for n in plan["nhip"]], [5, 2, 0])
        self.assertEqual([n["trang_thai"] for n in plan["nhip"]], ["ĐÃ THI CÔNG", "ĐANG TRIỂN KHAI", "CHƯA THI CÔNG"])

    def test_cycle_change_moves_every_date(self):
        plan = M.ke_hoach_dam(with_plan(chu_ky_ngay=6)["ke_hoach_duc_dam"])
        self.assertEqual(plan["nhip"][0]["ket_thuc"], dt.date(2026, 11, 24))      # 25/10 + 5 × 6 ngày
        self.assertEqual(plan["nhip"][2]["ket_thuc"], dt.date(2027, 1, 23))       # 25/10 + 15 × 6 ngày


class OutputTest(unittest.TestCase):
    def test_html_reflects_progress(self):
        html = M.fill_tokens("@@DA_DUC@@/@@TONG_PHIEN@@ | @@N2_TT@@ | @@N1_LICH@@", M.html_tokens(with_plan(da_duc=7)))
        self.assertEqual(html.split(" | ")[0], "7/15")
        self.assertEqual(html.split(" | ")[1], "Đang đúc (2/5 phiến)")
        self.assertTrue(html.split(" | ")[2].startswith("D1 (25/10-01/11) → D2 (01/11-08/11)"))

    def test_full_html_has_no_unfilled_token(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "a.html")
            M.generate_html_dashboard(path, with_plan(da_duc=7))
            with open(path, encoding="utf-8") as f:
                text = f.read()
        self.assertNotIn("@@", text)
        self.assertIn("Hiện tại: 7/15 phiến (Đang đúc)", text)

    def test_html_template_needs_three_spans(self):
        with self.assertRaisesRegex(M.AsBuiltConfigError, "3 nhịp"):
            M.html_tokens(with_plan(so_nhip=4))

    def test_excel_beam_rows_and_summary_come_from_config(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "a.xlsx")
            M.generate_excel_dashboard(path, with_plan(da_duc=7))
            ws = openpyxl.load_workbook(path).active
        rows = {r[0]: r for r in ws.iter_rows(min_row=7, values_only=True) if r[0]}
        self.assertEqual(rows["Đúc dầm Super-T Nhịp 2"][8], "ĐANG TRIỂN KHAI")
        self.assertEqual(rows["Đúc dầm Super-T Nhịp 2"][9], "Đã đúc (2/5)")
        self.assertEqual(rows["Đúc dầm Super-T Nhịp 1"][5], "25/10 - 29/11/2026")
        self.assertIn("Dầm Super-T đã đúc 7/15 phiến", ws["A2"].value)

    def test_no_hardcoded_plan_or_personal_path_in_source(self):
        with open(SCRIPT, encoding="utf-8") as f:
            src = f.read()
        for literal in ("0/15", "25/10/2026", "29/11/2026", "03/01/2027", "07/02/2027", "144.95", ".gemini"):
            self.assertNotIn(literal, src, literal)


if __name__ == "__main__":
    unittest.main()
