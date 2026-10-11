# -*- coding: utf-8 -*-
"""
Test sửa file theo dõi tần suất thí nghiệm theo bản vẽ (tools/tan_suat_theo_ban_ve.py,
examples/sua_tan_suat_km19_theo_ban_ve.py). File thật của đơn vị không có trong repo, nên test dựng một workbook
nhỏ cùng bố cục (sheet bê tông, thép, tổng hợp, DM) rồi chạy công cụ trên dữ liệu bản vẽ + BBS thật.
"""

import os
import tempfile
import unittest

import openpyxl

from tools.tan_suat_theo_ban_ve import _lan, _split

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DU_LIEU = os.path.join(ROOT, "examples", "HO_SO_CAU_KM19_529", "05_DU_LIEU_GOC_SCAN_MARKER", "du_lieu.json")
BT, THEP, TH, DM = "Theo_Doi_Tan_Suat_Be_Tong", "Theo_Doi_Tan_Suat_Cot_Thep", "Tong_Hop_Tan_Suat_Vat_Lieu", "DM"


class RulesTest(unittest.TestCase):
    def test_first_lot_is_tested_when_use_starts(self):
        """File cũ dùng INT(lũy kế/20): Ø dưới 20 t có 0 lần thí nghiệm. Đúng: lần 1 ngay khi bắt đầu dùng."""
        self.assertEqual([_lan(x, 20) for x in (0, 0.1, 13.4, 20, 20.001, 40, 41)], [0, 1, 1, 1, 2, 2, 3])

    def test_split_keeps_drawing_total(self):
        parts = _split(234.92, [42, 42, 38])
        self.assertAlmostEqual(sum(parts), 234.92, places=6)
        self.assertEqual(parts[0], parts[1])


def build_workbook(path):
    wb = openpyxl.Workbook()
    bt = wb.active
    bt.title = BT
    bt["A1"], bt["B1"], bt["C1"], bt["E1"] = "STT", "HẠNG MỤC CÔNG VIỆC", "N/T/N", "KL/T.tế (m3)"
    bt["Q2"], bt["R2"], bt["S2"] = "Xi măng 50 tấn", "Cát 200m3", "Đá 350m3"
    rows = [("Bê tông cọc khoan nhồi - C1 trụ T1 (L=40m)", 47.01, True),
            ("Bê tông thân trụ T1 - Đốt 1 (C30 đá 1x2)", 42, False),
            ("Bê tông thân trụ T1 - Đốt 2 (C30 đá 1x2)", 42, False),
            ("Bê tông thân trụ T1 - Đốt 3 (C30 đá 1x2)", 38, False),
            ("Bê tông bản mặt cầu Nhịp 1 (C35 đá 1x2)", 85, False)]
    prev = None
    for i, (name, e, pile) in enumerate(rows, start=4):
        bt.cell(i, 1, i - 3)
        bt.cell(i, 2, name)
        bt.cell(i, 5, e)
        for c, v in zip((8, 9, 10) if pile else (11, 12, 13), (0.445, 0.525, 0.696) if pile else (0.425, 0.555, 0.7)):
            bt.cell(i, c, v)
        for col, a, b in (("N", "H", "K"), ("O", "I", "L"), ("P", "J", "M")):
            base = f"{col}{prev}+" if prev else ""
            bt[f"{col}{i}"] = f'=IF(E{i}="","",ROUND({base}IF({a}{i}>0,E{i}*{a}{i},E{i}*{b}{i}),2))'
        bt[f"R{i}"] = f'=IF(O{i}="","",IF(INT(O{i}/350)>0,"Lần " & INT(O{i}/350),""))'      # chia ngược như file cũ
        bt[f"S{i}"] = f'=IF(P{i}="","",IF(INT(P{i}/200)>0,"Lần " & INT(P{i}/200),""))'
        prev = i
    st = wb.create_sheet(THEP)
    st["A1"], st["P2"] = "STT", "Ø10 (kèm Ø8)"
    for i, (k, n) in enumerate((("Coc", 1), ("Be", 3), ("Than", 2), ("Dam", 4), ("HoanThien", 2)), start=2):
        st.cell(i, 43, k)
        st.cell(i, 44, n)
    st["A3"], st["A4"] = "PHẦN I", "--- CỌC ---"
    for c, v in enumerate([1, "x", "x", "P1", 10, "CB240-T", "", 1.0, 1, 0.617], start=1):
        st.cell(5, c, v)
    st["A6"] = "TỔNG CỘNG"
    th = wb.create_sheet(TH)
    th["B6"], th["F6"], th["G6"] = "Xi măng", f"={BT}!N8", '=IF(F6="","",IF(F6=0,"1 lần",INT(F6/50)+1 & " lần"))'
    th["B12"], th["F12"], th["G12"] = "Ø10", f"={THEP}!P6", '=IF(F12="","",IF(F12=0,"1 lần",INT(F12/20)+1 & " lần"))'
    dm = wb.create_sheet(DM)
    for c, v in enumerate(["STT", "Hạng mục", "Đối tượng"], start=1):
        dm.cell(13, c, v)
    dm["A14"], dm["B14"], dm["C14"], dm["G14"] = 1, "Phần cầu", f'="Lấy mẫu: "&{TH}!B6', f"={BT}!C4"
    dm["N14"] = '=IFERROR(VALUE(MID(K14,FIND("Lần ",K14)+4,FIND(" -",K14)-FIND("Lần ",K14)-4)),"")'
    wb.save(path)


@unittest.skipUnless(os.path.isfile(DU_LIEU), "thiếu dữ liệu OCR HSTK Km19")
class FrequencyWorkbookTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from examples.sua_tan_suat_km19_theo_ban_ve import run
        cls.tmp = tempfile.TemporaryDirectory()
        src, cls.out = os.path.join(cls.tmp.name, "vao.xlsx"), os.path.join(cls.tmp.name, "ra.xlsx")
        build_workbook(src)
        cls.stats = run(src, cls.out, ngay="11/10/2026")
        cls.wb = openpyxl.load_workbook(cls.out)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_concrete_volumes_follow_drawing(self):
        bt = self.wb[BT]
        self.assertEqual(bt["E4"].value, 45.88)                       # tr.68, đã gồm bê tông đầu cọc
        self.assertAlmostEqual(sum(bt.cell(r, 5).value for r in (5, 6, 7)), 234.92, places=2)   # tr.48 chia 3 đốt
        self.assertEqual(bt["E8"].value, 102.53)
        notes = [bt.cell(r, bt.max_column).value or "" for r in range(4, 9)]
        self.assertTrue(all(notes), notes)
        self.assertIn("C35", notes[-1])                               # mác chưa có cấp phối → nghi vấn

    def test_sand_and_stone_divisors_and_roundup(self):
        bt = self.wb[BT]
        for r in range(4, 9):
            self.assertIn("/200", bt.cell(r, 18).value)
            self.assertIn("/350", bt.cell(r, 19).value)
            self.assertIn("ROUNDUP", bt.cell(r, 17).value)

    def test_steel_rebuilt_from_bbs_and_summary_points_to_new_total(self):
        st, th = self.wb[THEP], self.wb[TH]
        total = next(r for r in range(3, st.max_row + 1) if str(st.cell(r, 1).value or "").startswith("TỔNG CỘNG"))
        self.assertGreater(self.stats["dong_thep"], 700)
        self.assertEqual(th["F12"].value, f"={THEP}!P{total}")
        self.assertIn("ROUNDUP(F12/20,0)", th["G12"].value)
        self.assertEqual(st.cell(total, 16).value, f"=P{total - 1}")
        self.assertIsNone(st["AL5"].value)                            # không ghi sẵn 'ĐẠT'

    def test_dm_has_one_row_per_sampling_event(self):
        dm = self.wb[DM]
        rows = [r for r in range(14, dm.max_row + 1) if dm.cell(r, 2).value]
        self.assertEqual(len(rows), self.stats["dm"])
        self.assertEqual(self.stats["dm"], self.stats["lan_bt"] + self.stats["lan_thep"])
        self.assertTrue(all(str(dm.cell(r, 11).value).startswith('="Lần ') for r in rows))
        self.assertNotIn("K14", dm["N15"].value)                       # công thức dòng mẫu đổi theo dòng mới
        self.assertIn("K15", dm["N15"].value)

    def test_review_sheet(self):
        nv = self.wb["NGHI_VAN"]
        self.assertEqual(nv["I3"].value, "Đã soát / Ghim")
        self.assertGreater(nv.max_row, 10)


if __name__ == "__main__":
    unittest.main()
