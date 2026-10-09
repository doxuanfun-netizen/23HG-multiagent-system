# -*- coding: utf-8 -*-
"""
Test sổ nghiệm thu KCS (tools/kcs_register.py + examples/generate_kcs_word_package.py).

Bắt lỗi cũ: cột "Đánh giá logic" từng là chữ "Hợp lệ" gõ tay cho cả 22 bản ghi, kể cả những bản ghi
mà chính ngày trong sổ mâu thuẫn với điều kiện tiên quyết. Các test dưới đây bảo đảm trạng thái
được TÍNH ra, và biên bản không kết luận "đủ điều kiện" cho bản ghi chưa đạt.
"""

import csv
import datetime as dt
import os
import tempfile
import unittest

import docx

from tools.kcs_register import (CHUA_KIEM, HOP_LE, MAU_THUAN, NGAY_SAI, THAM_CHIEU_SAI, KcsRecord,
                                RegisterError, build_docx, check_register, load_register, load_signers,
                                parse_date)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REGISTER = os.path.join(ROOT, "examples", "data", "kcs_register_km19.csv")
SIGNERS = os.path.join(ROOT, "examples", "data", "kcs_nguoi_ky_km19.json")


def d(s):
    return dt.date(*map(int, reversed(s.split("/"))))


def rec(code, start, finish, refs=()):
    return KcsRecord(code, f"việc {code}", d(start), d(finish), "", tuple(refs))


class RegisterLoadTest(unittest.TestCase):
    def test_real_register_has_22_records_and_no_duplicates(self):
        recs = load_register(REGISTER)
        self.assertEqual(len(recs), 22)
        self.assertEqual(len({r.code for r in recs}), 22)
        self.assertEqual(recs[0].code, "BBNT-01")
        self.assertEqual(recs[-1].code, "BBNT-22")

    def test_references_parsed(self):
        by = {r.code: r for r in load_register(REGISTER)}
        self.assertEqual(by["BBNT-05"].refs, ("BBNT-02", "BBNT-04"))
        self.assertEqual(by["BBNT-01"].refs, ())

    def test_bad_date_is_rejected_with_row(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = os.path.join(tmp, "s.csv")
            with open(p, "w", encoding="utf-8", newline="") as f:
                f.write("ma;cong_viec;bat_dau;nghiem_thu;dieu_kien;tham_chieu\nX-1;việc;2026-10-01;03/10/2026;;\n")
            with self.assertRaises(RegisterError) as cm:
                load_register(p)
            self.assertIn("X-1", str(cm.exception))

    def test_duplicate_code_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = os.path.join(tmp, "s.csv")
            with open(p, "w", encoding="utf-8", newline="") as f:
                f.write("ma;cong_viec;bat_dau;nghiem_thu\nA;v;01/10/2026;02/10/2026\nA;v;01/10/2026;02/10/2026\n")
            with self.assertRaisesRegex(RegisterError, "trùng"):
                load_register(p)

    def test_parse_date_format(self):
        self.assertEqual(parse_date("5/9/2026"), dt.date(2026, 9, 5))
        with self.assertRaises(ValueError):
            parse_date("2026-09-05")


class LogicCheckTest(unittest.TestCase):
    def test_real_register_findings(self):
        """Bốn mâu thuẫn thật trong sổ hiện hành — không được tự động thành HỢP LỆ."""
        chk = check_register(load_register(REGISTER))
        for code in ("BBNT-06", "BBNT-07", "BBNT-09", "BBNT-13"):
            self.assertEqual(chk[code].status, MAU_THUAN, code)
        self.assertIn("BBNT-05 nghiệm thu 28/10/2026", chk["BBNT-06"].detail)
        self.assertIn("BBNT-06 nghiệm thu 05/11/2026", chk["BBNT-07"].detail)
        self.assertIn("BBNT-08 nghiệm thu 07/11/2026", chk["BBNT-09"].detail)
        self.assertIn("BBNT-12 nghiệm thu 08/01/2027", chk["BBNT-13"].detail)

    def test_real_register_counts(self):
        from collections import Counter
        counts = Counter(c.status for c in check_register(load_register(REGISTER)).values())
        self.assertEqual(counts[HOP_LE], 13)
        self.assertEqual(counts[MAU_THUAN], 4)
        self.assertEqual(counts[CHUA_KIEM], 5)

    def test_no_reference_is_never_valid(self):
        self.assertEqual(check_register([rec("A", "01/10/2026", "02/10/2026")])["A"].status, CHUA_KIEM)

    def test_start_before_predecessor_finish_is_conflict(self):
        chk = check_register([rec("A", "01/10/2026", "05/10/2026"),
                              rec("B", "04/10/2026", "06/10/2026", refs=["A"])])
        self.assertEqual(chk["B"].status, MAU_THUAN)

    def test_start_on_or_after_predecessor_finish_is_valid(self):
        chk = check_register([rec("A", "01/10/2026", "05/10/2026"),
                              rec("B", "06/10/2026", "07/10/2026", refs=["A"])])
        self.assertEqual(chk["B"].status, HOP_LE)

    def test_finish_before_start_flagged(self):
        self.assertEqual(check_register([rec("A", "05/10/2026", "01/10/2026")])["A"].status, NGAY_SAI)

    def test_unknown_reference_flagged(self):
        chk = check_register([rec("A", "01/10/2026", "02/10/2026", refs=["ZZ-9"])])
        self.assertEqual(chk["A"].status, THAM_CHIEU_SAI)


class WordOutputTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.recs = load_register(REGISTER)
        cls.chk = check_register(cls.recs)
        cls.path = build_docx(cls.recs, cls.chk, os.path.join(cls.tmp.name, "kcs.docx"),
                              project_name="TEST", location="TEST", signers=load_signers(SIGNERS))
        cls.doc = docx.Document(cls.path)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_one_biên_bản_per_record(self):
        n = sum("BIÊN BẢN NGHIỆM THU" in p.text for p in self.doc.paragraphs)
        self.assertEqual(n, 22)

    def test_check_table_status_matches_computation(self):
        rows = self.doc.tables[0].rows[1:]
        self.assertEqual(len(rows), 22)
        for r in rows:
            code, status = r.cells[0].text, r.cells[5].text
            self.assertTrue(status.startswith(self.chk[code].status), f"{code}: {status!r}")

    def test_mismatched_records_do_not_claim_valid(self):
        body = "\n".join(p.text for p in self.doc.paragraphs)
        for code in ("BBNT-06", "BBNT-07"):
            seg = body.split(f"Số: {code}/NT-GXD", 1)[1].split("BIÊN BẢN NGHIỆM THU", 1)[0]
            self.assertIn("CHƯA đủ điều kiện", seg, code)
            self.assertNotIn("Đồng ý nghiệm thu", seg, code)
            self.assertNotIn("Đã kiểm tra logic chéo", seg, code)

    def test_valid_record_can_pass(self):
        body = "\n".join(p.text for p in self.doc.paragraphs)
        seg = body.split("Số: BBNT-02/NT-GXD", 1)[1].split("BIÊN BẢN NGHIỆM THU", 1)[0]
        self.assertIn("Đồng ý nghiệm thu", seg)

    def test_signer_names_come_from_config_not_code(self):
        signers = load_signers(SIGNERS)
        table_text = " ".join(c.text for t in self.doc.tables[1:3] for r in t.rows for c in r.cells)
        self.assertIn(signers["nt_ten"], table_text)
        with open(os.path.join(ROOT, "tools", "kcs_register.py"), encoding="utf-8") as f:
            src = f.read()
        self.assertNotIn(signers["nt_ten"], src)


class SignerConfigTest(unittest.TestCase):
    def test_missing_signers_become_placeholders(self):
        s = load_signers(None)
        self.assertEqual(s["nt_ten"], "[.........]")


class CsvShapeTest(unittest.TestCase):
    def test_register_csv_uses_semicolon_and_header(self):
        with open(REGISTER, encoding="utf-8-sig", newline="") as f:
            header = next(csv.reader(f, delimiter=";"))
        self.assertEqual(header, ["ma", "cong_viec", "bat_dau", "nghiem_thu", "dieu_kien", "tham_chieu"])


if __name__ == "__main__":
    unittest.main()
