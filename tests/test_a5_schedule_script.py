# -*- coding: utf-8 -*-
"""
Test script cập nhật tiến độ Cống hộp A5 (tools/apply_a5_schedule_0509_1211.py).

Script ghi đè bảng tiến độ ca máy bằng công thức sống. Các test dưới đây chạy trên BẢN SAO TẠM của
workbook thật, không đụng vào file trong repo:
  - mốc thời gian 05/09/2026 → 12/11/2026 = 69 ngày;
  - chạy lại nhiều lần cho kết quả giống nhau (idempotent), không sinh thêm công thức rác;
  - toàn bộ công thức tính được, không ô nào ra lỗi Excel (#REF!, #DIV/0!...), không sheet bị thiếu;
  - các bản sao cùng tên được đồng bộ, không bản nào bị bỏ lại bản cũ.
"""

import datetime
import os
import shutil
import tempfile
import unittest
from unittest import mock

import openpyxl

from tools import apply_a5_schedule_0509_1211 as a5
from tools.audit_excels_static import audit_file

SRC = a5.SRC_FILES[0]


def _read(path):
    with open(path, "rb") as f:
        return f.read()


def _formula_count(path):
    wb = openpyxl.load_workbook(path)
    return sum(1 for ws in wb for row in ws.iter_rows() for c in row
               if isinstance(c.value, str) and c.value.startswith("="))


class ScheduleWindowTest(unittest.TestCase):
    def test_window_is_69_days(self):
        self.assertEqual(a5.D_START, datetime.date(2026, 9, 5))
        self.assertEqual(a5.D_FINISH, datetime.date(2026, 11, 12))
        self.assertEqual(a5.TOTAL_DAYS, 69)
        self.assertEqual(len(a5.DATES), 69)
        self.assertEqual((a5.DATES[0], a5.DATES[-1]), (a5.D_START, a5.D_FINISH))
        self.assertTrue(all(b - a == datetime.timedelta(days=1) for a, b in zip(a5.DATES, a5.DATES[1:])))


@unittest.skipUnless(os.path.exists(SRC), "thiếu workbook A5 mẫu")
class UpdateWorkbookTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp()
        cls.path = os.path.join(cls.tmp, "a5.xlsx")
        shutil.copyfile(SRC, cls.path)
        with mock.patch("builtins.print"):
            a5.update_workbook_live_formulas(cls.path)
        cls.after_first = _formula_count(cls.path)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_is_idempotent(self):
        with mock.patch("builtins.print"):
            a5.update_workbook_live_formulas(self.path)
        self.assertEqual(_formula_count(self.path), self.after_first)

    def test_workbook_is_mostly_live_formulas(self):
        self.assertGreater(self.after_first, 3000)           # bản đã chuyển 100% sang công thức sống

    def test_no_formula_errors_after_evaluation(self):
        r = audit_file(self.path, try_eval=True)
        self.assertIsNone(r["fatal"])
        self.assertEqual(r["cached_errors"], 0)
        self.assertEqual(r["formula_ref_errors"], 0)
        self.assertEqual(r["missing_sheet_refs"], 0)
        self.assertEqual(r["eval_error_results"], 0, r["details"][:3])
        self.assertEqual(r["eval_unsupported"], 0)           # hàm nào chưa hỗ trợ sẽ lộ ra ở đây

    def test_expected_sheets_present(self):
        names = openpyxl.load_workbook(self.path).sheetnames
        for expected in ("01_TienDo_CaMay_Master", "02_TongHop_CaXe_CaMay_MMTB", "03_KeHoach_Dau_Diezel"):
            self.assertIn(expected, names)

    def test_title_states_the_69_day_window(self):
        ws = openpyxl.load_workbook(self.path)["01_TienDo_CaMay_Master"]
        self.assertIn("05/09/2026", ws["A3"].value)
        self.assertIn("12/11/2026", ws["A3"].value)
        self.assertIn("69", ws["A3"].value)


class SyncCopiesTest(unittest.TestCase):
    def test_copies_overwrite_only_same_named_files(self):
        with tempfile.TemporaryDirectory() as tree:
            for sub in ("goiA", "goiE", "master"):
                os.makedirs(os.path.join(tree, sub))
            newest = os.path.join(tree, "goiA", "x.xlsx")
            stale = [os.path.join(tree, "goiE", "x.xlsx"), os.path.join(tree, "master", "x.xlsx")]
            other = os.path.join(tree, "goiE", "khac.xlsx")
            for p, data in [(newest, b"MOI"), (stale[0], b"CU"), (stale[1], b"CU"), (other, b"GIU-NGUYEN")]:
                with open(p, "wb") as f:
                    f.write(data)
            with mock.patch.object(a5, "A5_TREE", tree):
                synced = a5.sync_same_named_copies(newest)
            self.assertEqual(sorted(synced), sorted(stale))
            for p in stale:
                self.assertEqual(_read(p), b"MOI")
            self.assertEqual(_read(other), b"GIU-NGUYEN")     # file khác tên không bị đụng
            self.assertEqual(_read(newest), b"MOI")

    def test_no_copies_means_nothing_synced(self):
        with tempfile.TemporaryDirectory() as tree:
            only = os.path.join(tree, "only.xlsx")
            with open(only, "wb") as f:
                f.write(b"1")
            with mock.patch.object(a5, "A5_TREE", tree):
                self.assertEqual(a5.sync_same_named_copies(only), [])


if __name__ == "__main__":
    unittest.main()
