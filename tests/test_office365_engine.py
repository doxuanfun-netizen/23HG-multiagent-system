# -*- coding: utf-8 -*-
"""
UNIT TESTS — OFFICE 365 TAKEOFF ENGINE
Kiểm thử toàn diện động cơ Microsoft 365 Enterprise Engine trong 23HG:
  - Đăng ký hàm AEC LAMBDA vào Name Manager
  - Khởi tạo công thức tự diễn giải LET()
  - Khởi tạo công thức truy vấn liên kết XLOOKUP()
  - Dựng trang bìa Bảng điều hành Executive Dashboard 365 & Grid Kiểm toán
  - Lưu và nạp lại OpenXML Workbook đảm bảo tính toàn vẹn cấu trúc
"""

import os
import sys
import unittest
import tempfile
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from tools.office365_takeoff_engine import (
    AEC_LAMBDA_DEFINITIONS,
    register_aec_lambdas,
    build_let_formula,
    build_xlookup_formula,
    build_executive_365_dashboard,
    embed_cad_proof_images,
)


class TestOffice365TakeoffEngine(unittest.TestCase):

    def setUp(self):
        self.wb = openpyxl.Workbook()
        self.tmp_dir = tempfile.mkdtemp()

    def test_aec_lambda_registration(self):
        """Đảm bảo 8 hàm kỹ thuật AEC LAMBDA được đăng ký đầy đủ vào Name Manager."""
        registered = register_aec_lambdas(self.wb)
        self.assertEqual(len(registered), 8)
        self.assertIn("V_PRISM", self.wb.defined_names)
        self.assertIn("V_CYLINDER", self.wb.defined_names)
        self.assertIn("V_FRUSTUM", self.wb.defined_names)
        self.assertIn("S_FORMWORK_BOX", self.wb.defined_names)
        self.assertIn("S_FORMWORK_TRI", self.wb.defined_names)
        self.assertIn("STEEL_RATIO", self.wb.defined_names)
        self.assertIn("V_AVERAGE_END", self.wb.defined_names)
        self.assertIn("REBAR_WEIGHT", self.wb.defined_names)

        # Kiểm tra cú pháp OpenXML
        dn = self.wb.defined_names["V_PRISM"]
        self.assertTrue(dn.attr_text.startswith("_xlfn.LAMBDA(_xlpm."))

    def test_build_let_formula(self):
        """Kiểm tra tạo công thức LET() với namespace _xlfn. và prefix _xlpm."""
        vars_dict = {"qty": "D10", "L": "E10", "W": "F10", "H": "G10"}
        expr = "qty * L * W * H"
        f = build_let_formula(vars_dict, expr)
        self.assertTrue(f.startswith("=_xlfn.LET("))
        self.assertIn("_xlpm.qty, D10", f)
        self.assertIn("_xlpm.L, E10", f)
        self.assertIn("_xlpm.W, F10", f)
        self.assertIn("_xlpm.H, G10", f)
        self.assertIn("_xlpm.qty * _xlpm.L * _xlpm.W * _xlpm.H", f)

    def test_build_let_formula_word_boundaries(self):
        """Đảm bảo không bị thay thế chuỗi con khi biến số có tên tương tự (vd H và H_extra)."""
        vars_dict = {"H": "A1", "H_extra": "B1"}
        expr = "H + H_extra"
        f = build_let_formula(vars_dict, expr)
        self.assertIn("_xlpm.H + _xlpm.H_extra", f)

    def test_build_xlookup_formula(self):
        """Kiểm tra công thức XLOOKUP() với wildcard mode 2 mặc định."""
        f = build_xlookup_formula('"*TỔNG BÊ TÔNG*"', "'01_MO'!$B:$B", "'01_MO'!$D:$D")
        self.assertEqual(f, '=_xlfn.XLOOKUP("*TỔNG BÊ TÔNG*", \'01_MO\'!$B:$B, \'01_MO\'!$D:$D, 0, 2)')

    def test_build_executive_365_dashboard(self):
        """Kiểm tra tạo dựng Dashboard 365 với các thẻ KPI và Audit Grid."""
        kpis = [
            {"title": "BÊ TÔNG C30", "formula": "=100", "unit": "m3", "fmt": "#,##0.00"},
            {"title": "CỐT THÉP", "formula": "=20", "unit": "Tấn", "fmt": "#,##0.00"},
        ]
        audit_items = [
            {"stt": "1", "name": "Bê tông móng", "unit": "m3", "cad_formula": "=100", "design_val": 100.0, "note": "Khớp"},
        ]
        ws = build_executive_365_dashboard(
            self.wb,
            project_title="Cầu Thử Nghiệm",
            kpis=kpis,
            audit_items=audit_items
        )
        self.assertEqual(ws.title, "00_DASHBOARD_365")
        self.assertIn("CẦU THỬ NGHIỆM", ws["A1"].value)
        self.assertEqual(ws["A5"].value, "BÊ TÔNG C30")
        self.assertEqual(ws["A6"].value, "=100")
        self.assertEqual(ws["C5"].value, "CỐT THÉP")
        self.assertEqual(ws["C6"].value, "=20")

        # Kiểm tra hàng kiểm toán
        # Audit headers nằm ở dòng 10, item ở dòng 11
        self.assertEqual(ws.cell(11, 2).value, "Bê tông móng")
        self.assertEqual(ws.cell(11, 4).value, "=100")
        self.assertEqual(ws.cell(11, 5).value, 100.0)
        self.assertEqual(ws.cell(11, 6).value, "=D11-E11")
        self.assertIn("TRÙNG KHỚP 100%", ws.cell(11, 8).value)

    def test_roundtrip_save_and_reload(self):
        """Kiểm tra lưu workbook chứa LAMBDA & LET ra đĩa rồi nạp lại không bị mất dữ liệu."""
        register_aec_lambdas(self.wb)
        ws = self.wb.active
        ws.title = "TEST_SHEET"
        ws["A1"] = build_let_formula({"a": "10", "b": "20"}, "a + b")

        out_path = os.path.join(self.tmp_dir, "test_roundtrip.xlsx")
        self.wb.save(out_path)

        wb_loaded = openpyxl.load_workbook(out_path)
        self.assertIn("V_PRISM", wb_loaded.defined_names)
        self.assertIn("REBAR_WEIGHT", wb_loaded.defined_names)
        self.assertEqual(wb_loaded["TEST_SHEET"]["A1"].value, "=_xlfn.LET(_xlpm.a, 10, _xlpm.b, 20, _xlpm.a + _xlpm.b)")


if __name__ == "__main__":
    unittest.main()
