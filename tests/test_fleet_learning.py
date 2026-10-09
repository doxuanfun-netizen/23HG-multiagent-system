import os
import tempfile
import unittest

import openpyxl

from tools import fleet_learning as fl


def _make_fleet(path):
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "01_TienDo_CaMay_Master"
    rows = [  # wbs, tên, đvt, năng suất/ca, ca/ngày, máy, NC
        ("WBS 1.1", "Đổ bê tông bệ móng", "m3", 20.0, 2, "Xe bơm cần 42m + 3 Xe bồn 9m3", 16),
        ("WBS 1.2", "Đổ bê tông thân", "m3", 30.0, 2, "Xe bơm cần 42m + 3 Xe bồn 9m3", 18),
        ("WBS 1.3", "Đổ bê tông xà mũ", "m3", 10.0, 2, "Xe bơm cần 42m + 3 Xe bồn 9m3", 14),
        ("WBS 2.1", "Khoan cọc nhồi D1200", "m", 8.0, 2, "Máy khoan Bauer", 10),
        ("WBS 3.1", "Đào hố móng", "m3", 100.0, 1, "Máy xúc PC200 + Ô tô Howo 15T", 8),
    ]
    for i, (w, n, u, p, s, m, c) in enumerate(rows):
        r = 8 + i
        ws1.cell(r, 2, w); ws1.cell(r, 3, n); ws1.cell(r, 4, u); ws1.cell(r, 6, p)
        ws1.cell(r, 12, s); ws1.cell(r, 14, m); ws1.cell(r, 15, c)
    ws2 = wb.create_sheet("02_TongHop_CaXe_CaMay_MMTB")
    for i, (code, name, u, lit) in enumerate([("BM", "Xe bơm BT", "xe", 46), ("XB", "Xe bồn BT", "xe", 42),
                                              ("MD", "Máy đào", "máy", 68), ("OT", "Ô tô tự đổ", "xe", 52),
                                              ("MK", "Máy khoan cọc nhồi", "máy", 145)]):
        r = 4 + i
        ws2.cell(r, 2, code); ws2.cell(r, 3, name); ws2.cell(r, 4, u); ws2.cell(r, 5, lit)
    wb.save(path)


def _task(wbs, name, qty, unit, days=1):
    return {"wbs": wbs, "name": name, "qty": qty, "unit": unit, "days": days, "start": None, "finish": None}


def _lib():
    return {"status": "PENDING_APPROVAL",
            "source": {"project": "ĐT mẫu", "file": "dt.xlsm", "sha256_16": "x", "sheets": ["Don gia XD"]},
            "tasks": [
                {"code": "AF.12110", "name": "Bê tông tường C30", "unit": "m³", "machines": {"M1": 0.095, "M2": 0.18}, "labor": 2.0},
                {"code": "AF.12110", "name": "Bê tông tường C16", "unit": "m³", "machines": {"M1": 0.095, "M2": 0.18}, "labor": 2.0},
                {"code": "AF.86211", "name": "Ván khuôn thép bt tường", "unit": "100m²", "machines": {"M3": 1.5}, "labor": 28.0},
                {"code": "AB.31133", "name": "Đào đất", "unit": "100m³", "machines": {"M4": 0.3}, "labor": 0.5},
            ],
            "machines": {"M1": {"name": "Máy trộn", "fuel_per_ca": None, "fuel": "kwh", "crew": None},
                         "M2": {"name": "Đầm dùi", "fuel_per_ca": None, "fuel": "kwh", "crew": None},
                         "M3": {"name": "Máy hàn", "fuel_per_ca": None, "fuel": "kwh", "crew": None},
                         "M4": {"name": "Máy đào", "fuel_per_ca": 83.0, "fuel": "diezel", "crew": "1x3/7"}}}


class FleetLearningTest(unittest.TestCase):
    def test_learn_and_infer_project_norms(self):
        with tempfile.TemporaryDirectory() as d:
            p = os.path.join(d, "fleet.xlsx")
            _make_fleet(p)
            k = fl.build_knowledge([fl.learn_from_fleet_workbook(p, "DA mẫu")])
            self.assertEqual(k["status"], "PENDING_APPROVAL")
            bt = k["norms"]["BE_TONG_KET_CAU"]
            self.assertEqual((bt["n"], bt["prod_median"], bt["machines"]), (3, 20.0, {"BM": 1, "XB": 3}))
            rows = fl.infer([_task("1", "Đổ bê tông thân cống", 70.0, "m3", 5),
                             _task("2", "Khoan cọc nhồi", 10, "m"),
                             _task("3", "Lắp dựng ván khuôn", 500, "m2"),
                             _task("4", "Đổ bê tông gối", 5, "cái")], k)
            self.assertEqual(rows[0]["norm"]["source"], "BẢNG CA MÁY DỰ ÁN")
            self.assertAlmostEqual(rows[0]["norm"]["machines"]["XB"], 3 / 20.0)
            self.assertIn("đặc thù cầu", rows[1]["status"])
            self.assertIsNone(rows[2]["norm"])      # chưa học => không bịa máy
            self.assertIsNone(rows[3]["norm"])      # khác ĐVT

    def test_library_takes_priority_and_units_scale(self):
        k, lib = None, _lib()
        rows = fl.infer([_task("1", "Đào bóc nền", 300.0, "m3"),
                         _task("2", "Đổ bê tông thân cống", 70.0, "m3"),
                         _task("3", "Lắp dựng ván khuôn", 500.0, "m2"),
                         _task("4", "Gia công, lắp dựng cốt thép", 70.0, "m3 BT"),
                         _task("5", "Bảo dưỡng, tháo ván khuôn", 7, "ngày")], k, lib)
        self.assertEqual(rows[0]["norm"]["machines"], {"M4": 0.3})   # không có "đào bóc" => dự phòng "Đào đất", ghi rõ ở ref
        self.assertIn("Đào đất", rows[0]["norm"]["ref"])
        self.assertEqual(rows[1]["norm"]["source"], "ĐM DỰ TOÁN")
        self.assertEqual(rows[1]["norm"]["machines"], {"M1": 0.095, "M2": 0.18})
        self.assertEqual(rows[2]["norm"]["scale"], 100.0)
        self.assertIn("cốt thép", rows[3]["status"].lower())
        self.assertIsNone(rows[4]["norm"])
        with tempfile.TemporaryDirectory() as d:
            out = fl.write_fleet_workbook(os.path.join(d, "o.xlsx"), "Cống thử", rows, k, lib)
            wb = openpyxl.load_workbook(out)
            self.assertEqual(wb.sheetnames, ["01_TienDo_CaMay_Master", "02_TongHop_CaXe_CaMay_MMTB",
                                             "03_KeHoach_Dau_Diezel", "04_KeHoach_NhanLuc", "05_DoiChieu_BocTach"])
            ws = wb["01_TienDo_CaMay_Master"]
            self.assertEqual(ws["G7"].value, "=IF(F7>0,ROUND(E7/F7,2),0)")
            self.assertEqual(ws["M7"].value, "=IF(AND(G7>0,I7>0,L7>0),ROUNDUP(G7/(I7*L7),1),0)")
            text = " ".join(str(c.value) for s in wb.worksheets for row in s.iter_rows() for c in row if c.value)
            self.assertIn("SUY LUẬN", text)

    def test_split_unit(self):
        self.assertEqual(fl.split_unit("100m³"), (100.0, "m³"))
        self.assertEqual(fl.split_unit("tấn"), (1.0, "tấn"))
        self.assertEqual(fl.split_unit("m³"), (1.0, "m³"))

    def test_vincons_learning_and_matching(self):
        # Kiểm tra nạp tri thức Vincons và ưu tiên định mức nội bộ nhà thầu
        vc = fl.load_vincons()
        if not vc:
            self.skipTest("chưa có knowledge/dinh_muc_vincons.json")
        self.assertEqual(vc["status"], "PENDING_APPROVAL")
        self.assertIn("MHT", vc["source"]["sheets"])
        self.assertIn("VC", vc["source"]["sheets"])
        self.assertIn("May", vc["source"]["sheets"])
        self.assertGreaterEqual(len(vc["works"]), 180)
        self.assertIn("transport", vc)
        self.assertIn("organization", vc)
        self.assertEqual(vc["organization"]["driver_per_machine_per_shift"], 1.0)

        # Suy luận với cả 3 nguồn
        k, lib = fl.load_knowledge(), fl.load_library()
        tasks = [
            _task("1.1", "Đào bóc 30 cm nền hiện trạng", 100.0, "m3"),
            _task("1.2", "Đệm cát đầm chặt K0.9", 50.0, "m3"),
            _task("1.3", "Đổ bê tông thân cống", 70.0, "m3"),
            _task("1.4", "Đắp cát hai bên, hoàn trả", 30.0, "m3"),
        ]
        rows = fl.infer(tasks, k, lib, vc)
        # Đào bóc -> Vincons MHT (M1 máy đào + M3 máy ủi)
        self.assertEqual(rows[0]["norm"]["source"], "ĐM NỘI BỘ VINCONS (MHT)")
        self.assertIn("VC.M1", rows[0]["norm"]["machines"])
        self.assertIn("VC.M3", rows[0]["norm"]["machines"])
        # Đệm cát -> Vincons MHT (M3 máy ủi + M5 lu rung)
        self.assertEqual(rows[1]["norm"]["source"], "ĐM NỘI BỘ VINCONS (MHT)")
        self.assertIn("VC.M3", rows[1]["norm"]["machines"])
        self.assertIn("VC.M5", rows[1]["norm"]["machines"])
        # Đổ bê tông thân cống -> Library ĐM dự toán (máy trộn + đầm dùi)
        if lib:
            self.assertEqual(rows[2]["norm"]["source"], "ĐM DỰ TOÁN")
        # Đắp hoàn trả -> Vincons MHT (M1 máy đào)
        self.assertEqual(rows[3]["norm"]["source"], "ĐM NỘI BỘ VINCONS (MHT)")
        self.assertIn("VC.M1", rows[3]["norm"]["machines"])


if __name__ == "__main__":
    unittest.main()

