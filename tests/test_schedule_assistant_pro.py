# -*- coding: utf-8 -*-
"""
UNIT TESTS — 23HG SCHEDULE ASSISTANT PRO INTEGRATION
Kiểm thử tích hợp phân hệ Sổ tính Master V3, Add-in Ribbon và Động cơ CPM chuyên nghiệp.
"""

import os
import sys
import unittest
import zipfile
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
APP_DIR = os.path.join(ROOT, "apps", "23hg_schedule_assistant_pro")
SCRIPTS_DIR = os.path.join(APP_DIR, "1_Scripts_TuDongHoa")
if SCRIPTS_DIR not in sys.path:
    sys.path.insert(0, SCRIPTS_DIR)


class TestScheduleAssistantProIntegration(unittest.TestCase):

    def setUp(self):
        self.master_xlsm = os.path.join(APP_DIR, "23HG_DU_AN_MAU_TIEN_DO_CHUAN_G1_PRO.xlsm")
        self.addin_xlam = os.path.join(APP_DIR, "23HG_Schedule_Assistant_Pro.xlam")

    def test_master_xlsm_structure_and_sheets(self):
        """Kiểm tra sự tồn tại và 9 sheet nghiệp vụ chuyên sâu của Master XLSM."""
        self.assertTrue(os.path.exists(self.master_xlsm), "File Master XLSM phải tồn tại")
        wb = openpyxl.load_workbook(self.master_xlsm, read_only=True)
        expected_sheets = [
            "NGAY_NGHI_LE",
            "DB_DINH_MUC",
            "BOQ_TIEN_DO",
            "TIEN_DO",
            "EVM_5D_QUAN_TRI",
            "HUY_DONG_XMTB",
            "TIEN_DO_GIAI_NGAN",
            "KE_HOACH_QLCL",
            "HUONG_DAN",
        ]
        for name in expected_sheets:
            self.assertIn(name, wb.sheetnames, f"Sheet {name} phải có trong Master")

    def test_addin_openxml_and_ribbon(self):
        """Kiểm tra Add-in XLAM chứa đúng Ribbon XML và mã macro."""
        self.assertTrue(os.path.exists(self.addin_xlam), "File Add-in XLAM phải tồn tại")
        with zipfile.ZipFile(self.addin_xlam, "r") as zf:
            namelist = zf.namelist()
            self.assertIn("customUI/customUI14.xml", namelist, "Phải chứa customUI14.xml")
            self.assertIn("xl/vbaProject.bin", namelist, "Phải chứa vbaProject.bin")
            xml_content = zf.read("customUI/customUI14.xml").decode("utf-8")
            self.assertIn("23HG SCHEDULE ASSISTANT", xml_content)
            self.assertIn('id="grp_export"', xml_content)
            self.assertIn('onAction="XuatXlsxClean"', xml_content)
            self.assertIn('onAction="XuatPdfA3BaoCao"', xml_content)

    def test_cpm_engine_pure_python_matches_37_tasks_with_rain(self):
        """Kiểm tra động cơ CPM Python đối soát khớp 37/37 công tác khi xét hệ số mưa K_tt."""
        import cpm_compare as cc

        bad, n = cc.compare(self.master_xlsm, use_rain=True)
        self.assertEqual(n, 37, "Phải có đúng 37 công tác trong bảng tiến độ")
        self.assertEqual(len(bad), 0, "Không được có ngày lệch khi áp dụng hệ số mưa K_tt")

    def test_installer_and_documentation_presence(self):
        """Kiểm tra các tệp cài đặt và tài liệu hồ sơ bàn giao nghiệm thu."""
        required_files = [
            "CAI_DAT_23HG_ENTERPRISE_PRO.bat",
            "GO_CAI_DAT_23HG.bat",
            "README.md",
            "CHANGELOG.md",
            "BAN_GIAO_DANH_GIA_DU_AN.md",
        ]
        for fn in required_files:
            fp = os.path.join(APP_DIR, fn)
            self.assertTrue(os.path.exists(fp), f"Tệp {fn} phải tồn tại trong apps/23hg_schedule_assistant_pro")

    def test_bridge_km19_showcase_model_and_xml(self):
        """Kiểm tra dự án mẫu Cầu Km19+529.080 đồng bộ trong phân hệ 23HG Schedule Assistant Pro."""
        bridge_xlsx = os.path.join(APP_DIR, "2_BaoCao_XuatBan", "23HG_DU_AN_MAU_CAU_KM19_529_PRO.xlsx")
        bridge_xml = os.path.join(APP_DIR, "2_BaoCao_XuatBan", "Du_An_Mau_Cau_Km19_529.xml")
        bridge_master = os.path.join(ROOT, "examples", "HO_SO_CAU_KM19_529", "00_BAN_CHI_HUY_MASTER", "23HG_MASTER_TIEN_DO_EVM_CAU_KM19_PRO.xlsx")

        self.assertTrue(os.path.exists(bridge_xlsx), "File dự án mẫu Cầu Km19 phải tồn tại trong 2_BaoCao_XuatBan")
        self.assertTrue(os.path.exists(bridge_xml), "File MSPDI XML Cầu Km19 phải tồn tại trong 2_BaoCao_XuatBan")
        self.assertTrue(os.path.exists(bridge_master), "File Master Tiến độ EVM Cầu Km19 phải tồn tại trong 00_BAN_CHI_HUY_MASTER")

        # Kiểm tra 9 sheet chuẩn doanh nghiệp
        wb = openpyxl.load_workbook(bridge_xlsx, read_only=True)
        expected_sheets = [
            "NGAY_NGHI_LE",
            "DB_DINH_MUC",
            "BOQ_TIEN_DO",
            "TIEN_DO",
            "EVM_5D_QUAN_TRI",
            "HUY_DONG_XMTB",
            "TIEN_DO_GIAI_NGAN",
            "KE_HOACH_QLCL",
            "HUONG_DAN",
        ]
        for name in expected_sheets:
            self.assertIn(name, wb.sheetnames, f"Sheet {name} phải có trong dự án mẫu Cầu Km19")

        # Kiểm tra cấu trúc XML MSPDI
        import xml.etree.ElementTree as ET
        tree = ET.parse(bridge_xml)
        root = tree.getroot()
        tasks = [elem for elem in root.iter() if elem.tag.endswith("Task")]
        self.assertEqual(len(tasks), 37, "File XML phải chứa đúng 37 công tác tiến độ")


if __name__ == "__main__":
    unittest.main()
