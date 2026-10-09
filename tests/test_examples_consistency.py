# -*- coding: utf-8 -*-
"""Bảo vệ thư mục hồ sơ mẫu: không file rỗng, bản sao không lệch nhau, gói không có đơn giá thì không chứa sheet giá,
và các số đối chiếu giữa bóc tách / hợp đồng / đắp lưng cống nằm trong giới hạn đã biết."""

import hashlib
import os
import sys
import unittest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

import openpyxl

from tools.audit_excels_static import audit_file, find_xlsx
from tools.excel_eval import WorkbookEvaluator

A5 = os.path.join(ROOT, "examples", "HO_SO_CONG_HOP_TUYEN_A5")
HUB = os.path.join(A5, "HO_SO_THUC_CHIEN_HUB_AND_SPOKE_CONG_A5")
MASTER = os.path.join(A5, "BO_HO_SO_01_MACRO_MASTER_14_SHEET", "01_Ho_So_KCS_QS_TienDo_Master_14_Sheets_Cong_Hop_A5.xlsx")
PRICE_SHEET_HINTS = ("GXD", "03A", "THANH_TOAN", "DU_TOAN")


def _md5(p):
    with open(p, "rb") as f:
        return hashlib.md5(f.read()).hexdigest()


class NoPersonalPathsTest(unittest.TestCase):
    """Không ghi đường dẫn máy cá nhân vào repo (dùng examples/_paths.py và biến môi trường AEC_PROJECTS_DIR)."""

    def test_no_personal_machine_paths(self):
        import re
        pat = re.compile(r"[A-Za-z]:\\Users\\(?!<)[^\\\s\"']+\\|[A-Za-z]:\\Code\\\w", re.I)
        hits = []
        for dp, dirs, files in os.walk(ROOT):
            dirs[:] = [d for d in dirs if d not in (".git", "__pycache__", "du_lieu_du_an")]
            for fn in files:
                if fn.endswith((".py", ".md", ".txt", ".json", ".toml", ".yml", ".cfg")):
                    path = os.path.join(dp, fn)
                    with open(path, encoding="utf-8", errors="ignore") as f:
                        for no, line in enumerate(f, 1):
                            if pat.search(line):
                                hits.append(f"{os.path.relpath(path, ROOT)}:{no}")
        self.assertEqual(hits, [], "đường dẫn máy cá nhân trong repo")


# File dữ liệu xuất từ bộ giải cắt thép (RebarCut: tools/rebarcut_export.py không ghi công thức nào, các file có
# 0 công thức) không cần quét công thức; quét các file lớn mất tới ~40 giây. Bỏ qua file > 1 MB và cả thư mục kết quả.
_HEAVY_BYTES = 1024 * 1024
_SKIP_DIR_NAMES = {
    "01_HE_THONG_CAT_THEP_REBARCUT",
    "01_HIEN_TRUONG_QLCL_KCS",
    "02_XUONG_TIEN_CHE_COT_THEP",
    "05_DU_LIEU_GOC_SCAN_MARKER",
}


def _skip_audit(path):
    parts = set(path.replace("\\", "/").split("/"))
    return os.path.getsize(path) > _HEAVY_BYTES or bool(parts & _SKIP_DIR_NAMES)


class FormulaCoverageTest(unittest.TestCase):
    """Mọi ô công thức trong hồ sơ mẫu phải tính được bằng tools/excel_eval và không ra lỗi Excel."""

    def test_all_sample_formulas_evaluate_without_errors(self):
        problems = []
        by_content = {}      # các bản sao giống hệt nhau chỉ cần tính một lần
        for folder in ("examples", "templates"):
            for p in find_xlsx(os.path.join(ROOT, folder)):
                if _skip_audit(p):
                    continue
                key = _md5(p)
                if key not in by_content:
                    by_content[key] = audit_file(p)
                r = by_content[key]
                if r["eval_unsupported"] or r["eval_error_results"]:
                    problems.append(f"{os.path.relpath(p, ROOT)}: chưa hỗ trợ {r['eval_unsupported']}, "
                                    f"lỗi {r['eval_error_results']}")
        self.assertEqual(problems, [])


class ReadmeLinksTest(unittest.TestCase):

    def test_relative_links_point_to_existing_files(self):
        import re
        with open(os.path.join(ROOT, "README.md"), encoding="utf-8") as f:
            text = f.read()
        links = [l for l in re.findall(r"\]\(([^)#\s]+)\)", text) if not l.startswith(("http://", "https://"))]
        missing = [l for l in links if not os.path.exists(os.path.join(ROOT, l))]
        self.assertTrue(links)
        self.assertEqual(missing, [], "README trỏ tới file không tồn tại")


class ExamplesLayoutTest(unittest.TestCase):
    """Hồ sơ mẫu nằm trong thư mục dự án của nó, không để file rời ở gốc examples/ hay lẫn sang dự án khác."""
    EXAMPLES = os.path.join(ROOT, "examples")
    DATA_EXT = (".xlsx", ".xml", ".mpp", ".docx", ".csv")

    def test_no_loose_dossier_files_at_examples_root(self):
        loose = [f for f in os.listdir(self.EXAMPLES) if f.lower().endswith(self.DATA_EXT)]
        self.assertEqual(loose, [], "file hồ sơ nằm rời ở gốc examples/")

    def test_a5_files_only_in_a5_folder(self):
        stray = []
        for dp, _, files in os.walk(self.EXAMPLES):
            if os.path.abspath(dp).startswith(os.path.abspath(A5)):
                continue
            stray += [os.path.join(dp, f) for f in files if "_A5" in f and f.lower().endswith(self.DATA_EXT)]
        self.assertEqual(stray, [], "file của cống A5 nằm ngoài thư mục A5")


@unittest.skipUnless(os.path.isdir(A5), "thiếu thư mục ví dụ A5")
class A5ExamplesTest(unittest.TestCase):

    def test_no_empty_shell_workbooks(self):
        empties = [p for p in find_xlsx(os.path.join(ROOT, "examples"))
                   if not _skip_audit(p) and audit_file(p, try_eval=False)["empty_shell"]]
        self.assertEqual(empties, [], "file Excel chỉ có tiêu đề, không dữ liệu")

    def test_same_named_copies_are_identical(self):
        by_name = {}
        for p in find_xlsx(A5):
            by_name.setdefault(os.path.basename(p), []).append(p)
        for name, paths in by_name.items():
            if len(paths) > 1:
                self.assertEqual(len({_md5(p) for p in paths}), 1, f"bản sao của {name} đã lệch nhau")

    def test_goi_without_prices_have_no_price_sheets(self):
        for pkg in ("GOI_A_CO_GIOI_VA_DAU_DIEZEL", "GOI_B_XUONG_TIEN_CHE_COT_THEP", "GOI_C_HIEN_TRUONG_QLCL_KCS"):
            for p in find_xlsx(os.path.join(HUB, pkg)):
                for name in openpyxl.load_workbook(p).sheetnames:
                    self.assertFalse(any(h in name.upper() for h in PRICE_SHEET_HINTS),
                                     f"{pkg}/{os.path.basename(p)} chứa sheet giá '{name}'")

    def test_master_inputs_are_labelled_not_fake_formulas(self):
        ws = openpyxl.load_workbook(MASTER)["TONG_HOP_DU_TOAN_GXD"]
        self.assertEqual(ws["D5"].value, 42500000000)                      # số nhập, không phải "=42500000000"
        self.assertIn("ĐẦU VÀO", ws["F5"].value)
        for formula_cell in ("D6", "D7", "D9"):
            self.assertNotRegex(ws[formula_cell].value, r"\*0\.\d")         # tỷ lệ nằm ở ô đầu vào, không trong công thức

    def test_backfill_by_culvert_type_matches_hand_values(self):
        ev = WorkbookEvaluator(MASTER)
        s = "KHOI_LUONG_DAO_DAP"
        # (hào − cống bao ngoài) × số đốt × dài: 34×11.3×3.75 ; 50×11.3×5.4 ; 108×11.3×5.4
        self.assertAlmostEqual(ev.value(s, 15, 9), 1440.75, places=3)
        self.assertAlmostEqual(ev.value(s, 16, 9), 3051.0, places=3)
        self.assertAlmostEqual(ev.value(s, 17, 9), 6590.16, places=3)
        self.assertAlmostEqual(ev.value(s, 18, 9), 11081.91, places=3)
        self.assertAlmostEqual(ev.value(s, 19, 9), 20804.814, places=3)    # số cũ giữ nguyên để đối chiếu

    def test_contract_vs_takeoff_reconciliation(self):
        ev = WorkbookEvaluator(MASTER)
        t = "THANH_TOAN_KY_PHU_LUC_03A"
        for row in range(5, 10):                                           # đào, cọc tre, BT lót, BT thân, ván khuôn
            self.assertLessEqual(abs(ev.value(t, row, 11)), 0.005, f"dòng {row}")
        # cốt thép: hợp đồng 753.4 tấn lệch +14.4% so với bảng thống kê thép — sai khác ĐÃ BIẾT, chờ QS xử lý
        self.assertGreater(ev.value(t, 10, 11), 0.10)


if __name__ == "__main__":
    unittest.main()
