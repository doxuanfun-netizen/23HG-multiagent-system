# -*- coding: utf-8 -*-
"""Kiểm thử đơn vị cho examples/build_khai_hoang_2_full_pipeline.py, tools/generate_khai_hoang_2_elite_qlcl.py
và tính toán động tỷ lệ trùng lặp trong tools/web_portal_1click.py."""

import os
import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET
from unittest.mock import patch

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

import openpyxl

from examples.build_khai_hoang_2_full_pipeline import (
    WBS_TASKS,
    clean_dia,
    clean_length_mm,
    clean_num,
    create_master_rebar_dashboard,
    create_ms_project_xml,
    generate_kcs_word_dossier,
    generate_machine_diesel_workbook,
)
from tools.audit_excels_static import audit_file
from tools.excel_eval import WorkbookEvaluator
from tools.generate_khai_hoang_2_elite_qlcl import (
    build_tan_suat_khai_hoang_2,
    build_visual_asbuilt_khai_hoang_2,
)
from tools.web_portal_1click import DREAM_TEAM_6_FOLDERS, get_portal_state


class KhaiHoang2FullPipelineTest(unittest.TestCase):
    """Kiểm thử các bộ làm sạch OCR, bảng tổng hợp cắt thép, ca máy, XML và KCS Word."""

    def test_ocr_cleaners_handle_vn_and_corrupted_formats(self):
        self.assertAlmostEqual(clean_num("7,900,00"), 7900.0)
        self.assertAlmostEqual(clean_num("1,153,00"), 1153.0)
        self.assertAlmostEqual(clean_num("2.260.00"), 2260.0)
        self.assertAlmostEqual(clean_length_mm("11.700"), 11700.0)
        self.assertEqual(clean_dia("Φ32"), 32)
        self.assertEqual(clean_dia("Ø16"), 16)
        self.assertIsNone(clean_dia("D99"))

    def test_master_rebar_dashboard_computes_status_and_live_formulas(self):
        rows = [
            {
                "dia": 16,
                "desc": "Thép sườn dầm",
                "station": "Trạm 1",
                "bars_qty": 100,
                "total_len_m": 500.0,
                "bars_11m7": 45,
                "weight_kg": 830.8,
                "waste_m": 26.5,
                "waste_pct": 5.03,
                "patterns_count": 4,
                "reusable_m": 12.0,
                "status": "OPTIMAL",
            },
            {
                "dia": 25,
                "desc": "Thép chủ mố",
                "station": "Trạm 2",
                "bars_qty": 60,
                "total_len_m": 320.0,
                "bars_11m7": 30,
                "weight_kg": 1352.4,
                "waste_m": 31.0,
                "waste_pct": 8.83,
                "patterns_count": 3,
                "reusable_m": 8.5,
                "status": "FEASIBLE",
            },
        ]
        with tempfile.TemporaryDirectory() as tmp:
            out_xlsx = os.path.join(tmp, "rebar_summary.xlsx")
            create_master_rebar_dashboard(out_xlsx, rows)

            wb = openpyxl.load_workbook(out_xlsx, data_only=False)
            ws = wb["TONG_HOP_CAT_THEP_11M7"]
            # Dòng tổng cộng là dòng 9 (2 dòng dữ liệu 7, 8)
            self.assertEqual(ws["E9"].value, "=SUM(E7:E8)")
            self.assertEqual(ws["G9"].value, "=SUM(G7:G8)")
            self.assertEqual(ws["N9"].value, "FEASIBLE (1/2 OPTIMAL)")

            ev = WorkbookEvaluator(out_xlsx)
            self.assertEqual(ev.value("TONG_HOP_CAT_THEP_11M7", 9, 5), 160)
            self.assertEqual(ev.value("TONG_HOP_CAT_THEP_11M7", 9, 7), 75)
            rep = audit_file(out_xlsx)
            self.assertEqual(rep["eval_error_results"], 0)
            self.assertEqual(rep["eval_unsupported"], 0)

    def test_machine_diesel_xml_and_word_generators(self):
        with tempfile.TemporaryDirectory() as tmp:
            diesel_xlsx = os.path.join(tmp, "diesel.xlsx")
            generate_machine_diesel_workbook(diesel_xlsx)
            rep = audit_file(diesel_xlsx)
            self.assertFalse(rep["empty_shell"])
            self.assertEqual(rep["eval_error_results"], 0)

            xml_out = os.path.join(tmp, "schedule.xml")
            create_ms_project_xml(WBS_TASKS[:4], xml_out)
            tree = ET.parse(xml_out)
            self.assertIsNotNone(tree.getroot())

            docx_out = os.path.join(tmp, "kcs.docx")
            generate_kcs_word_dossier(docx_out)
            self.assertGreater(os.path.getsize(docx_out), 5000)


class KhaiHoang2EliteQlclTest(unittest.TestCase):
    """Kiểm thử bộ tạo Bảng tần suất thí nghiệm 4 sheet và Sơ đồ tiến độ trực quan As-Built."""

    def test_build_tan_suat_and_visual_asbuilt_workbooks(self):
        sample_bbs = [
            {
                "comp": "Bệ mố - bảng 1 (tr.14)",
                "mark": "M1",
                "dia": 25,
                "length_mm": 4200,
                "qty_per_unit": 40,
                "units": 1,
                "total_qty": 40,
                "kg_total": 647.3,
                "sheet": "PD-04",
                "page": 14,
            },
            {
                "comp": "Dầm T L=15m - dầm biên (1 dầm)",
                "mark": "A1",
                "dia": 32,
                "length_mm": 14900,
                "qty_per_unit": 10,
                "units": 2,
                "total_qty": 20,
                "kg_total": 1881.3,
                "sheet": "KC-05",
                "page": 22,
            },
            {
                "comp": "Neo mố vào đá (toàn cầu)",
                "mark": "N1",
                "dia": 32,
                "length_mm": 3000,
                "qty_per_unit": 98,
                "units": 1,
                "total_qty": 98,
                "kg_total": 1856.0,
                "sheet": "PD-01",
                "page": 11,
            },
        ]
        with tempfile.TemporaryDirectory() as tmp:
            tan_suat_xlsx = os.path.join(tmp, "tan_suat.xlsx")
            build_tan_suat_khai_hoang_2(tan_suat_xlsx, bbs_rows=sample_bbs)
            wb_ts = openpyxl.load_workbook(tan_suat_xlsx, data_only=False)
            self.assertEqual(
                wb_ts.sheetnames,
                [
                    "Theo_Doi_Tan_Suat_Be_Tong",
                    "Theo_Doi_Tan_Suat_Cot_Thep",
                    "Tong_Hop_Tan_Suat_Vat_Lieu",
                    "NGHI_VAN",
                ],
            )
            rep_ts = audit_file(tan_suat_xlsx)
            self.assertEqual(rep_ts["eval_error_results"], 0)
            self.assertEqual(rep_ts["eval_unsupported"], 0)

            asbuilt_xlsx = os.path.join(tmp, "asbuilt.xlsx")
            asbuilt_html = os.path.join(tmp, "asbuilt.html")
            build_visual_asbuilt_khai_hoang_2(asbuilt_xlsx, asbuilt_html)
            rep_ab = audit_file(asbuilt_xlsx)
            self.assertEqual(rep_ab["eval_error_results"], 0)
            self.assertEqual(rep_ab["eval_unsupported"], 0)
            self.assertTrue(os.path.isfile(asbuilt_html))


class WebPortalDuplicationRateTest(unittest.TestCase):
    """Đảm bảo duplication_rate trong tools/web_portal_1click.py được tính thực tế từ cây thư mục."""

    def test_duplication_rate_reflects_actual_duplicate_filenames(self):
        with tempfile.TemporaryDirectory() as tmp:
            f0 = os.path.join(tmp, DREAM_TEAM_6_FOLDERS[0])
            f1 = os.path.join(tmp, DREAM_TEAM_6_FOLDERS[1])
            os.makedirs(f0)
            os.makedirs(f1)
            # Tạo 2 file trùng tên ở 2 phân hệ và 2 file duy nhất -> tổng 4 file, 1 tên bị lặp (dup_count=1 -> 25.0%)
            for p in (
                os.path.join(f0, "dup.xlsx"),
                os.path.join(f0, "unique1.xlsx"),
                os.path.join(f1, "dup.xlsx"),
                os.path.join(f1, "unique2.xlsx"),
            ):
                with open(p, "wb") as fh:
                    fh.write(b"x")

            fake_catalog = {
                "khai_hoang_2": {
                    "id": "khai_hoang_2",
                    "name": "Test",
                    "category": "Test",
                    "path": tmp,
                    "pipeline_script": "",
                    "highlights": {
                        "excel_count": 4,
                        "live_formulas": 10,
                        "formula_errors": 0,
                        "bbs_rows": 10,
                        "kcs_reports": 1,
                    },
                }
            }
            with patch("tools.web_portal_1click.get_projects_catalog", return_value=fake_catalog):
                state = get_portal_state("khai_hoang_2")
                self.assertEqual(state["active_project"]["duplication_rate"], "25.0%")


if __name__ == "__main__":
    unittest.main()
