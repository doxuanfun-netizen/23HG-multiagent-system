# -*- coding: utf-8 -*-
"""Test tools/office_scrub.py: xóa absPath và liên kết ngoài thừa mà không đổi nội dung bảng tính."""

import os
import tempfile
import unittest
import zipfile

import openpyxl

from tools.office_scrub import ScrubError, inspect, scrub

ABS = ('<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
       '<mc:Choice Requires="x15"><x15ac:absPath url="C:\\Users\\ai_do\\Downloads\\Du_an\\" '
       'xmlns:x15ac="http://schemas.microsoft.com/office/spreadsheetml/2010/11/ac"/></mc:Choice>'
       '</mc:AlternateContent>')
EXT_XML = ('<externalLink xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
           '<externalBook xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="rId1"/>'
           '</externalLink>')
EXT_RELS = ('<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
            'externalLinkPath" Target="file:///D:\\May_khac\\Mau.xlsx" TargetMode="External"/></Relationships>')


def _rewrite(path, edit):
    with zipfile.ZipFile(path) as z:
        items = [(i, z.read(i.filename)) for i in z.infolist()]
    parts = {i.filename: d for i, d in items}
    extra = edit(parts) or {}
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as out:
        for i, _ in items:
            if i.filename in parts:
                out.writestr(i.filename, parts[i.filename])
        for name, data in extra.items():
            out.writestr(name, data)


def make_workbook(path, formula_uses_link=False):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["B1"] = "Khối lượng", 12.5
    ws["B2"] = "=[1]Sheet1!A1*2" if formula_uses_link else "=B1*2"
    wb.save(path)

    def edit(parts):
        w = parts["xl/workbook.xml"].decode()
        w = w.replace("<workbookPr/>", "<workbookPr/>" + ABS, 1)
        w = w.replace("</sheets>", '</sheets><externalReferences><externalReference r:id="rId99"/>'
                                   '</externalReferences><definedNames><definedName name="Ngoai">'
                                   "'[1]Sheet1'!$A$1</definedName><definedName name=\"TrongFile\">"
                                   "Sheet!$B$1</definedName></definedNames>", 1)
        if "xmlns:r=" not in w.split(">", 2)[1]:
            w = w.replace("<workbook ", '<workbook xmlns:r="http://schemas.openxmlformats.org/officeDocument/'
                                       '2006/relationships" ', 1)
        parts["xl/workbook.xml"] = w.encode()
        r = parts["xl/_rels/workbook.xml.rels"].decode()
        parts["xl/_rels/workbook.xml.rels"] = r.replace(
            "</Relationships>", '<Relationship Id="rId99" Type="http://schemas.openxmlformats.org/officeDocument/'
                                '2006/relationships/externalLink" Target="externalLinks/externalLink1.xml"/>'
                                "</Relationships>").encode()
        ct = parts["[Content_Types].xml"].decode()
        parts["[Content_Types].xml"] = ct.replace(
            "</Types>", '<Override PartName="/xl/externalLinks/externalLink1.xml" ContentType="application/'
                        'vnd.openxmlformats-officedocument.spreadsheetml.externalLink+xml"/></Types>').encode()
        return {"xl/externalLinks/externalLink1.xml": EXT_XML,
                "xl/externalLinks/_rels/externalLink1.xml.rels": EXT_RELS}
    _rewrite(path, edit)


class OfficeScrubTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = os.path.join(self.tmp.name, "mau.xlsx")

    def tearDown(self):
        self.tmp.cleanup()

    def test_removes_abs_path_and_unused_external_link_keeping_cells(self):
        make_workbook(self.path)
        before = inspect(self.path)
        self.assertTrue(before["absPath"])
        self.assertEqual(before["lien_ket_ngoai"], 1)
        self.assertEqual(before["ten_vung_tro_ra_ngoai"], ["Ngoai"])
        report = scrub(self.path, lien_ket_ngoai=True)
        self.assertTrue(report["absPath_da_xoa"])
        self.assertEqual(report["ten_vung_da_xoa"], ["Ngoai"])
        after = inspect(self.path)
        self.assertFalse(after["absPath"])
        self.assertEqual(after["lien_ket_ngoai"], 0)
        with zipfile.ZipFile(self.path) as z:
            self.assertFalse([n for n in z.namelist() if "externalLink" in n])
            text = "".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist())
        self.assertNotIn("ai_do", text)
        self.assertNotIn("May_khac", text)
        self.assertIn("TrongFile", text)                       # tên vùng trong file vẫn giữ
        wb = openpyxl.load_workbook(self.path)
        self.assertEqual(wb.active["B1"].value, 12.5)
        self.assertEqual(wb.active["B2"].value, "=B1*2")

    def test_abs_path_only_by_default_and_idempotent(self):
        make_workbook(self.path)
        self.assertTrue(scrub(self.path)["absPath_da_xoa"])
        self.assertEqual(inspect(self.path)["lien_ket_ngoai"], 1)    # không bật --lien-ket-ngoai thì giữ
        with open(self.path, "rb") as f:
            once = f.read()
        self.assertFalse(scrub(self.path)["absPath_da_xoa"])
        with open(self.path, "rb") as f:
            self.assertEqual(f.read(), once)

    def test_refuses_when_a_formula_uses_the_link(self):
        make_workbook(self.path, formula_uses_link=True)
        with self.assertRaises(ScrubError):
            scrub(self.path, lien_ket_ngoai=True)
        self.assertTrue(inspect(self.path)["absPath"])                # không sửa gì khi từ chối

    def test_same_input_gives_same_bytes(self):
        make_workbook(self.path)
        other = os.path.join(self.tmp.name, "ban_sao.xlsx")
        with open(self.path, "rb") as a, open(other, "wb") as b:
            b.write(a.read())
        scrub(self.path)
        scrub(other)
        with open(self.path, "rb") as a, open(other, "rb") as b:
            self.assertEqual(a.read(), b.read())


if __name__ == "__main__":
    unittest.main()
