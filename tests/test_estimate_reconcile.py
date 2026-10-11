# -*- coding: utf-8 -*-
"""
Test sửa dự toán / 03a theo bản vẽ thẩm định (tools/estimate_reconcile.py, examples/sua_du_toan_03a_km19_theo_ban_ve.py).

Quy tắc chủ dự án (11/10/2026): có dấu thẩm định thì lấy theo thẩm định, lấy theo thiết kế; sai khác thì NGHI VẤN.
"""

import json
import os
import re
import tempfile
import unittest

import openpyxl

from tools.estimate_reconcile import _drawing_number, apply_03a, apply_qs, ocr_value, resolve_f

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, "examples", "data")
CONFIG = os.path.join(DATA, "km19_du_toan_theo_ban_ve.json")
DT_GOC = os.path.join(DATA, "km19_du_toan_goc_truoc_sua.xlsx")
P03A_GOC = os.path.join(DATA, "km19_03a_goc_truoc_sua.xlsx")
OUT_DIR = os.path.join(ROOT, "examples", "HO_SO_CAU_KM19_529", "03_KINH_TE_QS_DU_TOAN_THANH_TOAN")
DT_MOI = os.path.join(OUT_DIR, "08_Du_Toan_GXD_Thong_Tu_11_2021.xlsx")
P03A_MOI = os.path.join(OUT_DIR, "09_Thanh_Toan_Khoi_Luong_Phu_Luc_03a.xlsx")
DU_LIEU = os.path.join(ROOT, "examples", "HO_SO_CAU_KM19_529", "05_DU_LIEU_GOC_SCAN_MARKER", "du_lieu.json")

PAGE = {"page": 175, "tables": [
    {"rows": [["", "Bê tông 45 Mpa", "m3", "434.803", ""],
              ["", "Cốt thép D<=10", "Tấn", "9.000", ""],
              ["", "Bàn mặt cầu, Liên tục nhiệt", "", "", ""],
              ["", "Cốt thepD<=10", "Tấn", "1.573", ""],
              ["", "Cốt thép 10<D<=18", "Tấn", "78.951", ""]]}]}


class DrawingNumbersTest(unittest.TestCase):
    def test_dot_is_decimal_comma_is_thousands(self):
        """Bảng tổng hợp ghi '148.992' m³ — không được hiểu thành 148.992 nghìn như cột 'values' chuẩn hóa VN."""
        self.assertEqual(_drawing_number("148.992"), 148.992)
        self.assertEqual(_drawing_number("1,359.80"), 1359.8)
        self.assertEqual(_drawing_number("15.000"), 15.0)
        self.assertIsNone(_drawing_number("(M3)"))

    def test_label_lookup_ignores_accents_and_spaces_and_respects_anchor(self):
        self.assertEqual(ocr_value(PAGE, "BE TONG 45 MPA"), 434.803)
        self.assertEqual(ocr_value(PAGE, "COT THEP D<=10"), 9.0)                         # lần xuất hiện đầu tiên
        self.assertEqual(ocr_value(PAGE, "COT THEP D<=10", sau="BAN MAT CAU"), 1.573)   # sau mốc bản mặt cầu
        with self.assertRaises(ValueError):
            ocr_value(PAGE, "KHONG CO DONG NAY")                                            # không đoán

    def test_sum_spec_is_written_as_auditable_formula(self):
        spec = {"tong": [{"he_so": 1, "trang": 175, "sau": "BAN MAT CAU", "nhan": "COT THEPD<=10"},
                         {"he_so": 2, "trang": 175, "sau": "BAN MAT CAU", "nhan": "COT THEP 10<D<=18"}]}
        value, cell = resolve_f(spec, {175: PAGE})
        self.assertAlmostEqual(value, 1.573 + 2 * 78.951)
        self.assertEqual(cell, "=1.573+2*78.951")


class ApplyQsTest(unittest.TestCase):
    def setUp(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "QS"
        ws["A7"], ws["C7"], ws["J7"] = "1", "Bê tông dầm", 999.0               # tiêu đề ghi cứng
        for c, v in zip("CDEFGHIJ", ["- dầm", "m3", 16, 38.2, 0.759, 1, 1, "=E8*F8*G8*H8*I8"]):
            ws[f"{c}8"] = v
        self.ws = ws
        self.cfg = {"tieu_de_thanh_cong_thuc": {"7": "=SUM(J8:J8)"},
                    "dong": [{"dong": 8, "C": "- Bê tông dầm (tr.175)", "E": 1,
                              "F": {"trang": 175, "nhan": "BE TONG 45 MPA"}}]}

    def test_changed_row_gets_drawing_value_old_value_and_flag(self):
        flags = apply_qs(self.ws, self.cfg, {175: PAGE}, "QS", "ghi chú")
        ws = self.ws
        self.assertEqual([ws[f"{c}8"].value for c in "EFGHI"], [1, 434.803, 1, 1, 1])
        self.assertEqual(ws["J8"].value, "=E8*F8*G8*H8*I8")
        self.assertAlmostEqual(ws["N8"].value, round(16 * 38.2 * 0.759, 3))
        self.assertIn("434,803", ws["M8"].value)
        self.assertEqual(ws["J7"].value, "=SUM(J8:J8)")
        self.assertEqual(ws["N7"].value, 999.0)
        self.assertEqual({f.dong for f in flags}, {7, 8})


@unittest.skipUnless(os.path.isfile(DT_GOC) and os.path.isfile(DU_LIEU), "thiếu dự toán gốc hoặc dữ liệu OCR Km19")
class Km19EstimateTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from examples.sua_du_toan_03a_km19_theo_ban_ve import main
        cls.tmp = tempfile.TemporaryDirectory()
        cls.dt = os.path.join(cls.tmp.name, "dt.xlsx")
        cls.p03 = os.path.join(cls.tmp.name, "03a.xlsx")
        main(["--out-du-toan", cls.dt, "--out-03a", cls.p03, "--ngay", "11/10/2026"])
        with open(CONFIG, encoding="utf-8") as f:
            cls.cfg = json.load(f)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def _cells(self, path, sheet):
        return [[c.value for c in r] for r in openpyxl.load_workbook(path)[sheet].iter_rows()]

    def test_committed_workbooks_are_exactly_what_the_script_produces(self):
        for new, committed in ((self.dt, DT_MOI), (self.p03, P03A_MOI)):
            for sheet in openpyxl.load_workbook(committed).sheetnames:
                self.assertEqual(self._cells(new, sheet), self._cells(committed, sheet), f"{committed}:{sheet}")

    def test_drawing_values(self):
        qs = openpyxl.load_workbook(self.dt)["QS_DIEN_GIAI_CHI_TIET"]
        expect = {18: 23.58, 21: 41.56, 42: 184.24, 51: 234.92, 52: 266.79, 54: 72.51, 63: 434.803, 65: 29.911,
                  74: 19.38, 78: 305.465, 80: 21.868, 87: 29.10, 88: 38.28}
        for r, v in expect.items():
            self.assertAlmostEqual(qs.cell(r, 6).value, v, places=3, msg=f"dòng {r}")
        self.assertEqual(qs["F67"].value, "=3*6.07125+3*6.06528+9*5.92008")       # 89,690 t = tr.175
        self.assertEqual(qs["F82"].value, "=1.573+78.951+5.071")
        self.assertEqual(qs["F59"].value, "=16.17783+32.55852+13.01311")                # không cắt chữ số
        self.assertEqual(qs["F58"].value, "=19.64873+0.23705")       # chủ dự án: tính W6, W7, W9 (bản vẽ ghi 0 kg)
        for r in (22, 56, 64, 66):                                                   # hết số ghi cứng ở tiêu đề
            self.assertTrue(str(qs.cell(r, 10).value).startswith("=SUM("), r)

    def test_every_change_is_flagged_for_review(self):
        qs = openpyxl.load_workbook(self.dt)["QS_DIEN_GIAI_CHI_TIET"]
        from tools.bbs_reconcile import load_pages
        pages = load_pages(DU_LIEU)
        for item in self.cfg["dong"]:
            r, old = item["dong"], qs.cell(item["dong"], 14).value
            new = item["E"] * (resolve_f(item["F"], pages)[0] if "F" in item else 0.0)
            if item.get("nghi_van") or old is None or abs(new - old) > max(0.005, 0.001 * abs(old)):
                self.assertTrue(qs.cell(r, 13).value, f"dòng {r} thiếu nghi vấn")
        nv = openpyxl.load_workbook(self.dt)["NGHI_VAN"]
        self.assertEqual(nv["I3"].value, "Đã soát / Ghim")
        missing = [nv.cell(r, 4).value for r in range(4, nv.max_row + 1) if nv.cell(r, 2).value == "(thiếu trong dự toán)"]
        self.assertIn("Cốt thép gờ lan can", missing)

    def test_03a_points_at_item_rows_of_its_own_qs_sheet(self):
        wb = openpyxl.load_workbook(self.p03)
        ws, qs = wb["THANH_TOAN_KY_PHU_LUC_03A"], wb["QS_DIEN_GIAI_CHI_TIET"]
        for r in range(8, 28):
            d = ws.cell(r, 4).value
            if not (isinstance(d, str) and d.startswith("=")):
                continue
            for qs_row in map(int, re.findall(r"QS_DIEN_GIAI_CHI_TIET!J(\d+)", d)):
                self.assertTrue(str(qs.cell(qs_row, 1).value).isdigit(), f"03a dòng {r} trỏ dòng diễn giải QS {qs_row}")
            self.assertTrue(ws.cell(r, 13).value, f"03a dòng {r} thiếu nghi vấn tỷ lệ hoàn thành")
        dt_qs = openpyxl.load_workbook(self.dt)["QS_DIEN_GIAI_CHI_TIET"]
        self.assertEqual(self._cells(self.p03, "QS_DIEN_GIAI_CHI_TIET"), [[c.value for c in r] for r in dt_qs.iter_rows()])

    def test_apply_03a_is_idempotent_on_references(self):
        wb = openpyxl.load_workbook(self.p03)
        flags = []
        apply_03a(wb, self.cfg["phu_luc_03a"], self.cfg["qs_sheet"], flags)
        self.assertFalse([f for f in flags if f.noi_dung.startswith("Sửa tham chiếu")])


if __name__ == "__main__":
    unittest.main()
