# -*- coding: utf-8 -*-
"""
Test sửa BBS theo bản vẽ thẩm định (tools/bbs_reconcile.py, examples/sua_bbs_km19_theo_ban_ve.py).

Quy tắc chủ dự án (11/10/2026): có dấu thẩm định thì lấy theo thẩm định; sai khác thì đánh dấu NGHI VẤN.
"""

import os
import tempfile
import unittest

import openpyxl

from tests.test_drawing_rebar_tables import table
from tools.bbs_loader import load_bbs
from tools.bbs_reconcile import (GIU, KHOP, THAY, THAY_MOT_PHAN, load_config, load_pages, read_bbs, reconcile,
                                 unit_changes, write_bbs)
from tools.cutting_stock_solver import rebar_unit_weight_kg_m

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
KM19_BBS_GOC = os.path.join(ROOT, "examples", "data", "km19_bbs_goc_truoc_sua.xlsx")
KM19_CONFIG = os.path.join(ROOT, "examples", "data", "km19_doi_chieu_ban_ve.json")
KM19_BBS_MOI = os.path.join(ROOT, "examples", "HO_SO_CAU_KM19_529", "02_XUONG_TIEN_CHE_COT_THEP",
                            "04_Thong_Ke_Thep_Chi_Tiet_BBS_396_Dong.xlsx")

HDR = ["TÊN THANH", "ĐƯỜNG KÍNH", "CHIỀU DÀI MM", "SỐ LƯỢNG", "KG/M", "KHỐI LƯỢNG KG"]


def page(no, rows, total):
    return {"page": no, "tables": [table(HDR, rows + [["", "TỔNG HỢP KHỐI LƯỢNG CỐT THÉP", "", "", "", str(total)]])]}


def write_old_bbs(path, items):
    """items: [(mục, [(ký hiệu, Ø, dài m, SL, số CK, kg/m, nhóm)])] — đủ cột như BBS thật."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "THONG_KE_THEP_CHI_TIET"
    ws["B2"], ws["B3"] = "DỰ ÁN MẪU", "BẢNG THỐNG KÊ CỐT THÉP"
    ws.append([])
    for c, h in enumerate(["TT", "Hạng mục kết cấu", "Bộ phận cấu kiện", "Ký hiệu thanh", "Đường kính Ø (mm)",
                           "Mác thép", "Hình dạng thanh", "Chiều dài 1 thanh (m)", "Số thanh / cấu kiện",
                           "Số cấu kiện", "Tổng số thanh", "Tổng chiều dài (m)", "Trọng lượng 1m (kg/m)",
                           "Tổng khối lượng (kg)", "Tổng khối lượng (Tấn)", "Phân nhóm Ø", "Ghi chú"], start=1):
        ws.cell(5, c, h)
    ws.cell(6, 1, "02. KẾT CẤU")
    r, tt = 7, 0
    for name, rows in items:
        for mark, d, L, n, ck, w, nhom in rows:
            tt += 1
            for c, v in enumerate([tt, name, name, mark, d, "CB400-V", "S01", L, n, ck, f"=I{r}*J{r}", f"=H{r}*K{r}",
                                   w, f"=L{r}*M{r}", f"=N{r}/1000", nhom, ""], start=1):
                ws.cell(r, c, v)
            r += 1
    ws.cell(r, 1, "TỔNG CỘNG")
    wb.save(path)


class ReconcileRulesTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        bbs = os.path.join(cls.tmp.name, "bbs.xlsx")
        write_old_bbs(bbs, [
            ("Mố X", [("F1", 16, 5.0, 10, 1, 1.58, "10<D<=18mm")]),                       # sai so với bản vẽ
            ("Thân Y", [("T1", 20, 6.0, 10, 1, 2.47, "D>18mm")]),                         # khớp bản vẽ
            ("Dầm Z", [("B1", 16, 10.0, 50, 15, 1.58, "10<D<=18mm")]),                    # chưa biết số phiến
            ("LTN", [("Z1", 128, 0.002, 525, 1, 106.18, "D>18mm")]),                      # dòng rác
        ])
        pages = {
            1: page(1, [["A1", "D16", "2000", "20", "1.580", "63.20"], ["A2", "D20", "3000", "10", "2.470", "74.10"]], 137.30),
            2: page(2, [["T1", "D20", "6000", "10", "2.470", "148.20"], ["T9", "D10", "100", "1", "0.617", "0.06"]], 148.26),
            3: page(3, [["B1", "D16", "10000", "10", "1.580", "158.00"], ["B2", "D12", "1000", "10", "0.888", "8.88"]], 166.88),
            4: page(4, [["K1", "D16", "1000", "5", "1.580", "7.90"], ["K2", "D12", "1000", "5", "0.888", "4.44"]], 12.34),
        }
        cfg = {"khop": {"tuong_doi": 0.003, "tuyet_doi_kg": 5.0}, "hang_muc": [
            {"ten": "Mố X", "trang": [{"so": 1, "he_so": 1, "bo_phan": "Mố X"}]},
            {"ten": "Thân Y", "trang": [{"so": 2, "he_so": 1, "bo_phan": "Thân Y"}]},
            {"ten": "Dầm Z", "trang": [{"so": 3, "he_so": None, "bo_phan": "Dầm Z"}], "ghi_chu_he_so": "chưa biết số phiến"},
            {"ten": "LTN", "tong_ban_ve_kg": 500.0, "trang_tham_chieu": [9], "ghi_chu": "nhập theo PDF"},
            {"ten": "Khe", "moi": True, "chen_sau": "Mố X", "trang": [{"so": 4, "he_so": 2, "bo_phan": "Khe"}]},
        ]}
        titles, cls.old_sections = read_bbs(bbs)
        cls.sections, cls.reports, cls.flags = reconcile(cls.old_sections, cfg, pages)
        cls.by = {r.hang_muc: r for r in cls.reports}
        cls.items = {k: v for s in cls.sections for k, v in s.items.items()}
        cls.out = os.path.join(cls.tmp.name, "moi.xlsx")
        cls.stats = write_bbs(cls.out, titles, cls.sections, cls.reports, cls.flags, unit_changes(cls.old_sections), "ĐÃ SỬA")

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_actions(self):
        self.assertEqual(self.by["Mố X"].hanh_dong, THAY)
        self.assertEqual(self.by["Thân Y"].hanh_dong, KHOP)
        self.assertEqual(self.by["Dầm Z"].hanh_dong, GIU)
        self.assertEqual(self.by["LTN"].hanh_dong, GIU)
        self.assertEqual(self.by["Khe"].hanh_dong, THAY)

    def test_replaced_rows_come_from_drawing_with_project_unit_weight(self):
        rows = self.items["Mố X"]
        self.assertEqual([(r.mark, r.d, r.L_m, r.n_per) for r in rows], [("A1", 16, 2.0, 20), ("A2", 20, 3.0, 10)])
        self.assertEqual([r.unit for r in rows], [rebar_unit_weight_kg_m(16), rebar_unit_weight_kg_m(20)])
        self.assertTrue(all(any(n.startswith("Thay theo tr.1") for n in r.nghi_van) for r in rows))

    def test_matching_item_keeps_bbs_rows_without_flag(self):
        rows = self.items["Thân Y"]
        self.assertEqual([r.mark for r in rows], ["T1"])
        self.assertEqual(rows[0].nghi_van, [])
        self.assertEqual(rows[0].unit, rebar_unit_weight_kg_m(20))          # chỉ đổi kg/m theo quy ước

    def test_unknown_member_count_keeps_bbs_and_flags_every_row(self):
        rows = self.items["Dầm Z"]
        self.assertEqual([r.mark for r in rows], ["B1"])
        self.assertTrue(rows[0].nghi_van and "chưa biết số cấu kiện" in rows[0].nghi_van[0])

    def test_broken_rows_removed_and_listed(self):
        rows = self.items["LTN"]
        self.assertEqual([r.placeholder for r in rows], [True])
        self.assertTrue(any("đã loại" in f.noi_dung for f in self.flags if f.hang_muc == "LTN"))

    def test_new_item_inserted_after_its_anchor(self):
        self.assertEqual(list(self.items)[:2], ["Mố X", "Khe"])
        self.assertEqual([r.n_ck for r in self.items["Khe"]], [2, 2])           # hệ số 2 cấu kiện

    def test_written_file_loads_cleanly_and_has_review_sheets(self):
        res = load_bbs(self.out, sheet="THONG_KE_THEP_CHI_TIET")
        self.assertEqual(res.errors, [])
        self.assertEqual(sorted(d.mark.split(":")[-1] for d in res.demands), ["A1", "A2", "B1", "K1", "K2", "T1"])
        wb = openpyxl.load_workbook(self.out)
        self.assertEqual(wb.sheetnames, ["THONG_KE_THEP_CHI_TIET", "TONG_HOP_THEP_THEO_DUONG_KINH",
                                         "DOI_CHIEU_BAN_VE", "NGHI_VAN"])
        self.assertEqual(wb["NGHI_VAN"]["I3"].value, "Đã soát / Ghim")
        self.assertGreater(self.stats["muc_nghi_van"], 3)


@unittest.skipUnless(os.path.isfile(KM19_BBS_GOC), "thiếu BBS gốc Km19")
class Km19ReproducibleTest(unittest.TestCase):
    def test_committed_bbs_is_exactly_what_the_script_produces(self):
        from examples.sua_bbs_km19_theo_ban_ve import main
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, "bbs.xlsx")
            main(["--out", out, "--ngay", "11/10/2026"])
            for sheet in ("THONG_KE_THEP_CHI_TIET", "TONG_HOP_THEP_THEO_DUONG_KINH", "DOI_CHIEU_BAN_VE", "NGHI_VAN"):
                a = [[c.value for c in r] for r in openpyxl.load_workbook(out)[sheet].iter_rows()]
                b = [[c.value for c in r] for r in openpyxl.load_workbook(KM19_BBS_MOI)[sheet].iter_rows()]
                self.assertEqual(a, b, sheet)

    def test_km19_actions(self):
        cfg = load_config(KM19_CONFIG)
        _, secs = read_bbs(KM19_BBS_GOC)
        _, reports, _ = reconcile(secs, cfg, load_pages(os.path.join(ROOT, cfg["du_lieu_ban_ve"])))
        actions = {r.hang_muc: r.hanh_dong for r in reports}
        self.assertEqual(actions["Mố M1"], THAY_MOT_PHAN)                  # BBS cũ lấy nhầm số của mố M2
        self.assertEqual(actions["Trụ T1, T2 (Xà mũ trụ)"], THAY)
        self.assertEqual(actions["Trụ T1 (Thân trụ)"], KHOP)
        self.assertEqual(actions["Dầm chủ Super-T (15 phiến L=38.2m)"], GIU)


if __name__ == "__main__":
    unittest.main()
